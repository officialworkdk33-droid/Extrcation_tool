# -*- coding: utf-8 -*-
"""Document extraction analysis - what was read from every document, how complete it is, whether its own figures add
up, and whether the documents agree with each other (invoice <-> packing list <-> B/L / AWB / customs instructions).

analyse(docs, ships) -> dict(documents=[...], shipments=[...], summary={...})
  documents : per file - type, reader, key fields, line completeness (% of lines with each field), checks, score 0-100
  shipments : cross-document checks + where each permit value came from (HS / COO / net weight / description)
Used by the web page (Check & files > Document analysis), the RAW_DATA.xlsx sheet DOCUMENT_ANALYSIS and the log."""
import collections
import re

KIND_NAME = {'INV': 'Invoice', 'PL': 'Packing list', 'AWB': 'Waybill / B/L', 'BL': 'Bill of lading', 'EMAIL': 'E-mail instructions',
             'OTHER': 'Not for the permit', '': 'Not recognised'}
# fields that matter per document kind: (key, label, weight)
INV_LINE_FIELDS = [('desc', 'Description', 3), ('qty', 'Quantity', 3), ('amount', 'Amount', 3), ('price', 'Unit price', 1),
                   ('item', 'Item / part no.', 1), ('hs', 'HS code', 2), ('coo', 'Origin (COO)', 2), ('uom', 'UOM', 1),
                   ('nw', 'Net weight', 0), ('brand', 'Brand', 0)]
PL_LINE_FIELDS = [('item', 'Item / part no.', 2), ('qty', 'Quantity', 2), ('nw', 'Net weight', 3), ('gw', 'Gross weight', 2),
                  ('desc', 'Description', 1)]
INV_HEAD = [('invoice_no', 'Invoice no.', 3), ('invoice_date', 'Invoice date', 1), ('currency', 'Currency', 3),
            ('total', 'Invoice total', 3), ('incoterm', 'Incoterm', 1), ('supplier', 'Seller', 1),
            ('invoice_to', 'Buyer / consignee', 1), ('hawb', 'HAWB / B/L', 1)]
PL_HEAD = [('total_gw', 'Gross weight', 3), ('total_nw', 'Net weight', 2), ('packages', 'Packages', 2), ('invoice_refs', 'Invoice ref.', 1),
           ('hawb', 'HAWB / B/L', 1)]
AWB_HEAD = [('hawb', 'HAWB / HBL', 3), ('mawb', 'MAWB / OBL', 1), ('gw', 'Gross weight', 3), ('pieces', 'Packages', 2),
            ('consignee', 'Consignee', 0.5), ('shipper', 'Shipper', 0.5), ('departure', 'From', 0.5), ('destination', 'To', 0.5)]
SRC_FIELDS = [('hs_src', 'HS code'), ('coo_src', 'Origin'), ('nw_src', 'Net weight'), ('uom_src', 'HS UOM'), ('casc_src', 'CASC code')]


def _f(v):
    try:
        if v in (None, ''):
            return None
        return float(str(v).replace(',', ''))
    except (TypeError, ValueError):
        return None


def _has(v):
    if v is None:
        return False
    if isinstance(v, (list, tuple, dict, set)):
        return len(v) > 0
    if isinstance(v, (int, float)):
        return True
    return str(v).strip() not in ('', '0', 'None')


def _fmt(v, n=2):
    v = _f(v)
    return '' if v is None else '{:,.{n}f}'.format(v, n=n)


def _near(a, b, rel=0.01, ab=0.05):
    if a is None or b is None:
        return True
    return abs(a - b) <= max(ab, rel * max(abs(a), abs(b)))


def _norm(name):
    n = re.sub(r'[^A-Z0-9 ]+', ' ', str(name or '').upper())
    n = re.sub(r'\b(PTE|LTD|PRIVATE|LIMITED|CO|COMPANY|INC|CORP|GMBH|AG|BV|SA|SG|SINGAPORE|THE|AND)\b', ' ', n)
    return set(w for w in n.split() if len(w) > 1)


def _same_party(a, b):
    A, B = _norm(a), _norm(b)
    if not A or not B:
        return None
    return len(A & B) / float(min(len(A), len(B))) >= 0.5


def _kind_label(d):
    k = d.get('kind') or ''
    r = (d.get('reader') or d.get('doc_type') or '').lower()
    if k == 'BL':
        return 'Bill of lading'
    if k == 'EMAIL':
        return 'Booking form' if 'booking' in r else 'E-mail instructions'
    if k == 'AWB':
        if 'customs instruction' in r:
            return 'Customs instructions'
        if 'arrival' in r:
            return 'Arrival notice'
        if 'b/l' in r or 'bill' in r or 'sea' in r or 'hbl' in r:
            return 'Bill of lading'
        return 'Air waybill / B/L'
    return KIND_NAME.get(k, k or 'Not recognised')


def _coverage(lines, fields):
    n = len(lines)
    out = []
    for k, lab, w in fields:
        got = sum(1 for l in lines if _has(l.get(k)))
        out.append(dict(field=k, label=lab, got=got, n=n, pct=round(100.0 * got / n) if n else 0, weight=w))
    return out


def _head(d, fields):
    out = []
    for k, lab, w in fields:
        v = d.get(k)
        if k == 'invoice_refs' and isinstance(v, list):
            v = ', '.join(str(x) for x in v[:6])
        out.append(dict(field=k, label=lab, value=v if _has(v) else '', ok=_has(v), weight=w))
    return out


def _score(head, cov, checks):
    tot = sum(h['weight'] for h in head) + sum(c['weight'] for c in cov)
    got = sum(h['weight'] for h in head if h['ok']) + sum(c['weight'] * c['pct'] / 100.0 for c in cov)
    s = 100.0 * got / tot if tot else 100.0
    for lvl, _ in checks:
        s -= {'ERROR': 15, 'MISMATCH': 10, 'CHECK': 4, 'MISSING': 6}.get(lvl, 0)
    return max(0, min(100, int(round(s))))


def _norm_doc(d):
    """MSD tool documents -> the common shape (BL = waybill, PL materials = lines, pallets = packages)."""
    k = d.get('kind') or ''
    if k == 'BL':
        d = dict(d, hawb=d.get('bl_no') or d.get('hawb'), pieces=d.get('packages'), departure=d.get('pol'),
                 destination=d.get('pod'), consignee=d.get('consignee') or '', mawb=d.get('mawb') or '',
                 cargo=dict(d.get('cargo') or {}, containers=[c.get('no') for c in d.get('containers') or [] if c.get('no')]))
    if k == 'PL' and not d.get('lines') and isinstance(d.get('materials'), dict):
        d = dict(d, lines=[dict(m) for m in d['materials'].values()], packages=d.get('packages') or d.get('pallets'),
                 invoice_refs=d.get('invoice_refs') or ([d['delivery']] if d.get('delivery') else []))
    if k == 'INV' and ('date' in d or 'exporter' in d):
        d = dict(d, invoice_date=d.get('invoice_date') or d.get('date'), supplier=d.get('supplier') or d.get('exporter'),
                 invoice_to=d.get('invoice_to') or d.get('consignee'))
    return d


def _doc(d):
    d = _norm_doc(d)
    k = d.get('kind') or ''
    if k == 'BL':
        k = 'AWB'
    lines = d.get('lines') or []
    checks = []
    head, cov = [], []
    if k == 'INV':
        head, cov = _head(d, INV_HEAD), _coverage(lines, INV_LINE_FIELDS)
        s = sum(_f(l.get('amount')) or 0 for l in lines)
        tot = _f(d.get('total'))
        if not lines:
            checks.append(('ERROR', 'no item lines read'))
        chg = (_f(d.get('charges')) or 0) + (_f(d.get('freight')) or 0)
        gv = _f(d.get('goods_value'))
        if tot is not None and lines and not _near(s, tot, 0.0005, 0.05) and chg and _near(s + chg, tot, 0.0005, 0.05):
            checks.append(('OK', 'line amounts %s + %s %s = invoice total %s' % (_fmt(s), d.get('charges_label') or 'charges', _fmt(chg), _fmt(tot))))
        elif tot is not None and lines and not _near(s, tot, 0.0005, 0.05) and gv and _near(s, gv, 0.0005, 0.05):
            checks.append(('OK', 'line amounts %s = goods value (invoice total %s includes charges)' % (_fmt(s), _fmt(tot))))
        elif tot is not None and lines and not _near(s, tot, 0.0005, 0.05):
            checks.append(('MISMATCH', 'line amounts %s <> invoice total %s (difference %s)' % (_fmt(s), _fmt(tot), _fmt(tot - s))))
        bad = []
        for l in lines:
            q, p, a = _f(l.get('qty')), _f(l.get('price')), _f(l.get('amount'))
            if q and p and a and not _near(q * p, a, 0.005, 0.05):
                per = None
                for div in (10, 100, 1000):          # price per 10 / 100 / 1000 pieces
                    if _near(q * p / div, a, 0.005, 0.05):
                        per = div
                if not per:
                    bad.append('%s (%s x %s = %s, amount %s)' % (l.get('item') or l.get('line') or '?', _fmt(q, 3), _fmt(p, 4), _fmt(q * p), _fmt(a)))
        if bad:
            checks.append(('CHECK', '%d line(s) where qty x unit price <> amount: %s' % (len(bad), '; '.join(bad[:4]))))
        neg = [l for l in lines if (_f(l.get('amount')) or 0) < 0]
        if neg:
            checks.append(('CHECK', '%d line(s) with a negative amount (discount / credit)' % len(neg)))
        zero = [l for l in lines if _f(l.get('amount')) == 0]
        if zero:
            checks.append(('CHECK', '%d line(s) with amount 0 (free of charge - declare a value)' % len(zero)))
        nd = [l for l in lines if not _has(l.get('desc'))]
        if nd:
            checks.append(('MISSING', '%d line(s) without description: %s' % (len(nd), ', '.join(str(l.get('item') or l.get('line') or '?') for l in nd[:6]))))
        if not d.get('currency'):
            checks.append(('MISSING', 'currency not found'))
        if not d.get('invoice_no'):
            checks.append(('MISSING', 'invoice number not found'))
        tnw, tgw = _f(d.get('total_nw')), _f(d.get('total_gw'))
        if tnw and tgw and tnw > tgw * 1.001:
            checks.append(('MISMATCH', 'net weight %s > gross weight %s' % (_fmt(tnw, 3), _fmt(tgw, 3))))
        lnw = sum(_f(l.get('nw')) or 0 for l in lines)
        if tnw and lnw and not _near(lnw, tnw, 0.01, 0.05):
            checks.append(('CHECK', 'line net weights %s <> total net weight %s' % (_fmt(lnw, 3), _fmt(tnw, 3))))
    elif k == 'PL':
        head, cov = _head(d, PL_HEAD), _coverage(lines, PL_LINE_FIELDS)
        tgw, tnw = _f(d.get('total_gw')), _f(d.get('total_nw'))
        lgw = sum(_f(l.get('gw')) or 0 for l in lines)
        lnw = sum(_f(l.get('nw')) or 0 for l in lines)
        if tgw and lgw and not _near(lgw, tgw, 0.01, 0.05):
            checks.append(('CHECK', 'line gross weights %s <> total %s' % (_fmt(lgw, 3), _fmt(tgw, 3))))
        if tnw and lnw and not _near(lnw, tnw, 0.01, 0.05):
            checks.append(('CHECK', 'line net weights %s <> total %s' % (_fmt(lnw, 3), _fmt(tnw, 3))))
        if (tnw or lnw) and (tgw or lgw) and (tnw or lnw) > (tgw or lgw) * 1.001:
            checks.append(('MISMATCH', 'net weight %s > gross weight %s' % (_fmt(tnw or lnw, 3), _fmt(tgw or lgw, 3))))
        if not tgw and not lgw:
            checks.append(('MISSING', 'no gross weight'))
        if not tnw and not lnw:
            checks.append(('MISSING', 'no net weight'))
    elif k == 'AWB':
        head = _head(d, AWB_HEAD)
        cg = d.get('cargo') or {}
        for key, lab in (('permit_type', 'Permit type'), ('release', 'Place of release'), ('receipt', 'Place of receipt'),
                         ('containers', 'Containers'), ('special', 'Special instructions'), ('permit_remarks', 'Permit remarks')):
            v = cg.get(key)
            if _has(v):
                head.append(dict(field=key, label=lab, value=', '.join(v) if isinstance(v, list) else v, ok=True, weight=0))
        if not _has(d.get('gw')):
            checks.append(('MISSING', 'gross weight not found'))
        if not _has(d.get('hawb')):
            checks.append(('MISSING', 'HAWB / HBL number not found'))
    elif k == 'EMAIL':
        em = d.get('email') or {}
        head = [dict(field=a, label=a.replace('_', ' ').capitalize(), value=b, ok=True, weight=0) for a, b in em.items() if not a.startswith('_')]
        for l in d.get('booking_lines') or []:
            head.append(dict(field='line', label='Booking line', value='%s - %s %s %s, GW %s' % (
                l.get('marks'), l.get('packages'), l.get('unit'), l.get('desc') + (' - ' + l['batch'] if l.get('batch') else ''), l.get('gw')), ok=True, weight=0))
        if not em:
            checks.append(('CHECK', 'no instruction fields found'))
    elif k == 'OTHER':
        checks.append(('INFO', 'not used for the permit (%s)' % '; '.join(d.get('notes') or ['other document'])))
    else:
        checks.append(('ERROR', 'not recognised - %s' % '; '.join(d.get('notes') or ['no reader for this layout'])))
    if d.get('ocr'):
        checks.append(('CHECK', 'read by OCR (scanned) - check the figures'))
    for n in d.get('notes') or []:
        if k not in ('OTHER', '') and re.match(r'(CHECK|MISSING|ERROR|MISMATCH)\b', str(n)):
            checks.append((str(n).split()[0], str(n).split(' ', 1)[1] if ' ' in str(n) else str(n)))
    score = _score(head, cov, checks) if k in ('INV', 'PL', 'AWB') else (100 if k in ('OTHER', 'EMAIL') else 0)
    keyv = (d.get('invoice_no') or ', '.join(d.get('invoice_refs') or [])[:60] if k in ('INV', 'PL') else d.get('hawb') or '')
    return dict(file=d.get('file', ''), kind=k, type=_kind_label(d) if d.get('kind') != 'BL' else 'Bill of lading', reader=d.get('reader') or d.get('doc_type') or '',
                pages=d.get('page_range') or '', key=keyv or '', lines=len(lines), currency=d.get('currency', ''),
                total=_f(d.get('total')), head=head, coverage=cov, checks=[dict(level=a, text=b) for a, b in checks],
                score=score, ocr=bool(d.get('ocr')))


def _by_type(awbs, key):
    """[(label, value)] per waybill type: one value when all copies agree (same document twice), else the sum
    (several house B/Ls of one permit)."""
    grp = collections.OrderedDict()
    for a in awbs:
        v = _f(a.get(key))
        if v:
            grp.setdefault(_kind_label(dict(a, kind='AWB')).lower() if a.get('kind') != 'BL' else 'bill of lading', []).append(v)
    out = []
    for lab, vs in grp.items():
        if len(set(round(v, 3) for v in vs)) == 1:
            out.append((lab if len(vs) == 1 else '%s (%d copies)' % (lab, len(vs)), vs[0]))
        else:
            out.append(('%s x%d (sum)' % (lab, len(vs)), sum(vs)))
    return out


def _ship(sh, docs_by_file):
    out = []
    add = lambda lvl, grp, txt: out.append(dict(level=lvl, group=grp, text=txt))
    invs, pls = sh.get('invs') or [], [_norm_doc(p) for p in sh.get('pls') or []]
    awbs = [_norm_doc(b) for b in sh['bls']] if sh.get('bls') else (sh.get('awbs') or [])
    items = sh.get('items') or []
    # ---- value
    by_cur = collections.OrderedDict()
    for i in invs:
        c = i.get('currency') or '?'
        by_cur.setdefault(c, [0.0, 0.0, 0, 0.0])
        by_cur[c][0] += _f(i.get('total')) or sum(_f(l.get('amount')) or 0 for l in i.get('lines') or [])
        by_cur[c][2] += 1
        by_cur[c][3] += (_f(i.get('charges')) or 0) + (_f(i.get('freight')) or 0)
    for it in items:
        c = it.get('currency') or (invs[0].get('currency') if invs else '?') or '?'
        by_cur.setdefault(c, [0.0, 0.0, 0, 0.0])
        by_cur[c][1] += _f(it.get('amount')) or 0
    for c, (ti, ta, n, ch) in by_cur.items():
        if n and ch and not _near(ti, ta, 0.0005, 0.05) and _near(ti, ta + ch, 0.0005, 0.05):
            add('OK', 'Value', '%s: invoice total %s = items %s + charges %s (charges: %s)' % (
                c, _fmt(ti), _fmt(ta), _fmt(ch), {'other': 'invoice total / Other Taxable Charge', 'last': 'in the last item',
                                                  'otc': 'Other Taxable Charge'}.get(sh.get('charges_mode_used') or '', 'see Invoice tab')))
        elif n and not _near(ti, ta, 0.0005, 0.05):
            add('MISMATCH', 'Value', '%s: %d invoice(s) total %s, permit items %s (difference %s) - charges / discounts / lines not read?'
                % (c, n, _fmt(ti), _fmt(ta), _fmt(ti - ta)))
        elif n:
            add('OK', 'Value', '%s: %d invoice(s) %s = permit items' % (c, n, _fmt(ti)))
    if len(by_cur) > 1:
        add('CHECK', 'Value', '%d currencies (%s) - one invoice row per currency' % (len(by_cur), ', '.join(by_cur)))
    nos = [i.get('invoice_no') for i in invs if i.get('invoice_no')]
    dup = [n for n, c in collections.Counter(nos).items() if c > 1]
    if dup:
        add('CHECK', 'Invoices', 'invoice no. read more than once (copy / UPDATE version?): %s' % ', '.join(dup[:6]))
    # ---- invoice <-> packing list
    if invs and pls:
        refs = set()
        for p in pls:
            refs |= set(str(x) for x in (p.get('invoice_refs') or []))
        alias = {}                                   # invoice no. + its delivery / order numbers (MSD lists refer to the delivery)
        for i in invs:
            for k_ in ('invoice_no', 'delivery', 'sap_order', 'sto', 'order_no'):
                if i.get(k_):
                    alias[str(i[k_])] = i.get('invoice_no')
        no_pl = [n for n in nos if refs and not any(alias.get(r) == n for r in refs)]
        if refs and no_pl:
            add('CHECK', 'Invoice / PL', '%d invoice(s) with no packing list referring to them: %s' % (len(no_pl), ', '.join(no_pl[:8])))
        stray = [r for r in refs if r and nos and r not in alias]
        if stray:
            add('CHECK', 'Invoice / PL', 'packing list refers to invoice(s) not uploaded: %s' % ', '.join(sorted(stray)[:8]))
        iq, pq = collections.defaultdict(float), collections.defaultdict(float)
        for i in invs:
            for l in i.get('lines') or []:
                if l.get('item'):
                    iq[str(l['item']).upper()] += _f(l.get('qty')) or 0
        for p in pls:
            for l in p.get('lines') or []:
                if l.get('item'):
                    pq[str(l['item']).upper()] += _f(l.get('qty')) or 0
        if iq and pq:
            both = set(iq) & set(pq)
            only_i = sorted(set(iq) - set(pq))
            diff = ['%s (inv %s / PL %s)' % (k, _fmt(iq[k], 0), _fmt(pq[k], 0)) for k in sorted(both) if not _near(iq[k], pq[k], 0.0, 0.001)]
            add('OK' if not only_i and not diff else 'CHECK', 'Invoice / PL', '%d of %d invoice item(s) found on the packing list%s%s' % (
                len(both), len(iq), ('; not on PL: ' + ', '.join(only_i[:6])) if only_i else '',
                ('; qty differs: ' + '; '.join(diff[:5])) if diff else ''))
    elif invs and not pls:
        add('CHECK', 'Invoice / PL', 'no packing list - weights from the invoice / B/L (net weight may be estimated)')
    # ---- gross weight / packages from every source
    gws = []
    for p in pls:
        if _f(p.get('total_gw')):
            gws.append(('PL ' + p.get('file', ''), _f(p['total_gw'])))
    pl_gw = sum(v for _, v in gws)
    srcs = []
    if pl_gw:
        srcs.append(('packing list(s)', pl_gw))
    ci_gw = sum(_f(i.get('total_gw')) or 0 for i in invs)
    if ci_gw:
        srcs.append(('invoice(s)', ci_gw))
    for lab, vals in _by_type(awbs, 'gw'):
        srcs.append((lab, vals))
    used = _f(sh.get('total_gw'))
    if srcs:
        ref = used or srcs[0][1]
        off = [(n, v) for n, v in srcs if not _near(v, ref, 0.01, 0.5)]
        txt = '; '.join('%s %s' % (n, _fmt(v, 3)) for n, v in srcs)
        add('MISMATCH' if off else 'OK', 'Weight', 'gross weight on the permit %s KG - %s%s' % (
            _fmt(used, 3) or '?', txt, (' - differs: ' + ', '.join(n for n, _ in off[:4])) if off else ''))
    nw_items = sum(_f(it.get('nw')) or 0 for it in items)
    if used and nw_items and nw_items > used * 1.001:
        add('MISMATCH', 'Weight', 'total net weight of the items %s > gross weight %s' % (_fmt(nw_items, 3), _fmt(used, 3)))
    pk = []
    pl_pk = sum(int(_f(p.get('packages')) or 0) for p in pls)
    if pl_pk:
        pk.append(('packing list', pl_pk))
    for lab, v in _by_type(awbs, 'pieces'):
        pk.append((lab, int(v)))
    if pk:
        vals = set(v for _, v in pk)
        add('OK' if len(vals) == 1 else 'CHECK', 'Packages', 'packages on the permit %s - %s' % (
            sh.get('packages') or '?', '; '.join('%s %s' % (n, v) for n, v in pk)))
    # ---- HAWB / parties
    hb = set(str(x.get('hawb')).upper() for x in invs + pls + awbs if x.get('hawb'))
    if len(hb) > 1 and not sh.get('house_list'):
        add('CHECK', 'References', 'several HAWB / B/L numbers in the documents: %s (permit: %s)' % (', '.join(sorted(hb)[:6]), sh.get('hawb') or '?'))
    party = sh.get('party') or ''
    exp = sh.get('direction') == 'EXPORT'
    for x in invs[:3] + awbs:
        nm = (x.get('shipper') if exp else (x.get('consignee') or x.get('invoice_to') or x.get('ship_to'))) or ''
        if isinstance(nm, list):
            nm = ' '.join(nm)
        same = _same_party(party, nm)
        if same is False:
            add('CHECK', 'Parties', '%s on %s: "%s" - permit %s "%s"' % ('shipper' if exp else 'consignee', x.get('file', '')[:40],
                                                                         str(nm)[:60], 'exporter' if exp else 'importer', party))
    # ---- provenance of the permit values
    prov = []
    for key, lab in SRC_FIELDS:
        c = collections.Counter(str(it.get(key) or 'blank') for it in items)
        if items and (len(c) > 1 or 'blank' not in c):
            prov.append(dict(label=lab, counts=c.most_common(6)))
    hs_doc_diff = [it for it in items if it.get('doc_hs') and it.get('hs') and re.sub(r'\D', '', str(it['doc_hs']))[:6] != re.sub(r'\D', '', str(it['hs']))[:6]]
    if hs_doc_diff:
        add('CHECK', 'HS code', '%d item(s) with a different HS than the invoice (customs instructions / master / list): %s' % (
            len(hs_doc_diff), '; '.join('%s inv %s -> %s' % (it.get('item') or '?', it.get('doc_hs'), it.get('hs')) for it in hs_doc_diff[:4])))
    no_hs = [it for it in items if not it.get('hs')]
    if no_hs:
        add('MISSING', 'HS code', '%d item(s) without HS code' % len(no_hs))
    no_coo = [it for it in items if not it.get('coo')]
    if no_coo:
        add('MISSING', 'Origin', '%d item(s) without country of origin' % len(no_coo))
    est = [it for it in items if 'estimat' in str(it.get('nw_src') or '').lower() or 'share' in str(it.get('nw_src') or '').lower()]
    if est:
        add('CHECK', 'Weight', '%d item(s) with net weight shared / estimated (not from a document line)' % len(est))
    used_files = set(x.get('file') for x in invs + pls + awbs)
    return dict(key=sh.get('key'), hawb=sh.get('hawb', ''), party=party, direction=sh.get('direction', ''), items=len(items),
                invoices=len(invs), pls=len(pls), waybills=len(awbs), files=sorted(f for f in used_files if f),
                checks=out, provenance=prov)


def _merge(group):
    """Several shipments (e.g. MSD deliveries) under one B/L = one permit: documents, items, weights added up."""
    m = dict(group[0])
    for f in ('invs', 'pls', 'items'):
        m[f] = [x for s in group for x in s.get(f) or []]
    for f in ('bls', 'awbs'):
        seen, out = set(), []
        for s in group:
            for x in s.get(f) or []:
                if x.get('file') not in seen:
                    seen.add(x.get('file'))
                    out.append(x)
        m[f] = out
    for f in ('total_gw', 'packages'):
        vals = [_f(s.get(f)) for s in group if _f(s.get(f)) is not None]
        m[f] = sum(vals) if vals else None
    if m.get('packages') is not None:
        m['packages'] = int(m['packages'])
    m['key'] = ' + '.join(str(s.get('key')) for s in group)
    return m


def analyse(docs, ships):
    ds = [_doc(d) for d in docs or []]
    by_file = {d.get('file'): d for d in docs or []}
    groups = collections.OrderedDict()               # shipments of one permit (same HAWB / B/L) compared together
    for sh in (ships or {}).values():
        if sh.get('include', True):
            groups.setdefault(sh.get('hawb') or sh.get('key'), []).append(sh)
    shs = [_ship(g[0] if len(g) == 1 else _merge(g), by_file) for g in groups.values()]
    used = set()
    for s in shs:
        used |= set(s['files'])
    for d in ds:
        d['used'] = d['file'] in used
        if d['kind'] in ('INV', 'PL', 'AWB') and not d['used'] and shs:
            d['checks'].append(dict(level='CHECK', text='not matched to any shipment of this job'))
    kinds = collections.Counter(d['type'] for d in ds)
    lv = collections.Counter(c['level'] for d in ds for c in d['checks']) + collections.Counter(c['level'] for s in shs for c in s['checks'])
    sc = [d['score'] for d in ds if d['kind'] in ('INV', 'PL', 'AWB')]
    summary = dict(documents=len(ds), kinds=kinds.most_common(), score=int(round(sum(sc) / len(sc))) if sc else 0,
                   errors=lv.get('ERROR', 0) + lv.get('MISMATCH', 0), checks=lv.get('CHECK', 0) + lv.get('MISSING', 0),
                   ok=lv.get('OK', 0), lines=sum(d['lines'] for d in ds))
    return dict(documents=ds, shipments=shs, summary=summary)


def log_lines(a):
    s = a['summary']
    out = ['Document analysis: %d document(s) (%s), extraction score %d%%, %d mismatch / error, %d to check' % (
        s['documents'], ', '.join('%d %s' % (n, k.lower()) for k, n in s['kinds']), s['score'], s['errors'], s['checks'])]
    for sh in a['shipments']:
        for c in sh['checks']:
            if c['level'] in ('MISMATCH', 'ERROR'):
                out.append('  %s %s: %s' % (c['level'], c['group'], c['text']))
    return out


def write_sheet(path, a):
    """DOCUMENT_ANALYSIS (+ DOCUMENT_CHECKS) sheets added to an existing workbook (RAW_DATA.xlsx / CHECK_REPORT)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.load_workbook(path)
    for n in ('DOCUMENT_ANALYSIS', 'DOCUMENT_CHECKS'):
        if n in wb.sheetnames:
            del wb[n]
    ws = wb.create_sheet('DOCUMENT_ANALYSIS', 0)
    hdr = ['File', 'Type', 'Reader', 'Pages', 'Invoice / ref.', 'Lines', 'Currency', 'Total', 'Score %', 'Used',
           'Fields found', 'Fields missing', 'Line completeness', 'Checks']
    ws.append(hdr)
    for d in a['documents']:
        ws.append([d['file'], d['type'], d['reader'], d['pages'], d['key'], d['lines'], d['currency'], d['total'], d['score'],
                   'yes' if d.get('used') else 'no',
                   ', '.join('%s %s' % (h['label'], h['value']) for h in d['head'] if h['ok'])[:500],
                   ', '.join(h['label'] for h in d['head'] if not h['ok']),
                   ', '.join('%s %d%%' % (c['label'], c['pct']) for c in d['coverage']),
                   '\n'.join('%s %s' % (c['level'], c['text']) for c in d['checks'])[:1500]])
    ws2 = wb.create_sheet('DOCUMENT_CHECKS', 1)
    ws2.append(['HAWB / shipment', 'Importer / exporter', 'Level', 'Group', 'Check'])
    for s in a['shipments']:
        for c in s['checks']:
            ws2.append([s['hawb'], s['party'], c['level'], c['group'], c['text']])
        for p in s['provenance']:
            ws2.append([s['hawb'], s['party'], 'SOURCE', p['label'], ', '.join('%s: %d' % (k, n) for k, n in p['counts'])])
    fills = {'ERROR': 'F8D7DA', 'MISMATCH': 'F8D7DA', 'CHECK': 'FFF3CD', 'MISSING': 'FFF3CD', 'OK': 'D1E7DD'}
    for w, widths in ((ws, [38, 20, 28, 8, 22, 7, 9, 14, 8, 6, 60, 30, 60, 90]), (ws2, [22, 34, 10, 14, 110])):
        for c in w[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='2E3036')
        for i, wd in enumerate(widths):
            w.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
        w.freeze_panes = 'A2'
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
    for row in ws2.iter_rows(min_row=2):
        f = fills.get(row[2].value)
        if f:
            row[2].fill = PatternFill('solid', fgColor=f)
    for row in ws.iter_rows(min_row=2):
        v = row[8].value or 0
        row[8].fill = PatternFill('solid', fgColor='D1E7DD' if v >= 85 else 'FFF3CD' if v >= 60 else 'F8D7DA')
    wb.save(path)
