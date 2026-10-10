# -*- coding: utf-8 -*-
"""MSD (Merck Sharp & Dohme) EXPORT permits on the Permit Server.

The reading rules are the MSD PERMIT TOOL's own (msd_engine.py - not changed). This file only connects them to the
server, which expects the same functions as the other tools (parse_pdf, build_shipments, resolve, generate ...):
  - invoices + packing lists + bill of lading, matched by Delivery # / order no / B/L delivery list
  - DHL instruction e-mail: paste it in "Cargo details (pre-alert / e-mail)" or upload it (.txt / .eml / .pdf)
  - one Excel per B/L (HAWB), CHECK_REPORT, EXCEL_ALL.zip; the server adds FrtImport XML + RAW_DATA
  - HS list = the server's UPDATE HS CODE list; CASC rule: IDAMISC when the HS has it, else MISC when controlled
"""
import os, re, collections

import msd_engine as M
import danfoss_engine as _D            # shared helpers: HS list (json), HS candidates, CA product codes, cargo text

TOOL_KEY = 'msd'

# ---------------------------------------------------------------- what the server reads directly
PERMIT_SCHEME = M.PERMIT_SCHEME
DEFAULT_UOM_RULES = M.DEFAULT_UOM_RULES
DEFAULT_DESC_TEMPLATE = M.DEFAULT_DESC_TEMPLATE
DEFAULT_REMARKS = M.DEFAULT_REMARKS
OCR = dict(enabled=False, cmd='', status='')        # MSD documents are text PDFs
fnum, fmt_thou, fmt_plain, h8 = M.fnum, M.fmt_thou, M.fmt_plain, M.h8
set_uom_rules = M.set_uom_rules
template_lists = M.template_lists
load_master, save_master = M.load_master, M.save_master
file_name_for, ship_label = M.file_name_for, M.ship_label
Countries = M.Countries
hs_candidates = _D.hs_candidates


def set_tesseract(cmd):
    pass


def load_hs_list(path):
    """The server's HS list (hs_data.json): {hs: (matched, uom, desc, info)} - MSD uses [0] code and [1] UOM."""
    return _D.load_hs_list(path)


def parse_cargo(text):
    """The "Cargo details" box = the DHL instruction e-mail (CARRIER, CARRIER BOOKING REF, PLACE OF RELEASE ...)."""
    em = M.parse_email(text or '')
    if text and not em:                               # a normal pre-alert (MAWB / VESSEL / ETA ...) also works
        try:
            c = _D.parse_cargo(text)
            em = {k: v for k, v in (('booking_ref', c.get('mawb')), ('job_no', c.get('job_no')),
                                    ('permit_type', c.get('permit_type'))) if v}
        except Exception:
            pass
    return em


def make_ctx(cfg, lists):
    return dict(master=M.load_master(cfg.get('master')) if cfg.get('master') else {},
                hs_list=load_hs_list(cfg.get('hs_list')) if cfg.get('hs_list') else {},
                countries=M.Countries(lists.get('List-Country', [])), default_uom=cfg.get('default_uom') or 'KGM',
                blank_hs_not_in_list=cfg.get('blank_hs_not_in_list'),
                desc_template=cfg.get('desc_template') or M.DEFAULT_DESC_TEMPLATE, desc_max=cfg.get('desc_max', 512),
                brand=cfg.get('brand', ''), email=cfg.get('cargo') or {}, cargo=cfg.get('cargo') or {},
                control_casc=cfg.get('control_casc', 'MISC'), hs_map=cfg.get('hs_map') or {})


# ---------------------------------------------------------------- reading
SUPPORT_RE = re.compile(r'Certificate\s+of\s+Analysis|Batch\s+Certificate|Certificate\s+of\s+Conformance|'
                        r'To\s+Whom\s+It\s+May\s+Concern|Reference\s+Copy|Electronic\s+Signatures|TEMPTALE|'
                        r'Temperature\s+Excursion|Material\s+Safety|Safety\s+Data\s+Sheet', re.I)


def _other(path, why):
    return [dict(kind='OTHER', file=os.path.basename(path), path=path, reader='not for the permit', notes=[why])]


def _xl_rows(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=False, data_only=True)
        out = []
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                v = [x for x in r if x not in (None, '')]
                if v:
                    out.append(v)
        return out
    except Exception:
        try:
            import xlrd
            bk = xlrd.open_workbook(path)
            return [[x for x in sh.row_values(i) if x not in (None, '')] for sh in bk.sheets() for i in range(sh.nrows)]
        except Exception:
            return []


def parse_booking_form(path):
    """MSD 'OCEAN SHIPMENT BOOKING FORM' (Excel): COUNTRY OF ORIGIN, SG HTS CODE, ETD, freight terms, port of discharge,
    one line per PO / batch (packages, description - batch, gross weight). Used like the DHL instruction e-mail."""
    rows = _xl_rows(path)
    flat = [str(x).strip() for r in rows for x in r]
    if not any(re.search(r'SHIPMENT\s+BOOKING\s+FORM|BOOKING\s+FORM', x, re.I) for x in flat[:20]):
        return None
    em, lines = {}, []
    import datetime as _dt

    def after(label_re):
        for r in rows:
            for i, x in enumerate(r):
                if re.search(label_re, str(x), re.I):
                    rest = [y for y in r[i + 1:] if str(y).strip()]
                    if rest:
                        return rest[0]
        return None
    v = after(r'^\s*COUNTRY\s+OF\s+ORIGIN')
    if v and re.fullmatch(r'[A-Za-z]{2}', str(v).strip()):
        em['coo'] = str(v).strip().upper()
    v = after(r'HTS\s+CODE|HS\s+CODE')
    if v and re.sub(r'\D', '', str(v)):
        em['hs'] = re.sub(r'\D', '', str(int(v) if isinstance(v, float) else v))
    for key, lab in (('etd', r'^\s*ETD\s*:?$'), ('position_date', r'POSITION\s+DATE'), ('booking_date', r'BOOKING\s+DATE')):
        v = after(lab)
        if isinstance(v, (_dt.datetime, _dt.date)):
            em[key] = v.strftime('%d/%m/%Y')
        elif v:
            em[key] = str(v)
    v = after(r'FREIGHT\s+TERMS')
    if v:
        em['freight_terms'] = str(v).strip()
    for i, r in enumerate(rows):                          # PORT OF DISCHARGE | PLACE OF DELIVERY ... then values
        if any(re.match(r'PORT\s+OF\s+DISCHARGE', str(x), re.I) for x in r) and i + 1 < len(rows):
            em['pod'] = str(rows[i + 1][0]).strip()
        if any(re.match(r'MARKS\s*&\s*NUMBER', str(x), re.I) for x in r):
            for r2 in rows[i + 1:i + 12]:
                if len(r2) >= 4 and isinstance(r2[1], (int, float)) and any(isinstance(x, (int, float)) and x > 1 for x in r2[3:]):
                    m = re.match(r'(.*?)\s*-\s*([A-Z]\d{5,})\s*$', str(r2[3]))
                    lines.append(dict(marks=str(r2[0]), packages=int(r2[1]), unit=str(r2[2]), desc=(m.group(1) if m else str(r2[3])).strip(),
                                      batch=m.group(2) if m else '', gw=next((x for x in r2[4:] if isinstance(x, (int, float))), None)))
    m = next((re.search(r'DECLARATION\s+VALUE\s*-?\s*(?:EST\.?)?\s*([A-Z]{3})\s*([\d,.]+)', x, re.I) for x in flat
              if re.search(r'DECLARATION\s+VALUE', x, re.I)), None)
    if m:
        em['est_value'] = '%s %s' % (m.group(1).upper(), m.group(2))
    em['shipment_type'] = em.get('shipment_type') or ('FCL' if any(re.search(r'REEFER\s+CONTAINER|CONTAINER\s+TYPE', x, re.I) for x in flat) else '')
    em = {k: v for k, v in em.items() if v}
    return dict(kind='EMAIL', file=os.path.basename(path), path=path, email=em, reader='MSD ocean booking form',
                booking_lines=lines, notes=[], text='')


def parse_pdf(path, countries=None, log=None):
    """One file -> [document]. PDFs: invoice / packing list / bill of lading; an e-mail (.txt .eml .msg or a PDF that
    is only the e-mail) -> EMAIL document used like the pasted instructions; the MSD ocean booking form (Excel) ->
    instructions too (COO, SG HTS code, ETD); batch certificates, QA letters, TempTale notes -> not for the permit."""
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.txt', '.eml', '.msg'):
        text = M.read_email_file(path)
        return [dict(kind='EMAIL', file=os.path.basename(path), path=path, email=M.parse_email(text), text=text)]
    if ext in ('.xlsx', '.xlsm', '.xls'):
        d = parse_booking_form(path)
        return [d] if d else _other(path, 'Excel file that is not an MSD booking form - not used')
    if ext in ('.docx', '.doc'):
        return _other(path, 'Word document (instructions / notes) - not used for the permit')
    if ext != '.pdf':
        return [dict(kind='', file=os.path.basename(path), path=path, notes=['Not a PDF / e-mail - skipped'])]
    d = M.parse_pdf(path, countries)
    if not d.get('kind'):
        em = M.parse_email(d.get('text', ''))
        if len(em) >= 3:
            return [dict(kind='EMAIL', file=d['file'], path=path, email=em, text=d.get('text', ''))]
        if SUPPORT_RE.search(d.get('text', '') or ''):
            t = d.get('text', '')
            what = ('batch certificate (certificate of analysis)' if re.search(r'Batch\s+Certificate|Certificate\s+of\s+Analysis', t, re.I)
                    else 'quality / approval letter' if re.search(r'To\s+Whom|Reference\s+Copy|Electronic\s+Signatures', t, re.I)
                    else 'supporting document')
            return _other(path, '%s - not used for the permit' % what)
    if d.get('kind') == 'BL':                         # B/L number cut on the PDF (SINC60376) - file name has it whole
        no = d.get('bl_no') or ''
        m = re.search(r'\b(%s\w{1,3})\b' % re.escape(no), re.sub(r'[_\-.]', ' ', d['file'])) if no else None
        if m and m.group(1) != no:
            d['notes'] = (d.get('notes') or []) + ['B/L no. printed %s - file name has %s (used)' % (no, m.group(1))]
            d['bl_no'] = d['hawb'] = m.group(1)
    d.pop('text', None)
    return [d]


def build_shipments(docs):
    mails = [d for d in docs if d.get('kind') == 'EMAIL']
    ships, issues = M.build_shipments([d for d in docs if d.get('kind') not in ('EMAIL', 'OTHER')])
    em = {}
    for d in sorted(mails, key=lambda d: d.get('reader') == 'MSD ocean booking form'):   # DHL e-mail first
        for k, v in (d.get('email') or {}).items():
            em.setdefault(k, v)
    dhl = {}
    for d in mails:
        if d.get('reader') != 'MSD ocean booking form':
            dhl.update(d.get('email') or {})
    bk = [k for d in mails if d.get('reader') == 'MSD ocean booking form' for k in (d.get('email') or {}) if k not in dhl]
    if bk:
        em['_booking'] = ','.join(bk)                 # fields that came from the booking form
    for sh in ships.values():
        sh['email_doc'] = dict(em)
        sh['direction'] = 'EXPORT'
        sh.setdefault('pkg_override', None)
    return ships, issues


# ---------------------------------------------------------------- rules
def _email(sh, ctx):
    em = dict(sh.get('email_doc') or {})
    em.update({k: v for k, v in (ctx.get('email') or {}).items() if v})     # pasted text wins over the uploaded file
    return em


def resolve(sh, ctx):
    for f in ('party', 'party_uen'):                  # edits on the Shipments tab -> MSD exporter fields
        if sh.get(f + '_src') == 'Entered':
            g = 'exporter' if f == 'party' else 'exporter_uen'
            sh[g], sh[g + '_src'] = sh[f], 'Entered'
    for c in sh.get('containers_edit') or []:
        c.setdefault('weight', None)
    em = _email(sh, ctx)
    M.resolve(sh, dict(ctx, email=em))
    if sh.get('pkg_override'):
        sh['pallets'] = sh['pkg_override']
    hsl = ctx.get('hs_list') or {}
    for it in sh['items']:
        it['keys'] = ['%s|%s' % (it.get('invoice', ''), it['item'])]
        it['description'] = it.get('description') or ''
        row = hsl.get(it['hs']) if it.get('hs') else None
        it['hs_row'] = row or ()
        it['pcodes'] = _D.product_codes(it['hs']) if it.get('hs') else []
        info = row[3] if row and len(row) > 3 and isinstance(row[3], dict) else {}
        it['controlled'] = (info.get('ctl_exp') or '').strip().lower() == 'control'
        it['ca'] = info.get('ca_exp', '') if it['controlled'] else ''
        if 'casc' in (sh['overrides'].get(it['keys'][0]) or {}):
            it['casc_src'] = 'Edited'
        elif it.get('casc'):
            it['casc_src'] = 'Product master'
        elif it['controlled']:
            it['casc'] = 'IDAMISC' if any(c[0] == 'IDAMISC' for c in it['pcodes']) else (ctx.get('control_casc') or 'MISC')
            it['casc_src'] = 'Rule (controlled HS)'
        it['lines'] = []
        it['doc_hs'] = it.get('doc_hs', '')
    if em.get('_booking'):                            # values from the booking form, not a DHL e-mail
        for it in sh['items']:
            for f in ('hs_src', 'coo_src'):
                if it.get(f) == 'E-mail' and f[:-4] in em['_booking'].split(','):
                    it[f] = 'Booking form'
        sh['issues_all'] = [(l, m.replace('e-mail HS CODE', 'booking form / e-mail HS CODE')) for l, m in sh.get('issues_all') or []]
    # what the server pages show
    sh['party'], sh['party_uen'] = sh.get('exporter', ''), sh.get('exporter_uen', '')
    sh['party_src'] = sh.get('exporter_src', '') or ('Invoice' if sh.get('exporter') else '')
    sh['party_uen_src'] = sh.get('exporter_uen_src', '') or ('Invoice' if sh.get('exporter_uen') else '')
    if em.get('uen') and sh.get('exporter_uen_src') != 'Entered':
        sh['party_uen'], sh['party_uen_src'] = em['uen'], 'E-mail'
    if em.get('exporter') and sh.get('exporter_src') != 'Entered':
        sh['party'], sh['party_src'] = em['exporter'], 'E-mail'
    sh['direction'] = 'EXPORT'
    sh['packages'] = sh.get('pallets')
    sh['lines_n'] = sum(len(i['lines']) for i in sh['invs'])
    sh['awbs'] = [dict(doc_type='Bill of lading', file=b['file']) for b in sh.get('bls') or []]
    sh['hawb_src'] = sh.get('hawb_src') or ''
    bl = sh.get('bl') or {}
    sh['cargo'] = {k: v for k, v in dict(
        mawb=em.get('booking_ref', ''), transport_id=bl.get('vessel', ''), flight=bl.get('voyage', ''),
        discharge_code=M.port_code(bl['pod']) if bl.get('pod') else '', discharge_name=bl.get('pod', ''),
        permit_type=em.get('permit_type', ''), job_no=em.get('job_no', ''), carrier=em.get('carrier', ''),
        release=em.get('release', ''), receipt=em.get('receipt', ''), coo=em.get('coo', ''),
        hs=em.get('hs', ''), shipment_type=em.get('shipment_type', '')).items() if v}
    return sh


# ---------------------------------------------------------------- output
def generate(ships, cfg, log=print):
    cfg = dict(cfg)
    hd = dict(cfg.get('header') or {})
    if (hd.get('outer_pack_uom') or 'PKG').upper() == 'PKG':
        hd['outer_pack_uom'] = 'PAT'                  # MSD packing lists count pallets
    cfg['header'] = hd
    cfg['direction'] = 'Export'
    cfg['fill_term'] = False                           # as the MSD desktop tool: Term Type not filled
    cfg.setdefault('container_weight', 'TNE')
    cfg.setdefault('fill_consignee', False)
    em = {}
    for s in ships:
        for k, v in (s.get('email_doc') or {}).items():
            em.setdefault(k, v)
    em.update({k: v for k, v in (cfg.get('cargo') or {}).items() if v})
    if em.get('etd') and not hd.get('date'):           # departure date = ETD of the booking form
        hd['date'] = em['etd']
    cfg['email'] = em
    res = M.generate(ships, cfg, log=log)
    extra_h = {k: v for k, v in (cfg.get('extra_header') or {}).items() if str(v).strip()}
    extra_i = {k: v for k, v in (cfg.get('extra_invoice') or {}).items() if str(v).strip()}
    if extra_h or extra_i:
        for f in res[0]:
            _extra_fields(f['path'], extra_h, extra_i, log)
    return res


def _extra_fields(path, eh, ei, log):
    """"More permit fields" (any Header / Invoice column of the template) written into row 2."""
    tpl = M.Template(path)
    changed = {}
    for name, vals in (('Header', eh), ('Invoice', ei)):
        if not vals or name not in tpl.sheets:
            continue
        sh = tpl.sheet(name)
        H = tpl.headers(sh)
        for h, v in vals.items():
            v = str(v).strip()
            if name == 'Invoice' and re.fullmatch(r'-?\d[\d,]*(\.\d+)?', v) and any(w in h.lower() for w in ('amount', '%')):
                v = fnum(v)
            M._put(sh, H, h, 2, v, log=log)
        changed[tpl.sheets[name]] = sh.tostring()
    if changed:
        tpl.save(path, changed)


def write_raw(path, docs, ships=None):
    """RAW_DATA.xlsx - everything read from the documents."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Documents'
    ws.append(['File', 'Kind', 'Invoice / B/L no', 'Delivery', 'Order / PO', 'Currency', 'Total', 'Gross KG', 'Net KG',
               'Pallets', 'HAWB', 'Notes'])
    for d in docs:
        ws.append([d.get('file'), d.get('kind') or '?', d.get('invoice_no') or d.get('bl_no', ''), d.get('delivery', ''),
                   d.get('sap_order') or d.get('po', ''), d.get('currency', ''), d.get('total'),
                   d.get('total_gw') if d.get('kind') == 'PL' else d.get('gw'), d.get('total_nw'), d.get('pallets'),
                   d.get('hawb', ''), '; '.join(d.get('notes') or []) or ', '.join(
                       '%s: %s' % kv for kv in (d.get('email') or {}).items())])
    wl = wb.create_sheet('Invoice lines')
    wl.append(['File', 'Invoice', 'Item', 'Description', 'Qty', 'UOM', 'Unit price', 'Amount', 'HS', 'Batches'])
    for d in docs:
        for l in d.get('lines') or []:
            wl.append([d.get('file'), d.get('invoice_no', ''), l.get('item'), l.get('desc'), l.get('qty'), l.get('uom'),
                       l.get('price'), l.get('amount'), l.get('hs', ''), ', '.join(b for b in l.get('batches') or [] if b)])
    wp = wb.create_sheet('Packing list')
    wp.append(['File', 'Delivery', 'Material', 'Description', 'Qty', 'UOM', 'Net KG', 'Gross KG'])
    for d in docs:
        for k, m in (d.get('materials') or {}).items():
            wp.append([d.get('file'), d.get('delivery', ''), k, m.get('desc'), m.get('qty'), m.get('uom'), m.get('nw'), m.get('gw')])
    wi = wb.create_sheet('Items written')
    wi.append(['HAWB', 'Invoice', 'Item', 'Description', 'Qty', 'UOM', 'Amount', 'Net KG', 'Net from', 'HS', 'HS from',
               'HS UOM', 'HS Qty', 'COO', 'COO from', 'CASC'])
    for s in ships or []:
        for it in s['items']:
            wi.append([s.get('hawb'), it.get('invoice'), it.get('item'), it.get('description'), it.get('qty'), it.get('uom'),
                       it.get('amount'), it.get('nw'), it.get('nw_src'), it.get('hs'), it.get('hs_src'), it.get('hs_uom'),
                       it.get('hs_qty'), it.get('coo'), it.get('coo_src'), it.get('casc')])
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True)
    wb.save(path)
