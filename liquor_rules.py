# -*- coding: utf-8 -*-
"""Liquor (wine, champagne, spirits) classification for Singapore permits - Singapore HS code (AHTN) + CA/SC product
code chosen from the OFFICIAL descriptions in the UPDATE HS CODE list (Singapore Customs HSCA product codes), never
copied from the supplier's (EU) tariff code or from old permits.

  still grape wine : container <= 2 L / 2-10 L / > 10 L  x  alcohol <= 15% / 15-23% / > 23%  -> 22042111 ... 22042912
                     product code by colour: RED (BH) / WHITE (BJ) / BLUSH-ROSE (BA) / FORTIFIED (BF) / GRAPE OTHERS (BGT)
  sparkling        : 22041000 ZBP0BC000CS
  whisky           : container <= 5 L 22083090 / > 5 L 22083010; BLENDED / STRAIGHT x <= 46% / > 46%
  gin, rum, vodka, brandy (cognac / armagnac), liqueur, tequila, mezcal, sake: their own codes by alcohol band
Anything not sure (whisky type unknown, unknown spirit) -> no code + a note: type it on the Item page.

Duty on liquor = total litres x rate x alcohol % (Singapore Customs) -> HS Qty = litres (LTR), Alcohol Percentage,
Dutiable Qty = bottle size (LTR), Outer Pack = bottles (BOT) - the user's permit layout (dec-upload template sample)."""
import re

FORTIFIED = re.compile(r'\b(PORTO?|PORT WINE|TAWNY|MADEIRA|SHERRY|XERES|X[EÉ]R[EÈ]S|JEREZ|MANZANILLA|MARSALA|BANYULS|MAURY|'
                       r'RIVESALTES|VIN DOUX NATUREL|VDN|PINEAU|MUSCAT DE BEAUMES|RASTEAU VDN|MISTELLE)\b', re.I)
SPARKLING = re.compile(r'\b(CHAMPAGNE|CR[EÉ]MANT|PROSECCO|CAVA|SPUMANTE|SEKT|FRANCIACORTA|P[EÉ]TILLANT|SPARKLING|MOUSSEUX|'
                       r'METHODE TRADITIONNELLE|BRUT NATURE|EXTRA[- ]BRUT)\b', re.I)
SPIRITS = [  # (kind, regex) - order matters
    ('WHISKY', r'\bWHISK(?:E)?Y|\bSCOTCH\b|\bBOURBON\b|SINGLE MALT|\bCHICHIBU\b|\bYAMAZAKI\b|\bHAKUSHU\b|\bHIBIKI\b'),
    ('BRANDY', r'\bCOGNAC\b|\bARMAGNAC\b|\bBRANDY\b'),
    ('GRAPPA', r'\bGRAPPA\b|\bMARC DE\b'),
    ('GIN', r'\bGIN\b|\bGENEVER\b'),
    ('RUM', r'\bRHUM\b|\bRUM\b'),
    ('VODKA', r'\bVODKA\b'),
    ('TEQUILA', r'\bTEQUILA\b'),
    ('MEZCAL', r'\bMEZCAL\b'),
    ('LIQUEUR', r'\bLIQUEUR\b|\bLIQUOR CREAM\b|\bCR[EÈ]ME DE\b'),
    ('SAKE', r'\bSAKE\b'),
]
# country by appellation / region printed on the invoice (suggestion only - the permit keeps the document's origin)
COUNTRY_HINTS = [
    ('IT', r'\b(DOCG?|IGT)\b|LANGHE|BAROLO|BARBARESCO|TOSCANA|BOLGHERI|CHIANTI|MONTALCINO|ETNA|PIEMONTE|VALPOLICELLA|AMARONE|SICILIA|ALTO ADIGE|FRIULI'),
    ('ES', r'\bDO\b|\bDOCA\b|VALDEORRAS|PENED[EÈ]S|RIBEIRA SACRA|R[IÍ]AS BAIXAS|RIOJA|RIBERA DEL DUERO|PRIORAT|JEREZ|BIERZO|TORO'),
    ('PT', r'\bDOURO\b|\bPORTO\b|MADEIRA|VINHO VERDE|\bD[AÃ]O\b|ALENTEJO'),
    ('DE', r'\bQMP\b|\bQBA\b|MOSEL|RHEINGAU|PFALZ|NAHE|KABINETT|SP[AÄ]TLESE|AUSLESE|WEHLENER'),
    ('AT', r'WACHAU|KAMPTAL|KREMSTAL|BURGENLAND'),
    ('AU', r'EDEN VALLEY|BAROSSA|MCLAREN VALE|MARGARET RIVER|HUNTER VALLEY|COONAWARRA|YARRA|CLARE VALLEY|HENSCHKE|PENFOLDS'),
    ('NZ', r'MARLBOROUGH|CENTRAL OTAGO|HAWKE.?S BAY|MARTINBOROUGH'),
    ('US', r'NAPA|SONOMA|WILLAMETTE|PASO ROBLES|SANTA BARBARA|BOURBON|KENTUCKY'),
    ('JP', r'CHICHIBU|YAMAZAKI|HAKUSHU|HIBIKI|YOICHI|MIYAGIKYO|NIKKA|SUNTORY'),
    ('GB', r'\bISLAY\b|SPEYSIDE|SCOTCH|HIGHLAND|CAMPBELTOWN'),
    ('FR', r'\bAOC\b|\bAOP\b|\bIGP\b|VIN DE FRANCE|CHAMPAGNE|COGNAC|ARMAGNAC|BORDEAUX|BOURGOGNE|BURGUNDY'),
]


def kind_of(name, appellation='', eu_hs='', extra=''):
    """WINE / SPARKLING / FORTIFIED / WHISKY / GIN / ... from the invoice text and the supplier's tariff code."""
    t = ' '.join((name or '', appellation or '', extra or '')).upper()
    h = re.sub(r'\D', '', eu_hs or '')
    for k, rx in SPIRITS:
        if re.search(rx, t, re.I):
            return k
    if h.startswith('220830'):
        return 'WHISKY'
    if h.startswith('2206') and 'SAKE' in t:
        return 'SAKE'
    if h.startswith('220410') or SPARKLING.search(t):
        return 'SPARKLING'
    if FORTIFIED.search(t):
        return 'FORTIFIED'
    if h.startswith('2204') or re.search(r'\b(AOC|AOP|IGP|DOCG?|IGT|DO|QMP|VIN|WINE|CH[AÂ]TEAU|DOMAINE)\b', t):
        return 'WINE'
    if h.startswith('2208'):
        return 'SPIRIT'
    return ''


def colour_of(appellation='', name='', extra=''):
    t = ' '.join((appellation or '', name or '', extra or '')).upper()
    if re.search(r'\bROS[EÉ]\b|\bBLUSH\b|\(RS\)', t):
        return 'ROSE'
    if re.search(r'-\s*RED\b|\bROUGE\b|\bRED\b|\(R\)', t):
        return 'RED'
    if re.search(r'-\s*WHITE\b|\bBLANC\b|\bWHITE\b|\(B\)', t):
        return 'WHITE'
    return ''


def country_hint(text):
    t = (text or '').upper()
    for c, rx in COUNTRY_HINTS:
        if re.search(rx, t):
            return c
    return ''


def _band_ok(desc, abv):
    """Does the alcohol strength fit the band written in an official product description?"""
    d = re.sub(r'\s+', ' ', (desc or '').upper()).replace('EXCEEDING46', 'EXCEEDING 46')
    if abv is None:
        return True
    m = re.search(r'EXCEEDING ([\d.]+)% ?(?:VOL)? ?BUT NOT EXCEEDING ([\d.]+)%', d)
    if m:
        return float(m.group(1)) < abv <= float(m.group(2))
    m = re.search(r'NOT EXCEEDING ([\d.]+)%', d)
    if m:
        return abv <= float(m.group(1))
    m = re.search(r'EXCEEDING ([\d.]+)%', d)
    if m:
        return abv > float(m.group(1))
    return True


def hs_for(kind, abv, bottle_l):
    """Singapore HS (AHTN) for one liquor line - None when the rule cannot decide."""
    b = bottle_l if bottle_l else 0.75
    if kind == 'SPARKLING':
        return '22041000'
    if kind in ('WINE', 'FORTIFIED'):
        if abv is None:
            return None
        band = 0 if abv <= 15 else (1 if abv <= 23 else 2)
        if b <= 2:
            return ('22042111', '22042113', '22042114')[band]
        if b <= 10:
            return ('22042211', '22042212', '22042213')[band]
        return ('22042911', '22042912', '22042912')[band]
    if kind == 'WHISKY':
        return '22083090' if b <= 5 else '22083010'
    if kind == 'BRANDY' or kind == 'GRAPPA':
        return '22082050'
    if kind == 'GIN':
        return '22085000'
    if kind == 'RUM':
        return '22084000'
    if kind == 'VODKA':
        return '22086000'
    if kind == 'LIQUEUR':
        return None if abv is None else ('22087010' if abv <= 57 else '22087090')
    if kind in ('TEQUILA', 'MEZCAL'):
        return '22089099'
    if kind == 'SAKE':
        return '22060020'
    return None


def pick_code(kind, colour, abv, codes, whisky_type=''):
    """Product code of the HS chosen from the official descriptions. codes = [[code, description, uom, ...]]."""
    want = {'RED': 'WINE, RED', 'WHITE': 'WINE, WHITE', 'ROSE': 'BLUSH', 'FORTIFIED': 'FORTIFIED'}
    hits = []
    for c in codes or []:
        code, desc = c[0], str(c[1] or '').upper()
        if code in ('MISC', 'IDAMISC') or not _band_ok(desc, abv):
            continue
        if kind == 'SPARKLING' and 'SPARKLING' in desc:
            hits.append(c)
        elif kind == 'FORTIFIED' and 'FORTIFIED' in desc:
            hits.append(c)
        elif kind == 'WINE' and colour and want.get(colour, '#') in desc:
            hits.append(c)
        elif kind == 'WHISKY' and whisky_type and ('WHISKY, %s' % whisky_type) in desc:
            hits.append(c)
        elif kind == 'BRANDY' and desc.startswith('BRANDY, OF'):
            hits.append(c)
        elif kind == 'GRAPPA' and 'MARC' in desc:
            hits.append(c)
        elif kind == 'GIN' and desc.startswith('GIN'):
            hits.append(c)
        elif kind in ('RUM', 'VODKA') and desc.startswith(kind):
            hits.append(c)
        elif kind == 'LIQUEUR' and desc.startswith('LIQUEURS'):
            hits.append(c)
        elif kind in ('TEQUILA', 'MEZCAL', 'SAKE') and kind in desc:
            hits.append(c)
    return hits[0] if len(hits) == 1 else None


def whisky_type(text):
    t = (text or '').upper()
    if re.search(r'\bBLEND(ED)?\b|BLENDED MALT|PURE MALT|VATTED', t):
        return 'BLENDED'
    if re.search(r'SINGLE\s+(MALT|CASK|GRAIN|BARREL)|\bSTRAIGHT\b', t):
        return 'STRAIGHT'
    return ''


def classify(liq, product_codes):
    """liq = dict(kind, colour, abv, bottle_l, name, appellation) -> dict(hs, code, code_desc, why, notes)."""
    kind, abv, b = liq.get('kind') or '', liq.get('abv'), liq.get('bottle_l')
    notes = []
    hs = hs_for(kind, abv, b)
    if not hs:
        return dict(hs='', code='', code_desc='', why='', notes=['%s: no Singapore HS by the liquor rules (type %s, %s%% vol) - '
                                                                 'type the HS code' % (liq.get('name', ''), kind or '?', abv)])
    wt = whisky_type(' '.join((liq.get('name', ''), liq.get('appellation', '')))) if kind == 'WHISKY' else ''
    c = pick_code(kind, liq.get('colour'), abv, product_codes(hs), wt)
    why = '%s%s, %s%% vol, %s L bottle' % (kind.lower(), (' ' + liq['colour'].lower()) if liq.get('colour') else '',
                                           ('%g' % abv) if abv is not None else '?', ('%g' % b) if b else '?')
    if kind == 'WHISKY' and not wt:
        notes.append('%s: whisky - blended or straight (single malt / single cask) not on the invoice - CASC code left blank'
                     % liq.get('name', ''))
    elif kind == 'WINE' and not liq.get('colour'):
        notes.append('%s: wine colour (red / white / rose) not on the invoice - CASC code left blank' % liq.get('name', ''))
    elif not c:
        notes.append('%s: no single product code of HS %s fits (%s) - pick it on the Item page' % (liq.get('name', ''), hs, why))
    return dict(hs=hs, code=c[0] if c else '', code_desc=c[1] if c else '', why=why, notes=notes, whisky=wt)


# ---------------------------------------------------------------- past permits (Clasquin MASTER LIST): cross-check only
_MASTER = {}


def norm_name(s):
    s = re.sub(r'\(.*?\)', ' ', str(s or '').upper())
    s = re.sub(r'[^A-Z0-9 ]+', ' ', s.replace('É', 'E').replace('È', 'E').replace('Â', 'A'))
    return ' '.join(w for w in s.split() if w not in ('CHATEAU', 'CH', 'DOMAINE', 'BT', 'DEMI', 'THE', 'DE', 'DU', 'LA', 'LE'))


def load_master(path):
    """MASTER LIST (HS Code | Description Goods | CA/SC Product Code 1) - past permits of the customer."""
    if path in _MASTER:
        return _MASTER[path]
    out = {}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for r in list(wb.active.iter_rows(values_only=True))[1:]:
            if r and r[0] and r[1]:
                out.setdefault(norm_name(r[1]), set()).add((re.sub(r'\D', '', str(r[0]))[:8], str(r[2] or '').strip()))
        wb.close()
    except Exception:
        pass
    _MASTER[path] = out
    return out
