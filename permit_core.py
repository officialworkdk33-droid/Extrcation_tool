# -*- coding: utf-8 -*-
"""
JAPAN FOOD PERMIT TOOL - core logic (no GUI)
GLT CO.,LTD invoice Excel + ANA/JAL air waybill PDF + SFA "HS CODE & PRODUCT CODE" master list
-> dec-upload permit Excel in the built-in OUTPUT FILE format.

Rules:
  * HS code + CA/SC product code come ONLY from the master list (or codes you saved in the tool).
    Not found -> left empty and flagged.
  * HS Qty = net weight = invoice (Quantity) column in KG for every line, converted to the HS UOM
    from the built-in UPDATE HS CODE list (TNE = kg / 1000), 4 decimals, minimum 0.0001.
  * CASC Product Code 1 Qty / UOM = HS Qty / UOM.
  * Casc sheet Code 1 = Japan prefecture code of the production center.
  * Description = "NAME (ENGLISH)**PREFECTURE**".
"""
import base64, datetime as dt, io, json, os, re, sys, unicodedata, zlib
from collections import Counter, defaultdict

import openpyxl

import embedded_data as ED

APP_NAME = "JAPAN FOOD PERMIT TOOL"

# --------------------------------------------------------------------------------------------
# Built-in data
# --------------------------------------------------------------------------------------------
HS_TABLE = json.loads(zlib.decompress(base64.b64decode(ED.HS_DATA)))   # hs -> [uom, desc, imp agency, imp control, all CA codes]
PREFECTURES = dict(ED.PREFECTURES)                                      # 'HOKKAIDO' -> 'JPHKD'
PREF_ALIASES = {"IBARAGI": "IBARAKI", "HYOUGO": "HYOGO", "OOITA": "OITA", "KOUCHI": "KOCHI",
                "KYOUTO": "KYOTO", "TOUKYOU": "TOKYO", "OOSAKA": "OSAKA", "GUNMA": "GUNMA",
                "HOKKAIDOU": "HOKKAIDO", "SHIGA": "SHIGA"}

AIRPORTS = {  # loading port from AWB routing "NRT-SIN"
    "NRT": ("JPNRT", "NARITA APT/TOKYO"),
}


HS_OVERRIDE_FILE = None  # set after app_dir() is defined


def load_hs_excel(path):
    """Replace the built-in HS list with a newer UPDATE HS CODE Excel (sheet 'HS CODE')."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = next((w for w in wb.worksheets if w.title.strip().upper() == "HS CODE"), wb.worksheets[0])
    try:
        ws.reset_dimensions()
    except Exception:
        pass
    new = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not r or not r[0]:
            continue
        code = re.sub(r"\D", "", str(r[0]))
        if len(code) != 8:
            continue
        allpc = r[37] if len(r) > 37 and isinstance(r[37], str) else ""
        new[code] = [r[2] or "", r[1] or "", (r[8] if len(r) > 8 else "") or "",
                     (r[11] if len(r) > 11 else "") or "", allpc]
    wb.close()
    if len(new) < 1000:
        raise ValueError("This does not look like the UPDATE HS CODE workbook (only %d HS codes found)." % len(new))
    HS_TABLE.clear()
    HS_TABLE.update(new)
    with open(HS_OVERRIDE_FILE, "w", encoding="utf-8") as f:
        json.dump(new, f)
    return len(new)


def template_bytes():
    return zlib.decompress(base64.b64decode(ED.TEMPLATE_XLSX))


def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


SETTINGS_FILE = os.path.join(app_dir(), "permit_tool_settings.json")
LEARNED_FILE = os.path.join(app_dir(), "learned_codes.json")
HS_OVERRIDE_FILE = os.path.join(app_dir(), "hs_code_update.json")
try:   # a newer HS list loaded from the app overrides the built-in one
    with open(HS_OVERRIDE_FILE, encoding="utf-8") as _f:
        _hs = json.load(_f)
    if len(_hs) > 1000:
        HS_TABLE.clear(); HS_TABLE.update(_hs)
except Exception:
    pass

DEFAULT_HEADER = {
    "Importer": "GINZA SUSHI ICHI PTE. LTD.",
    "Importer Uen": "201435734W",
    "Cargo Pack Type Code": "5",
    "Declaration Type Code": "GST",
    "Inward Carrier Agent Name": "SINGAPORE AIRPORT TERMINAL SERVICES",
    "Inward Carrier Agent Uen": "197201770G",
    "Inward Transport Mode": "4",
    "Loading Port Code": "JPNRT",
    "Loading Port Name": "NARITA APT/TOKYO",
    "Release Location Code": "CZ",
    "Release Location Name": "CHANGI FTZ",
    "Receipt Location Code": "O",
    "Receipt Location Name": "OTHERS",
    "Total Gross Weight Uom": "KGM",
    "Total Outer Pack Uom": "PKG",
    "Reference Documents": "Y",
    "Licence 1": "IH15E0171",
    "Internal Remarks": "DK",
    "Sea Store": "N",
    "Cross Reference": "JP",
}
DEFAULT_INVOICE = {
    "Invoice Number": "-",
    "Term Type": "FOB",
    "Ad Valorem Indicator": "N",
    "Preferential Duty Indicator": "N",
    "Supplier/Manufacturer Name": "-",
    "Total Invoice Currency": "JPY",
    "Other Taxable Charge %": 0,
    "Other Taxable Charge Amount": 0,
    "Freight Charge %": 0,
    "Freight Charge Currency": "JPY",
    "Insurance Charge %": 1,
    "Insurance Charge Currency": "SGD",
    "Insurance Charge Amount": 0,
}
DEFAULT_SETTINGS = {"header": DEFAULT_HEADER, "invoice": DEFAULT_INVOICE, "brand": "GLT CO.",
                    "master_files": [], "last_dir": "", "airports": {}}


def load_settings():
    s = json.loads(json.dumps(DEFAULT_SETTINGS))
    try:
        with open(SETTINGS_FILE, encoding="utf-8") as f:
            saved = json.load(f)
        for k, v in saved.items():
            if isinstance(v, dict) and isinstance(s.get(k), dict):
                s[k].update(v)
            else:
                s[k] = v
    except Exception:
        pass
    return s


def save_settings(s):
    try:
        with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump(s, f, indent=1, ensure_ascii=False)
    except Exception:
        pass


def load_learned():
    try:
        with open(LEARNED_FILE, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_learned(d):
    try:
        with open(LEARNED_FILE, "w", encoding="utf-8") as f:
            json.dump(d, f, indent=1, ensure_ascii=False)
    except Exception:
        pass


# --------------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------------
def nfkc(s):
    return unicodedata.normalize("NFKC", str(s)) if s is not None else ""


def key(s):
    """Normalised match key: letters only, upper case."""
    return re.sub(r"[^A-Z]", "", nfkc(s).upper())


def clean_name(n):
    n = nfkc(n).upper().strip()
    n = re.sub(r"\s*\(\s*", " (", n)
    n = re.sub(r"\s*\)", ")", n)
    n = re.sub(r"\s+", " ", n)
    return n.strip()


def pref_code(pref):
    p = re.sub(r"[^A-Z]", "", nfkc(pref).upper())
    p = PREF_ALIASES.get(p, p)
    for suf in ("PREFECTURE", "KEN", "FU", "TO"):
        if p not in PREFECTURES and p.endswith(suf) and p[: -len(suf)] in PREFECTURES:
            p = p[: -len(suf)]
    return PREFECTURES.get(p, "")


def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(nfkc(v).replace(",", "").strip())
    except ValueError:
        return None


def fmt_qty(v):
    return float("%.4f" % v)


def parse_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = nfkc(v).strip().replace(",", ".")
    for f in ("%b.%d.%Y", "%b.%d. %Y", "%b %d %Y", "%d.%b.%Y", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%Y-%m-%d",
              "%B.%d.%Y", "%B %d.%Y", "%b. %d.%Y"):
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            pass
    m = re.search(r"([A-Za-z]{3})[A-Za-z]*\.?\s*(\d{1,2})\.?\s*(\d{4})", s)
    if m:
        try:
            return dt.datetime.strptime("%s %s %s" % m.groups(), "%b %d %Y").date()
        except ValueError:
            pass
    return None


# --------------------------------------------------------------------------------------------
# Master list (SFA "HS CODE & PRODUCT CODE" report)
# --------------------------------------------------------------------------------------------
class Master:
    def __init__(self):
        self.full = defaultdict(Counter)   # key(desc with prefecture) -> Counter((hs, pc1, pc2))
        self.name = defaultdict(Counter)   # key(desc without **PREF**) -> Counter
        self.records = 0
        self.files = []

    def load(self, path):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        n = 0
        for ws in wb.worksheets:
            try:
                ws.reset_dimensions()   # some system exports carry a wrong sheet size
            except Exception:
                pass
            cols = None
            for r in ws.iter_rows(values_only=True):
                if cols is None:
                    up = [nfkc(c).strip().upper() if c is not None else "" for c in r]
                    if "HS CODE" in up and any("DESCRIPTION" in c for c in up):
                        cols = {
                            "desc": next(i for i, c in enumerate(up) if "DESCRIPTION" in c),
                            "hs": up.index("HS CODE"),
                            "pc1": next((i for i, c in enumerate(up) if "PRODUCT CODE 1" in c), None),
                            "pc2": next((i for i, c in enumerate(up) if "PRODUCT CODE 2" in c), None),
                        }
                    continue
                desc = r[cols["desc"]] if cols["desc"] < len(r) else None
                hs = r[cols["hs"]] if cols["hs"] < len(r) else None
                if not desc or not hs:
                    continue
                hs = re.sub(r"\D", "", nfkc(hs)).zfill(8)[:8]
                pc1 = nfkc(r[cols["pc1"]]).strip() if cols["pc1"] is not None and r[cols["pc1"]] else ""
                pc2 = nfkc(r[cols["pc2"]]).strip() if cols["pc2"] is not None and r[cols["pc2"]] else ""
                if pc1 == "0":
                    pc1 = ""
                val = (hs, pc1, pc2)
                self.full[key(desc)][val] += 1
                self.name[key(re.sub(r"\*+.*$", "", nfkc(desc)))][val] += 1
                n += 1
        wb.close()
        if n == 0:
            raise ValueError("No 'Description Goods / HS Code' rows found in master list:\n" + path)
        self.records += n
        self.files.append(path)
        return n

    def lookup(self, name, pref):
        desc = "%s**%s**" % (name, pref)
        c = self.full.get(key(desc))
        if c:
            return c.most_common(1)[0][0], "FOUND"
        c = self.name.get(key(name))
        if c:
            return c.most_common(1)[0][0], "FOUND (OTHER PREFECTURE)"
        return None, "NOT IN MASTER"


# --------------------------------------------------------------------------------------------
# Invoice (GLT CO.,LTD Excel: cover sheet + "Attached sheet"s)
# --------------------------------------------------------------------------------------------
def read_invoice(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    info = {"invoice_no": "", "invoice_date": None, "term": "", "currency": "", "shipper": "", "consignee": "",
            "cartons": None, "lines": [], "sheet_totals": {}}
    # ---- cover sheet details
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                t = nfkc(v).strip().upper()
                right = [ws.cell(c.row, c.column + k).value for k in range(1, 5)]
                right = [x for x in right if x not in (None, "") and nfkc(x).strip()]
                if t == "INVOICE NO" and right and not info["invoice_no"]:
                    info["invoice_no"] = nfkc(right[0]).strip()
                elif t == "DATE" and right and not info["invoice_date"]:
                    info["invoice_date"] = parse_date(right[0])
                elif t == "SHIPPER" and not info["shipper"]:
                    info["shipper"] = nfkc(ws.cell(c.row + 1, c.column).value or "").strip()
                elif t == "CONSIGNEE" and not info.get("consignee"):
                    info["consignee"] = nfkc(ws.cell(c.row + 1, c.column).value or "").strip()
                elif re.match(r"^(FOB|CIF|CFR|C&F|EXW|FCA|DAP|DDP|CPT|CIP)\b", t) and not info["term"]:
                    info["term"] = t.split()[0].replace("C&F", "CFR")
                elif re.match(r"^IN [A-Z]{3}$", t) and not info["currency"]:
                    info["currency"] = t[3:]
                elif t.startswith("TOTAL:") and right and to_num(right[0]) and "CARTON" in nfkc(right[1] if len(right) > 1 else "").upper():
                    info["cartons"] = to_num(right[0])
                if not info["invoice_no"]:
                    m = re.search(r"INVOICE\s*NO\.?\s*[:：]?\s*([A-Z]{1,4}\d{6,})", t)
                    if m:
                        info["invoice_no"] = m.group(1)
    # ---- attached item sheets
    for ws in wb.worksheets:
        hdr_row, cols = None, {}
        for row in ws.iter_rows(max_row=15):
            texts = {c.column: nfkc(c.value).upper() for c in row if isinstance(c.value, str)}
            if any("DESCRIPTION" in t for t in texts.values()) and any("AMOUNT" in t for t in texts.values()):
                hdr_row = row[0].row
                for col, t in texts.items():
                    if "DESCRIPTION" in t: cols["desc"] = col
                    elif "PRODUCTION" in t or "CENTER" in t or "ORIGIN" in t: cols["pref"] = col
                    elif "STANDARD" in t: cols["std"] = col
                    elif "QUANTITY" in t or t.strip("() ") == "QTY": cols["qty"] = col
                    elif "UNIT PRICE" in t: cols["price"] = col
                    elif "UNIT" in t: cols["unit"] = col
                    elif "AMOUNT" in t: cols["amt"] = col
                break
        if not hdr_row or "desc" not in cols or "qty" not in cols:
            continue
        g = lambda r, k, off=0: ws.cell(r, cols[k] + off).value if k in cols else None
        tag = None
        for r in range(hdr_row + 1, ws.max_row + 1):
            d = g(r, "desc")
            if d is None or not nfkc(d).strip():
                continue
            ds = nfkc(d).strip()
            if re.match(r"^S\.?\s*TOTAL", ds.upper()) or ds.upper().startswith("TOTAL"):
                info["sheet_totals"][ws.title] = to_num(g(r, "amt"))
                tag = ws.title
                for rr in range(r + 1, min(r + 8, ws.max_row + 1)):   # customer name written under the total
                    nxt = ws.cell(rr, cols["desc"]).value
                    if isinstance(nxt, str) and nfkc(nxt).strip():
                        tag = nfkc(nxt).strip()
                        break
                break
            qty = to_num(g(r, "qty"))
            price = to_num(g(r, "price"))
            amt = to_num(g(r, "amt"))
            if amt is None and qty is not None and price is not None:
                amt = round(qty * price, 2)
            if qty is None and amt is None:
                continue
            std_v = to_num(g(r, "std"))
            std_u = nfkc(g(r, "std", 1) or "").strip().upper()
            info["lines"].append({
                "sheet": ws.title, "row": r, "name": clean_name(ds),
                "pref": nfkc(g(r, "pref") or "").strip().upper(),
                "std": std_v, "std_unit": std_u, "qty": qty or 0.0,
                "unit": nfkc(g(r, "unit") or "").strip().upper(), "price": price, "amount": amt or 0.0,
            })
        for ln in info["lines"]:
            if ln["sheet"] == ws.title:
                ln["customer"] = tag or ws.title
    wb.close()
    if not info["lines"]:
        raise ValueError("No item lines found in the invoice Excel (looking for '(Description of goods)' ... '(Amount)' sheets).")
    return info


def net_weight_kg(ln):
    """The invoice (Quantity) column IS the net weight in KG for every line (KG, PK or PC)."""
    return ln["qty"] if ln.get("qty") is not None else None


# --------------------------------------------------------------------------------------------
# Air waybill PDF
# --------------------------------------------------------------------------------------------
def read_awb(path, airports=None):
    import pdfplumber
    airports = dict(AIRPORTS, **(airports or {}))
    with pdfplumber.open(path) as pdf:
        text = "\n".join((p.extract_text() or "") for p in pdf.pages)
    t = nfkc(text)
    res = {"text": text, "mawb": "", "flight": "", "arrival": None, "gw": None, "pkgs": None,
           "freight": None, "weight_charge": None, "other_charges": {}, "invoice_no": "", "origin": "",
           "loading_code": "", "loading_name": ""}
    if not t.strip():
        return res
    m = re.search(r"\b(\d{3})[\s-]?(\d{4})\s?(\d{4})\b", t)
    if m:
        res["mawb"] = "".join(m.groups())
    exec_date = None
    m = re.search(r"(\d{1,2})[.\s-]([A-Z]{3})[.\s-](\d{4})", t.upper())
    if m:
        try:
            exec_date = dt.datetime.strptime("%s %s %s" % m.groups(), "%d %b %Y").date()
        except ValueError:
            pass
    m = re.search(r"\b([A-Z]{2}|[A-Z]\d|\d[A-Z])\s?(\d{3,4})\s*/\s*(\d{1,2})\b", t)
    if m:
        res["flight"] = m.group(1) + m.group(2).zfill(4)
        day = int(m.group(3))
        base = exec_date or dt.date.today()
        try:
            d = base.replace(day=day)
            if d < base - dt.timedelta(days=5):   # flight in next month
                nm = (base.replace(day=1) + dt.timedelta(days=32)).replace(day=day)
                d = nm
            res["arrival"] = d
        except ValueError:
            res["arrival"] = base
    elif exec_date:
        res["arrival"] = exec_date
    m = re.search(r"\b([A-Z]{3})\s?-\s?SIN\b", t)
    if m:
        res["origin"] = m.group(1)
        code, name = airports.get(m.group(1), ("JP" + m.group(1), ""))
        res["loading_code"], res["loading_name"] = code, name
    # rate line: "9 55.0K N 55.0 1,320 72,600"
    m = re.search(r"^\s*(\d+)\s+([\d.,]+)\s*K\b.*?([\d,]+)\s*$", t, re.M)
    if m:
        res["pkgs"] = int(m.group(1))
        res["gw"] = float(m.group(2).replace(",", ""))
        res["weight_charge"] = float(m.group(3).replace(",", ""))
    for code, amt in re.findall(r"\b([A-Z]{3})\s*:\s*([\d,]+)", t):
        if code in ("MYC", "SCC", "FSC", "SSC", "XDC", "AWC", "RAC", "CGC", "MOC"):
            res["other_charges"][code] = float(amt.replace(",", ""))
    if res["weight_charge"] is not None:
        res["freight"] = res["weight_charge"] + sum(res["other_charges"].values())
    m = re.search(r"INVOICE\s*NO\.?\s*:?\s*\n?\s*([A-Z]{1,4}\d{6,})", t.upper())
    if m:
        res["invoice_no"] = m.group(1)
    return res


# --------------------------------------------------------------------------------------------
# Build items
# --------------------------------------------------------------------------------------------
def build_items(inv, master, learned, brand):
    items = []
    for ln in inv["lines"]:
        name, pref = ln["name"], ln["pref"]
        desc = "%s**%s**" % (name, pref) if pref else name
        status = ""
        lk = learned.get(key(desc)) or learned.get(key(name))
        if lk:
            hs, pc1, pc2 = lk.get("hs", ""), lk.get("pc1", ""), lk.get("pc2", "")
            status = "SAVED CODE"
        else:
            val, status = master.lookup(name, pref) if master else (None, "NO MASTER LIST")
            hs, pc1, pc2 = val if val else ("", "", "")
        kg = net_weight_kg(ln)
        it = {"desc": desc, "name": name, "pref": pref, "pref_code": pref_code(pref), "hs": hs,
              "pc1": pc1, "pc2": pc2, "kg": kg, "amount": ln["amount"], "currency": inv.get("currency") or "JPY",
              "inv_qty": ln["qty"], "inv_unit": ln["unit"], "std": ln["std"], "std_unit": ln["std_unit"],
              "customer": ln.get("customer", ln["sheet"]), "status": status, "brand": brand, "origin": "JP"}
        recalc(it)
        items.append(it)
    return items


def recalc(it):
    """Fill HS UOM / qty, group letter and warnings from current hs / pc1 / kg."""
    hs = re.sub(r"\D", "", str(it.get("hs") or ""))
    it["hs"] = hs.zfill(8) if hs else ""
    row = HS_TABLE.get(it["hs"]) if it["hs"] else None
    it["uom"] = row[0] if row else ""
    kg = it.get("kg")
    q = None
    if kg is not None and it["uom"]:
        if it["uom"] == "TNE":
            q = kg / 1000.0
        elif it["uom"] == "KGM":
            q = kg
        elif it["uom"] == "GRM":
            q = kg * 1000
        else:   # NMB / PCE etc. -> invoice quantity
            q = it.get("inv_qty")
    elif kg is not None:
        q = kg
    if q is not None:
        q = fmt_qty(q)
        if q == 0 and (kg or 0) > 0:
            q = 0.0001
    it["hs_qty"] = q
    it["group"] = (it.get("pc1") or "Z")[:1].upper()
    warn = []
    if not it["hs"]:
        warn.append("NO HS")
    elif not row:
        warn.append("HS NOT IN HS LIST")
    elif it.get("pc1") and row[4] and it["pc1"] not in [x.strip() for x in row[4].split(",")]:
        warn.append("PRODUCT CODE NOT UNDER HS")
    if row and row[3] == "Control" and not it.get("pc1"):
        warn.append("CONTROLLED - NO PRODUCT CODE")
    if kg is None:
        warn.append("NO NET WEIGHT")
    if not it.get("pref_code"):
        warn.append("NO PREFECTURE CODE")
    it["warnings"] = warn
    return it


# --------------------------------------------------------------------------------------------
# Write OUTPUT FILE format
# --------------------------------------------------------------------------------------------
LABEL_LISTS = {  # label column -> (code column, list sheet)
    "Cargo Pack Type ": ("Cargo Pack Type Code", "List-CargoPacktype"),
    "Declaration Type": ("Declaration Type Code", "List-DeclarationType"),
    "Outward Transport": ("Outward Transport Mode", "List-TransportMode"),
    "Inward Transport": ("Inward Transport Mode", "List-TransportMode"),
    "Vessel Type List": ("Vessel Type", "List-VesselType"),
    "CO Type List": ("CO Type", "List-CoType"),
    "Certificate Detail #1 - Type List": ("Certificate Detail #1 - Type", "List-CoCertType"),
    "Certificate Detail #2 - Type List": ("Certificate Detail #2 - Type", "List-CoCertType"),
    "BG Indicator List": ("BG Indicator", "List-BgIndicator"),
}
DATE_FMT = "dd-mm-yyyy"


def _clear(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)


def write_permit(path, header, invoice, items, remarks):
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes()))
    # ---------- Header
    ws = wb["Header"]
    cols = {nfkc(c.value): c.column for c in ws[1] if c.value}
    hv = dict(header)
    hv["Trader's Remarks"] = remarks
    for lab, (codecol, lst) in LABEL_LISTS.items():
        code = str(hv.get(codecol, "") or "")
        if code and lst in wb.sheetnames:
            for a, b in wb[lst].iter_rows(values_only=True):
                if str(b) == code:
                    hv[lab] = a
                    break
    for k, v in hv.items():
        if k in cols and v not in (None, ""):
            c = ws.cell(2, cols[k])
            if isinstance(v, dt.date):
                c.value = v
                c.number_format = DATE_FMT
            else:
                c.value = v
    # ---------- Invoice
    ws = wb["Invoice"]
    cols = {nfkc(c.value): c.column for c in ws[1] if c.value}
    for k, v in invoice.items():
        if k in cols and v not in (None, ""):
            c = ws.cell(2, cols[k])
            c.value = v
            if isinstance(v, dt.date):
                c.number_format = DATE_FMT
    # ---------- Item + Casc
    wi, wc = wb["Item"], wb["Casc"]
    _clear(wi)
    _clear(wc)
    icol = {nfkc(c.value): c.column for c in wi[1] if c.value}
    zero_cols = ["Unit Price", "Total Dutiable Qty", "Dutiable Qty", "Outer Pack Qty", "In Pack Qty",
                 "Inner Pack Qty", "Inmost Pack Qty", "Last Selling Price", "Other Tax Rate", "Other Tax Amount"]
    for n, it in enumerate(items, 1):
        r = n + 1
        vals = {
            "Country of Origin": it.get("origin", "JP"), "HS Code": it["hs"], "Description": it["desc"],
            "HS Qty": it["hs_qty"], "Total Line Amount": it["amount"], "DG Indicator": "N", "Unbranded": "N",
            "Brand": it.get("brand", ""), "HS UOM": it["uom"], "Invoice Number": "-",
            "Item Currency": it.get("currency", "JPY"),
            "CASC Product Code 1": it.get("pc1", ""),
            # SFA team: CASC qty / UOM by the PRODUCT UOM (pc1_qty / pc1_uom); otherwise = HS qty / UOM
            "CASC Product Code 1 Qty": (it["pc1_qty"] if it.get("pc1_qty") is not None else it["hs_qty"]) if it.get("pc1") else None,
            "CASC Product Code 1 UOM": (it.get("pc1_uom") or it["uom"]) if it.get("pc1") else None,
            "CASC Product Code 2": it.get("pc2", ""),
            "CASC Product Code 2 Qty": it["hs_qty"] if it.get("pc2") else None,
            "CASC Product Code 2 UOM": it["uom"] if it.get("pc2") else None,
            "Shipping Marks 1": it["group"],
        }
        for zc in zero_cols:
            vals[zc] = 0
        for k, v in vals.items():
            if k in icol and v not in (None, ""):
                c = wi.cell(r, icol[k])
                c.value = v
                if k == "HS Code":
                    c.number_format = "@"
        if it.get("sfa") and not it.get("pc1"):
            continue                     # SFA team: Casc row only for items with a CASC product code
        cr = wc.max_row + 1 if (wc.max_row > 1 or wc.cell(1, 1).value) else 2
        for j, v in enumerate([n, 1, it.get("pref_code", "")], 1):
            wc.cell(cr, j).value = v
    wb.save(path)
    return path


def generate(out_dir, base_name, header, invoice, items, invoice_no, split=False):
    """Write one permit (or one per product group F/H/Z when split=True). Returns list of paths."""
    os.makedirs(out_dir, exist_ok=True)
    outs = []
    if not split:
        h = dict(header)
        p = os.path.join(out_dir, base_name + ".xlsx")
        outs.append(write_permit(p, h, invoice, items, "INVOICE NO : %s" % invoice_no))
    else:
        groups = []
        for it in items:
            if it["group"] not in groups:
                groups.append(it["group"])
        for g in sorted(groups, key=lambda x: "FHZ".find(x) if x in "FHZ" else 9):
            h = dict(header)
            h["Cross Reference"] = "%s-JP" % g
            p = os.path.join(out_dir, "%s_%s.xlsx" % (base_name, g))
            sel = [it for it in items if it["group"] == g]
            outs.append(write_permit(p, h, invoice, sel, "INVOICE NO : %s\n\nPARTS OF PERMIT" % invoice_no))
    return outs


def check_report(items, inv, awb):
    lines = []
    tot = sum(i["amount"] for i in items)
    st = sum(v for v in inv.get("sheet_totals", {}).values() if v)
    if st and abs(st - tot) > 0.5:
        lines.append("Invoice sub-totals %s <> item total %s" % (f"{st:,.0f}", f"{tot:,.0f}"))
    if awb and awb.get("invoice_no") and inv.get("invoice_no") and awb["invoice_no"] != inv["invoice_no"]:
        lines.append("AWB invoice no %s <> invoice %s" % (awb["invoice_no"], inv["invoice_no"]))
    for n, it in enumerate(items, 1):
        if it["warnings"]:
            lines.append("Line %d %s: %s" % (n, it["desc"], ", ".join(it["warnings"])))
        if it["status"].startswith("FOUND (OTHER"):
            lines.append("Line %d %s: code taken from same item with another prefecture" % (n, it["desc"]))
    return lines
