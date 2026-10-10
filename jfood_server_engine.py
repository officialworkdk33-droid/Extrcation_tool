# -*- coding: utf-8 -*-
"""JAPAN FOOD (GLT) permits on the Permit Server.

The reading and writing rules are the JAPAN FOOD PERMIT TOOL's own (permit_core.py - not changed). This file connects
them to the server, which expects the same functions as the other tools:
  - documents of one job: GLT CO.,LTD invoice Excel + ANA / JAL air waybill PDF
    (+ the SFA "HS CODE & PRODUCT CODE" master list the first time, + optionally a previous permit Excel for the header)
  - SFA master lists: Settings > Product masters > "Japan food" (upload), or added to a job - kept for every next job
  - codes typed on the Item tab with "Save to product master" = the tool's learned codes (next time filled by itself)
  - HS UOM / HS qty / control from the server's UPDATE HS CODE list (Settings > UPDATE HS CODE list)
  - output: the tool's OUTPUT FILE permit Excel (split by product group F / H / Z when "split" is ticked) + the
    server's FrtImport XML, CHECK_REPORT, RAW_DATA, sheet tabs, validation
"""
import os, re, json, glob, shutil, zipfile, datetime as dt, collections

import permit_core as C
import msd_engine as M
import sfa_rules as SFA

TOOL_KEY = 'jfood'
PERMIT_SCHEME = M.PERMIT_SCHEME
DEFAULT_UOM_RULES = M.DEFAULT_UOM_RULES
DEFAULT_DESC_TEMPLATE = '{DESC}'
DEFAULT_REMARKS = 'INVOICE NO : {INV}'
OCR = dict(enabled=False, cmd='', status='')
fnum, fmt_thou, fmt_plain, h8 = M.fnum, M.fmt_thou, M.fmt_plain, M.h8
set_uom_rules = M.set_uom_rules
template_lists = M.template_lists
Countries = M.Countries
LEARNED_NAME = 'JAPAN_FOOD_LEARNED_CODES.json'
SFA_DIR_NAME = 'JAPAN_FOOD_SFA'
STATE = dict(master_dir='', hs_path='')


def set_tesseract(cmd):
    pass


def hs_candidates(hs, hsl, n=40):
    import danfoss_engine as D
    return D.hs_candidates(hs, hsl, n)


# ---------------------------------------------------------------- HS list: the server's UPDATE HS CODE list
def load_hs_list(path):
    """Server HS list -> permit_core.HS_TABLE (hs -> [uom, desc, imp agency, imp control, CA product codes]) and the
    common {hs: (code, uom, desc, info)} for the server pages."""
    import danfoss_engine as D
    out = D.load_hs_list(path)
    STATE['hs_path'] = path                    # SFA team rules check HS | product code pairs against this list
    try:
        with open(path, encoding='utf-8') as f:
            d = json.load(f)
        pcs = d.get('pcodes') or {}
        tab = {k: [r[1] or '', r[2] or '', r[5] or '', r[3] or '', ','.join(p[0] for p in pcs.get(k, []))]
               for k, r in d['hs'].items()}
        if len(tab) > 1000:
            C.HS_TABLE.clear()
            C.HS_TABLE.update(tab)
    except Exception:
        pass
    return out


# ---------------------------------------------------------------- masters: SFA lists + learned codes
def _learned_path():
    return os.path.join(STATE['master_dir'] or '.', LEARNED_NAME)


def _sfa_dir():
    p = os.path.join(STATE['master_dir'] or '.', SFA_DIR_NAME)
    os.makedirs(p, exist_ok=True)
    return p


def _packed_masters():
    import sys
    return os.path.join(getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__))), 'masters')


def sfa_dir():
    """PERMIT_SERVER_DATA/masters/sfa (SFA team database + overrides) - made from the packed copy the first time."""
    return SFA.setup_folder(STATE['master_dir'] or '.', _packed_masters())


def sfa_rules():
    return SFA.get(sfa_dir(), STATE.get('hs_path') or os.path.join(os.path.dirname(STATE['master_dir'] or '.'), 'hs_data.json'))


def _read_learned():
    try:
        with open(_learned_path(), encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


_SFA_CACHE = {}


def sfa_master(extra=()):
    """All SFA master lists (Settings upload + kept from jobs + this job) as one permit_core.Master (cached)."""
    files = []
    main = os.path.join(STATE['master_dir'] or '.', 'JAPAN_FOOD_SFA_MASTER.xlsx')
    if os.path.exists(main):
        files.append(main)
    files += sorted(glob.glob(os.path.join(_sfa_dir(), '*.xlsx')))
    files += [p for p in extra if p not in files]
    sig = tuple((p, os.path.getmtime(p)) for p in files if os.path.exists(p))
    if _SFA_CACHE.get('sig') != sig:
        m = C.Master()
        bad = []
        for p in files:
            try:
                m.load(p)
            except Exception as e:
                bad.append('%s: %s' % (os.path.basename(p), e))
        _SFA_CACHE.update(sig=sig, master=m if m.records else None, bad=bad)
    return _SFA_CACHE['master'], _SFA_CACHE['bad']


def load_master(path):
    """Product master of this tool. path = masters/JAPAN_FOOD_SFA_MASTER.xlsx (the SFA list, Settings upload).
    Returns the learned codes as {item: {desc, hs, uom, coo, casc, pc2}} (what "Save to product master" edits)."""
    if path.endswith('.upload'):               # Settings > Product masters > Upload: check it is an SFA list
        import tempfile
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        tmp.close()
        shutil.copy(path, tmp.name)
        try:
            n = C.Master().load(tmp.name)      # raises when it is not an SFA "HS CODE & PRODUCT CODE" list
        finally:
            os.remove(tmp.name)
        return {i: {} for i in range(n)}       # count shown after the upload = master list records
    STATE['master_dir'] = os.path.dirname(os.path.abspath(path))
    out = {}
    for k, v in _read_learned().items():
        out[v.get('desc') or k] = dict(desc=v.get('desc', ''), hs=v.get('hs', ''), uom='', coo='', casc=v.get('pc1', ''),
                                       pc2=v.get('pc2', ''), brand='', remarks='')
    return out


def save_master(path, master):
    """Codes saved from the Item tab -> learned codes (the SFA list itself is never changed)."""
    STATE['master_dir'] = STATE['master_dir'] or os.path.dirname(os.path.abspath(path))
    learned = _read_learned()
    for item, d in master.items():
        learned[C.key(item)] = dict(desc=d.get('desc') or item, hs=re.sub(r'\D', '', d.get('hs') or ''),
                                    pc1=(d.get('casc') or '').upper(), pc2=(d.get('pc2') or '').upper())
    with open(_learned_path(), 'w', encoding='utf-8') as f:
        json.dump(learned, f, indent=1, ensure_ascii=False)


def parse_cargo(text):
    import permit_engine as P_
    return P_.parse_cargo(text or '')


def make_ctx(cfg, lists):
    return dict(master={}, hs_list={}, countries=M.Countries(lists.get('List-Country', [])), cargo=cfg.get('cargo') or {},
                party_name=cfg.get('party_name', ''), party_uen=cfg.get('party_uen', ''))


# ---------------------------------------------------------------- reading
def _xlsx_kind(path):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return ''
    names = [s.strip().upper() for s in wb.sheetnames]
    kind = ''
    if 'HEADER' in names and 'ITEM' in names:
        kind = 'PREV'
    else:
        for ws in wb.worksheets[:3]:
            for r in ws.iter_rows(max_row=15, values_only=True):
                up = [C.nfkc(c).strip().upper() for c in r if c is not None]
                if 'HS CODE' in up and any('DESCRIPTION' in c for c in up) and any('PRODUCT CODE' in c for c in up):
                    kind = 'MASTER'
                    break
                if any('DESCRIPTION' in c for c in up) and any('AMOUNT' in c for c in up):
                    kind = 'INV'
                    break
            if kind:
                break
    wb.close()
    return kind


def parse_pdf(path, countries=None, log=None):
    name = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        k = _xlsx_kind(path)
        if k == 'MASTER':
            dst = os.path.join(_sfa_dir(), name)
            if os.path.abspath(dst) != os.path.abspath(path):
                shutil.copy(path, dst)          # kept for every next job
            m = C.Master()
            n = m.load(path)
            return [dict(kind='MASTER', file=name, path=dst, records=n, notes=[])]
        if k == 'PREV':
            return [dict(kind='PREV', file=name, path=path, prev=_prev_permit(path), notes=[])]
        if k == 'INV':
            inv = C.read_invoice(path)
            inv['invoice_date'] = inv['invoice_date'].strftime('%d/%m/%Y') if inv.get('invoice_date') else ''
            return [dict(kind='INV', file=name, path=path, inv=inv, invoice_no=inv.get('invoice_no', ''),
                         currency=inv.get('currency', ''), total=sum(l['amount'] or 0 for l in inv['lines']),
                         lines=inv['lines'], notes=[])]
        return [dict(kind='', file=name, path=path, notes=['Excel not recognised (GLT invoice / SFA master list / previous permit)'])]
    if ext == '.pdf':
        a = C.read_awb(path)
        a['arrival'] = a['arrival'].strftime('%d/%m/%Y') if a.get('arrival') else ''
        a.pop('text', None)
        if not (a.get('mawb') or a.get('flight')):
            return [dict(kind='', file=name, path=path, notes=['PDF not recognised as an air waybill (scanned? - type MAWB, '
                                                               'flight, date, weight, packages and freight on the Header tab)'])]
        return [dict(kind='AWB', file=name, path=path, awb=a, hawb=a.get('mawb', ''), notes=[])]
    return [dict(kind='', file=name, path=path, notes=['not a GLT invoice Excel / air waybill PDF / SFA master list'])]


def _prev_permit(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = dict(header={}, invoice={}, brand='')
    for sheet, sec in (('Header', 'header'), ('Invoice', 'invoice')):
        if sheet in wb.sheetnames:
            rows = list(wb[sheet].iter_rows(max_row=2, values_only=True))
            if len(rows) == 2:
                out[sec] = {C.nfkc(a).strip(): b for a, b in zip(rows[0], rows[1]) if a and b not in (None, '')}
    if 'Item' in wb.sheetnames:
        rows = list(wb['Item'].iter_rows(max_row=2, values_only=True))
        if len(rows) == 2:
            d = {C.nfkc(a).strip(): b for a, b in zip(rows[0], rows[1]) if a}
            out['brand'] = str(d.get('Brand') or '')
    wb.close()
    return out


def build_shipments(docs):
    issues = []
    invs = [d for d in docs if d.get('kind') == 'INV']
    awbs = [d for d in docs if d.get('kind') == 'AWB']
    prevs = [d for d in docs if d.get('kind') == 'PREV']
    masters = [d for d in docs if d.get('kind') == 'MASTER']
    for d in docs:
        if not d.get('kind'):
            issues.append((d['file'], 'CHECK', '; '.join(d.get('notes') or ['not recognised'])))
    for d in masters:
        issues.append((d['file'], 'CHECK', 'SFA master list: %d records - kept for every next job' % d['records']))
    ships = collections.OrderedDict()
    for d in invs:
        key = 'INV:' + (d['invoice_no'] or d['file'])
        aw = next((a for a in awbs if a['awb'].get('invoice_no') == d['invoice_no']), None) or \
            (awbs[0] if len(awbs) == 1 or len(invs) == 1 else None)
        ships[key] = dict(key=key, invs=[d], pls=[], awbs=[aw] if aw else [], prev=prevs[0]['prev'] if prevs else None,
                          hawb='', hawb_src='', include=True, direction='IMPORT', overrides={}, items=[], issues=[],
                          party='', party_src='', party_uen='', party_uen_src='', dest='', dest_src='',
                          gw_override=None, pkg_override=None, extra_masters=[m['path'] for m in masters])
    if not invs:
        issues.append(('-', 'ERROR', 'No GLT invoice Excel in this job'))
    return ships, issues


def ship_label(sh):
    return sh['invs'][0]['invoice_no'] if sh['invs'] else sh['key']


# ---------------------------------------------------------------- rules
def _awb_of(sh):
    return (sh['awbs'][0]['awb'] if sh.get('awbs') else None) or {}


def resolve(sh, ctx):
    inv = sh['invs'][0]['inv']
    awb = _awb_of(sh)
    master, bad = sfa_master(sh.get('extra_masters') or [])
    learned = {k: v for k, v in _read_learned().items()}
    team = str(ctx.get('team') or '').strip().upper()
    rules = sfa_rules() if team == SFA.TEAM else None    # SFA team: the SFA project rules + database
    if rules is not None and not rules.ready():
        issues_sfa = [('MISSING', 'SFA team database is empty - Settings > Product masters > SFA team > Rebuild')]
        rules = None
    else:
        issues_sfa = []
    prev = sh.get('prev') or {}
    brand = prev.get('brand') or 'GLT CO.'
    issues = list(sh['issues']) + issues_sfa
    qty_nw = False
    if rules is not None:
        hd_prev = (prev.get('header') or {})
        imp_names = [sh.get('party') if sh.get('party_src') == 'Entered' else '', hd_prev.get('Importer'),
                     inv.get('consignee'), C.DEFAULT_HEADER.get('Importer') if not prev and not inv.get('consignee') else '']
        qty_nw = rules.qty_is_nw(*[x for x in imp_names if x])
        issues.append(('CHECK', 'SFA team rules: codes from saved codes > SFA overrides > SFA master data > past permits; '
                                'net weight = %s; HS / CASC qty by HS / product UOM (NMB, KGM, TNE, TEN)'
                       % ('invoice (Quantity) - importer rule (%s)' % (inv.get('consignee') or hd_prev.get('Importer') or 'importer')
                          if qty_nw else 'quantity x standard (KG lines = quantity)')))
    for b in bad:
        issues.append(('CHECK', 'SFA master list not read - ' + b))
    if not master and rules is None:
        issues.append(('MISSING', 'No SFA master list - add it on Settings > Product masters (Japan food) or with the '
                                  'documents; HS / product codes are empty'))
    raw_inv = dict(inv, invoice_date=None)
    items = C.build_items(raw_inv, master, learned, brand)
    for n, (it, ln) in enumerate(zip(items, inv['lines'])):
        k = '%s|%s' % (ln['sheet'], ln['row'])
        ov = sh['overrides'].get(k, {})
        if rules is not None:
            SFA.apply_codes(it, rules, keep_codes=(it.get('status') == 'SAVED CODE'), qty_as_nw=qty_nw)
        if 'hs' in ov:
            it['hs'], it['status'] = ov['hs'], 'EDITED'
        if 'casc' in ov:
            it['pc1'], it['status'] = ov['casc'], 'EDITED'
        if 'nw' in ov and ov['nw'] not in (None, ''):
            it['kg'] = fnum(ov['nw'])
        if 'coo' in ov and ov['coo']:
            it['origin'] = ov['coo'].upper()
        if 'desc' in ov and ov['desc']:
            it['desc'] = ov['desc']
        if ov.get('brand'):
            it['brand'] = it['brand_ov'] = ov['brand']
        C.recalc(it)
        if rules is not None:
            SFA.apply_qty(it, rules, (ov.get('uom') or '') if 'uom' in ov else '')
        elif 'uom' in ov and ov['uom']:                   # HS UOM typed: HS qty in that unit
            it['uom'] = ov['uom'].upper()
            kg = it.get('kg')
            q = kg / 1000.0 if it['uom'] == 'TNE' else kg if it['uom'] == 'KGM' else kg * 1000 if it['uom'] == 'GRM' \
                else it.get('inv_qty')
            it['hs_qty'] = C.fmt_qty(q) if q is not None else None
        row = C.HS_TABLE.get(it['hs']) if it['hs'] else None
        it.update(keys=[k], item=it['desc'].replace(', ', ' '), description=it['desc'], hs_uom=it['uom'],
                  hs_src='Edited' if 'hs' in ov else (it.get('sfa_src') or '') if rules is not None and it['hs']
                  else it['status'].title() if it['hs'] else '',
                  coo=it.get('origin') or 'JP', coo_src='Edited' if 'coo' in ov else 'Japan food rule',
                  casc=it.get('pc1', ''), casc_src='Edited' if 'casc' in ov else ((it.get('sfa_src') if rules is not None else
                                                                                  'Master list') if it.get('pc1') else ''),
                  nw=it.get('kg'), nw_src='Edited' if 'nw' in ov else (it.get('nw_rule') or 'Invoice quantity (net KG)'), gw=None,
                  qty=it.get('inv_qty'), uom_inv=it.get('inv_unit'), invoice=inv.get('invoice_no', ''),
                  currency=inv.get('currency') or 'JPY', price=None, lines=[ln['sheet'] + ' row %s' % ln['row']],
                  inv_desc=ln['name'], doc_hs='', controlled=bool(row and row[3] == 'Control'), hs_row=(), pcodes=[],
                  order=it.get('customer', ''), delivery='', src=sh['invs'][0]['file'])
        it['uom'] = it['inv_unit']                       # the server pages show the invoice unit as "UOM"
        it['hs_uom'] = it['hs_uom'] or ''
        for w in it.get('sfa_notes') or []:
            issues.append(('CHECK', 'Line %d %s: %s' % (n + 1, it['desc'], w)))
        for w in it['warnings']:
            issues.append(('MISSING' if w in ('NO HS', 'NO NET WEIGHT') else 'CHECK', 'Line %d %s: %s' % (n + 1, it['desc'], w)))
        if it['status'].startswith('FOUND (OTHER'):
            issues.append(('CHECK', 'Line %d %s: code taken from the same item of another prefecture' % (n + 1, it['desc'])))
    for x in C.check_report([dict(i, uom=i['hs_uom']) for i in items], inv, awb):
        if x.startswith('Invoice sub-totals') or x.startswith('AWB invoice'):
            issues.append(('MISMATCH', x))
    sh['items'] = items
    # ---- shipment values for the server pages
    hd0 = dict(C.DEFAULT_HEADER, **((prev.get('header')) or {}))
    sh['direction'] = 'IMPORT'
    if sh.get('party_src') != 'Entered':
        sh['party'], sh['party_src'] = str(hd0.get('Importer') or ''), 'Previous permit' if prev else 'Tool default'
    if sh.get('party_uen_src') != 'Entered':
        sh['party_uen'], sh['party_uen_src'] = str(hd0.get('Importer Uen') or ''), sh['party_src']
    sh['hawb'] = sh.get('hawb') if sh.get('hawb_src') == 'Entered' else (awb.get('mawb') or '')
    sh['hawb_src'] = sh.get('hawb_src') if sh.get('hawb_src') == 'Entered' else ('Air waybill' if awb.get('mawb') else '')
    sh['currency'] = inv.get('currency') or 'JPY'
    sh['incoterm'] = inv.get('term') or 'FOB'
    sh['invoice_total'] = round(sum(i['amount'] or 0 for i in items), 2)
    sh['total_gw'] = sh.get('gw_override') or awb.get('gw')
    sh['gw_src'] = 'typed' if sh.get('gw_override') else 'Air waybill' if awb.get('gw') else ''
    sh['packages'] = sh.get('pkg_override') or awb.get('pkgs') or inv.get('cartons')
    sh['lines_n'] = len(items)
    sh['containers'] = []
    sh['charges'] = 0
    sh['cargo'] = {k: v for k, v in dict(mawb=awb.get('mawb', ''), flight=awb.get('flight', ''), date=awb.get('arrival', ''),
                                         loading_code=awb.get('loading_code') or hd0.get('Loading Port Code', ''),
                                         loading_name=awb.get('loading_name') or hd0.get('Loading Port Name', ''),
                                         release=hd0.get('Release Location Code', ''), receipt=hd0.get('Receipt Location Code', ''),
                                         permit_type='IG', mode='Air').items() if v}
    import permit_engine as P_
    al = P_.AIRLINES.get((awb.get('flight') or '')[:2])
    if al:
        sh['cargo']['transport_id'] = al
    if not awb:
        issues.append(('MISSING', 'No air waybill - type MAWB, flight, arrival date, gross weight, packages and freight '
                                  'on the Header tab'))
    sh['issues_all'] = list(dict.fromkeys(issues))
    return sh


# ---------------------------------------------------------------- output
HDR_MAP = (('mawb', 'Inward MAWB'), ('flight', 'In Flight Number'), ('transport_id', 'In Transport Identifier'),
           ('loading_port_code', 'Loading Port Code'), ('loading_port_name', 'Loading Port Name'),
           ('release_code', 'Release Location Code'), ('release_name', 'Release Location Name'),
           ('receipt_code', 'Receipt Location Code'), ('receipt_name', 'Receipt Location Name'),
           ('carrier_agent', 'Inward Carrier Agent Name'), ('carrier_agent_uen', 'Inward Carrier Agent Uen'),
           ('forwarder', 'Freight Forwarder'), ('forwarder_uen', 'Freight Forwarder Uen'),
           ('outer_pack_uom', 'Total Outer Pack Uom'))


def file_name_for(group, cfg):
    s0 = group[0]
    if s0.get('file_name'):
        return M.safe_name(s0['file_name'])
    return M.safe_name((s0['invs'][0].get('invoice_no') or 'PERMIT') + '_PERMIT')


def _date(v):
    return C.parse_date(v) if v else None


def generate(ships, cfg, log=print):
    out = cfg['out_dir']
    os.makedirs(out, exist_ok=True)
    files = []
    lists = M.template_lists(cfg['template'])
    for sh in ships:
        if not sh.get('include', True):
            continue
        inv = sh['invs'][0]['inv']
        awb = _awb_of(sh)
        prev = sh.get('prev') or {}
        header = dict(C.DEFAULT_HEADER, **{k: v for k, v in (prev.get('header') or {}).items()
                                           if k not in ('Inward MAWB', 'Arrival Date', 'In Flight Number', 'Total Gross Weight',
                                                        'Total Outer Pack Qty', "Trader's Remarks", 'Cross Reference')})
        invoice = dict(C.DEFAULT_INVOICE, **{k: v for k, v in (prev.get('invoice') or {}).items()
                                             if k in ('Term Type', 'Ad Valorem Indicator', 'Preferential Duty Indicator',
                                                      'Supplier/Manufacturer Name', 'Insurance Charge %', 'Insurance Charge Currency',
                                                      'Supplier/Importer Relationship')})
        # documents
        if awb.get('mawb'):
            header['Inward MAWB'] = awb['mawb']
        if awb.get('arrival'):
            header['Arrival Date'] = _date(awb['arrival'])
        if awb.get('flight'):
            header['In Flight Number'] = awb['flight']
        if awb.get('loading_code'):
            header['Loading Port Code'] = awb['loading_code']
            if awb.get('loading_name'):
                header['Loading Port Name'] = awb['loading_name']
        if sh.get('total_gw') is not None:
            header['Total Gross Weight'] = sh['total_gw']
        if sh.get('packages') is not None:
            header['Total Outer Pack Qty'] = int(sh['packages']) if float(sh['packages']).is_integer() else sh['packages']
        header['Importer'], header['Importer Uen'] = sh.get('party') or header.get('Importer'), sh.get('party_uen') or header.get('Importer Uen')
        # Header tab / Requirements of the job (blank = documents)
        hd = cfg.get('header') or {}
        for k, col in HDR_MAP:
            if hd.get(k):
                header[col] = hd[k]
        if hd.get('date'):
            header['Arrival Date'] = _date(hd['date']) or header.get('Arrival Date')
        pt = (cfg.get('permit_type') or '').upper()
        if pt in PERMIT_SCHEME:
            header['Declaration Type Code'] = PERMIT_SCHEME[pt]
        if hd.get('declaration_type'):
            code = M.lookup(lists.get('List-DeclarationType', []), hd['declaration_type'])
            if code:
                header['Declaration Type Code'] = code
        for k, col in (('cargo_pack_type', 'Cargo Pack Type Code'), ('transport', 'Inward Transport Mode')):
            if hd.get(k):
                lst = 'List-CargoPacktype' if k == 'cargo_pack_type' else 'List-TransportMode'
                code = M.lookup(lists.get(lst, []), hd[k])
                if code:
                    header[col] = code
        header.update({k: v for k, v in (cfg.get('extra_header') or {}).items() if str(v).strip()})
        # invoice sheet
        if inv.get('invoice_date'):
            invoice['Invoice Date'] = _date(inv['invoice_date'])
        if inv.get('term'):
            invoice['Term Type'] = inv['term']
        if inv.get('currency'):
            invoice['Total Invoice Currency'] = invoice['Freight Charge Currency'] = inv['currency']
        invoice['Total Invoice Amount'] = round(sum(i['amount'] or 0 for i in sh['items']), 2)
        if awb.get('freight') is not None:
            invoice['Freight Charge Amount'] = awb['freight']
        invoice.update({k: v for k, v in (cfg.get('extra_invoice') or {}).items() if str(v).strip()})
        for it in sh['items']:
            it['brand'] = it.get('brand_ov') or prev.get('brand') or it.get('brand') or 'GLT CO.'
            it['currency'] = invoice.get('Total Invoice Currency') or 'JPY'
            it['uom_inv'], it['uom'] = it.get('uom'), it.get('hs_uom') or ''
        inv_no = inv.get('invoice_no') or 'PERMIT'
        base = file_name_for([sh], cfg)
        split = bool(cfg.get('max_items'))
        try:
            paths = C.generate(out, base, header, invoice, sh['items'], inv_no, split=split)
        finally:
            for it in sh['items']:
                it['uom'] = it.pop('uom_inv', it.get('uom'))
        groups = []
        for it in sh['items']:
            if it['group'] not in groups:
                groups.append(it['group'])
        groups = sorted(groups, key=lambda x: 'FHZ'.find(x) if x in 'FHZ' else 9)
        for n, p in enumerate(paths):
            sel = sh['items'] if not split else [it for it in sh['items'] if it['group'] == groups[n]]
            files.append(dict(name=os.path.basename(p), path=p, ships=[sh], items=sel, header_blank=[], part=None))
            log('%s  -  IMPORT, %d item row(s), %s %s' % (os.path.basename(p), len(sel), invoice.get('Total Invoice Currency', ''),
                                                          '{:,.0f}'.format(sum(i['amount'] or 0 for i in sel))))
    rep = os.path.join(out, cfg.get('report_name') or 'CHECK_REPORT.xlsx')
    _report(rep, ships)
    log('%s  -  check report' % os.path.basename(rep))
    zpath = os.path.join(out, 'EXCEL_ALL.zip')
    with zipfile.ZipFile(zpath, 'w', zipfile.ZIP_DEFLATED) as z:
        for f in files:
            z.write(f['path'], f['name'])
    log('EXCEL_ALL.zip  -  %d Excel file(s)' % len(files))
    return files, rep, zpath


def _report(path, ships):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Issues'
    ws.append(['Invoice', 'Level', 'Message'])
    for sh in ships:
        for l, m in sh.get('issues_all') or []:
            ws.append([ship_label(sh), l, m])
    wi = wb.create_sheet('Items')
    wi.append(['Invoice', '#', 'Description', 'HS', 'HS from', 'Product code 1', 'Product code 2', 'Net KG', 'HS Qty', 'HS UOM',
               'Prefecture', 'Amount', 'Group', 'Customer', 'Status', 'Warnings'])
    for sh in ships:
        for n, it in enumerate(sh['items'], 1):
            wi.append([ship_label(sh), n, it['desc'], it['hs'], it.get('hs_src'), it.get('pc1'), it.get('pc2'), it.get('kg'),
                       it.get('hs_qty'), it.get('hs_uom'), it.get('pref_code'), it['amount'], it.get('group'), it.get('customer'),
                       it.get('status'), ', '.join(it.get('warnings') or [])])
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True)
    wb.save(path)


def write_raw(path, docs, ships=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Invoice lines'
    ws.append(['File', 'Sheet', 'Row', 'Description', 'Prefecture', 'Standard', 'Std unit', 'Quantity (net KG)', 'Unit',
               'Unit price', 'Amount', 'Customer'])
    for d in docs:
        for l in d.get('lines') or []:
            ws.append([d['file'], l.get('sheet'), l.get('row'), l.get('name'), l.get('pref'), l.get('std'), l.get('std_unit'),
                       l.get('qty'), l.get('unit'), l.get('price'), l.get('amount'), l.get('customer')])
    wa = wb.create_sheet('Air waybill')
    for d in docs:
        if d.get('kind') == 'AWB':
            for k, v in d['awb'].items():
                wa.append([d['file'], k, json.dumps(v) if isinstance(v, (dict, list)) else v])
    wb.save(path)
