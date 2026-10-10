# -*- coding: utf-8 -*-
"""MSD PERMIT TOOL - engine
Reads MSD commercial / tax invoices + packing lists (PDF), matches them by Delivery
number, and fills the dec-upload template (Header / Item / Invoice / Casc, + Container)
- one Excel per HAWB - plus a CHECK REPORT and EXCEL_ALL.zip.
"""
import os, re, math, zipfile, datetime, collections, json
import logging as _logging
_logging.getLogger('pdfminer').setLevel(_logging.ERROR)   # harmless font warnings (FontBBox ...) not shown
from xml.sax.saxutils import escape


def tne_round(v):
    """Container weight in TNE as a round figure: x.5 and above -> up, below x.5 -> down (5.5 -> 6, 5.49 -> 5), at least 1."""
    return max(1, int(math.floor(round(float(v), 6) + 0.5)))


# ---------------------------------------------------------------- helpers (shared with RS COMPO TOOL)
def h8(v):
    """Normalise an HS code to an 8-digit string."""
    if v is None:
        return ''
    if isinstance(v, float):
        v = int(v)
    s = str(v).strip().replace('.', '').replace(' ', '')
    if s.isdigit():
        return s.zfill(8)
    return s


def num(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '').strip())
    except ValueError:
        return 0.0


def read_rows(path, sheet=None):
    """Return all rows (list of lists) of a sheet in .xlsx/.xlsm/.xls."""
    ext = os.path.splitext(path)[1].lower()
    if ext == '.xls':
        try:
            import xlrd
        except ImportError:
            raise RuntimeError('Reading .xls needs the xlrd package. Save the file as .xlsx or install xlrd.')
        book = xlrd.open_workbook(path)
        sh = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
        out = []
        for i in range(sh.nrows):
            row = []
            for j, c in enumerate(sh.row(i)):
                v = c.value
                if c.ctype == xlrd.XL_CELL_EMPTY:
                    v = None
                elif c.ctype == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                    v = int(v)
                row.append(v)
            out.append(row)
        return out
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    out = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return out


def find_header(rows, must, limit=40):
    """Find the header row containing all names in `must`; return (index, {name_lower: col})."""
    for i, r in enumerate(rows[:limit]):
        names = {str(v).strip().lower(): j for j, v in enumerate(r) if v is not None and str(v).strip()}
        if all(m.lower() in names for m in must):
            return i, names
    raise RuntimeError('Header row with %s not found' % ', '.join(must))




# ---------------------------------------------------------------- xlsx template writer (keeps every sheet/format)
def col2num(c):
    n = 0
    for ch in c:
        n = n * 26 + ord(ch) - 64
    return n


def num2col(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def xml_unescape(s):
    return s.replace('&lt;', '<').replace('&gt;', '>').replace('&quot;', '"').replace('&apos;', "'").replace('&amp;', '&')


class Sheet:
    ROW_RE = re.compile(r'<row\b[^>]*?(?:/>|>.*?</row>)', re.S)
    CELL_RE = re.compile(r'<c\b[^>]*?(?:/>|>.*?</c>)', re.S)

    def __init__(self, xml):
        self.xml = xml
        m = re.search(r'<sheetData\s*/>|<sheetData>(.*?)</sheetData>', xml, re.S)
        self.span = m.span()
        inner = m.group(1) or ''
        self.rows = collections.OrderedDict()
        for rm in self.ROW_RE.finditer(inner):
            rx = rm.group(0)
            rn = int(re.search(r'\br="(\d+)"', rx).group(1))
            head = re.match(r'<row\b[^>]*?(?=/?>)', rx).group(0)
            cells = collections.OrderedDict()
            for cm in self.CELL_RE.finditer(rx):
                cx = cm.group(0)
                ref = re.search(r'\br="([A-Z]+)\d+"', cx).group(1)
                cells[ref] = cx
            self.rows[rn] = [head, cells, False]

    def _row(self, rn):
        if rn not in self.rows:
            self.rows[rn] = ['<row r="%d"' % rn, collections.OrderedDict(), True]
        return self.rows[rn]

    def get(self, ref):
        col, rn = re.match(r'([A-Z]+)(\d+)', ref).groups()
        r = self.rows.get(int(rn))
        return r[1].get(col) if r else None

    def set(self, ref, value, style=None):
        if value is None or value == '':
            return
        col, rn = re.match(r'([A-Z]+)(\d+)', ref).groups()
        row = self._row(int(rn))
        old = row[1].get(col)
        s_attr = ''
        if old:
            if '<f>' in old or '<f ' in old:
                raise RuntimeError('Cell %s has a formula' % ref)
            m = re.search(r'\bs="(\d+)"', old)
            if m:
                s_attr = ' s="%s"' % m.group(1)
        if style is not None:
            s_attr = ' s="%d"' % style
        if isinstance(value, bool):
            value = str(value)
        if isinstance(value, (int, float)):
            cx = '<c r="%s"%s><v>%s</v></c>' % (ref, s_attr, repr(float(value)) if isinstance(value, float) else value)
        else:
            cx = '<c r="%s"%s t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (ref, s_attr, escape(str(value)))
        row[1][col] = cx
        row[2] = True

    def set_cached(self, ref, val):
        cx = self.get(ref)
        if not cx or '<f' not in cx:
            return
        col, rn = re.match(r'([A-Z]+)(\d+)', ref).groups()
        if re.search(r'<v>.*?</v>|<v/>', cx, re.S):
            cx = re.sub(r'<v>.*?</v>|<v/>', '<v>%s</v>' % escape(val), cx, count=1, flags=re.S)
        else:
            cx = cx.replace('</c>', '<v>%s</v></c>' % escape(val))
        self.rows[int(rn)][1][col] = cx

    def tostring(self):
        out = []
        for rn in sorted(self.rows):
            head, cells, changed = self.rows[rn]
            if changed:
                head = re.sub(r'\sspans="[^"]*"', '', head)
            body = ''.join(cells[c] for c in sorted(cells, key=col2num))
            out.append(head + ('>' + body + '</row>' if body else '/>'))
        sd = '<sheetData>' + ''.join(out) + '</sheetData>'
        xml = self.xml[:self.span[0]] + sd + self.xml[self.span[1]:]
        # dimension
        maxr = max(self.rows) if self.rows else 1
        maxc = max((col2num(c) for r in self.rows.values() for c in r[1]), default=1)
        xml = re.sub(r'<dimension ref="[^"]*"/>', '<dimension ref="A1:%s%d"/>' % (num2col(maxc), maxr), xml, count=1)
        return xml


class Template:
    def __init__(self, path):
        z = zipfile.ZipFile(path)
        self.infos = z.infolist()
        self.data = {i.filename: z.read(i.filename) for i in self.infos}
        z.close()
        self.sst = []
        if 'xl/sharedStrings.xml' in self.data:
            x = self.data['xl/sharedStrings.xml'].decode('utf-8')
            for si in re.findall(r'<si>(.*?)</si>', x, re.S):
                self.sst.append(xml_unescape(''.join(re.findall(r'<t[^>]*>(.*?)</t>', si, re.S))))
        wb = self.data['xl/workbook.xml'].decode('utf-8')
        rels = self.data['xl/_rels/workbook.xml.rels'].decode('utf-8')
        relmap = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship [^>]*?Id="([^"]+)"[^>]*?Target="([^"]+)"', rels)}
        relmap.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship [^>]*?Target="([^"]+)"[^>]*?Id="([^"]+)"', rels)})
        self.sheets = collections.OrderedDict()
        for m in re.finditer(r'<sheet [^>]*?name="([^"]+)"[^>]*?r:id="([^"]+)"', wb):
            tgt = relmap[m.group(2)].lstrip('/')
            self.sheets[xml_unescape(m.group(1))] = tgt if tgt.startswith('xl/') else 'xl/' + tgt
        self._add_styles()

    def _add_styles(self):
        st = self.data['xl/styles.xml'].decode('utf-8')
        ids = [int(x) for x in re.findall(r'numFmtId="(\d+)"', st)] + [163]
        a, b = max(ids) + 1, max(ids) + 2
        nf = '<numFmt numFmtId="%d" formatCode="0.0000"/><numFmt numFmtId="%d" formatCode="dd/mm/yyyy"/>' % (a, b)
        if '<numFmts' in st:
            st = re.sub(r'<numFmts count="(\d+)">', lambda m: '<numFmts count="%d">' % (int(m.group(1)) + 2), st, count=1)
            st = st.replace('</numFmts>', nf + '</numFmts>', 1)
        else:
            st = st.replace('<fonts', '<numFmts count="2">' + nf + '</numFmts><fonts', 1)
        m = re.search(r'<cellXfs count="(\d+)">', st)
        base = int(m.group(1))
        xfs = ('<xf numFmtId="%d" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
               '<xf numFmtId="2" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
               '<xf numFmtId="%d" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
               '<xf numFmtId="49" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
               '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1"><alignment wrapText="1"/></xf>') % (a, b)
        st = st.replace(m.group(0), '<cellXfs count="%d">' % (base + 5), 1)
        i = st.find('</cellXfs>')
        st = st[:i] + xfs + st[i:]
        self.data['xl/styles.xml'] = st.encode('utf-8')
        self.S4, self.S2, self.SDATE, self.STXT, self.SWRAP = base, base + 1, base + 2, base + 3, base + 4
        wb = self.data['xl/workbook.xml'].decode('utf-8')
        if 'fullCalcOnLoad' not in wb:
            wb = re.sub(r'<calcPr\b', '<calcPr fullCalcOnLoad="1"', wb, count=1)
        self.data['xl/workbook.xml'] = wb.encode('utf-8')

    def sheet(self, name):
        return Sheet(self.data[self.sheets[name]].decode('utf-8'))

    def headers(self, sh):
        """{header text lower: column letter} from row 1."""
        out = {}
        r = sh.rows.get(1)
        if not r:
            return out
        for col, cx in r[1].items():
            v = None
            m = re.search(r'<v>(.*?)</v>', cx, re.S)
            if 't="s"' in cx and m:
                v = self.sst[int(m.group(1))]
            elif 't="inlineStr"' in cx:
                v = xml_unescape(''.join(re.findall(r'<t[^>]*>(.*?)</t>', cx, re.S)))
            elif m:
                v = xml_unescape(m.group(1))
            if v:
                out.setdefault(v.strip().lower(), col)
        return out

    def save(self, path, changed):
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zo:
            for info in self.infos:
                n = info.filename
                data = changed.get(n, self.data[n])
                if isinstance(data, str):
                    data = data.encode('utf-8')
                zo.writestr(info, data)


def template_lists(path):
    """Dropdown lists from the template's List-* sheets: {sheet: [(text, code), ...]}."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        if ws.title.startswith('List-'):
            out[ws.title] = [(str(r[0]), '' if len(r) < 2 or r[1] is None else str(r[1]))
                             for r in ws.iter_rows(values_only=True) if r and r[0] is not None]
    wb.close()
    return out


def lookup(lst, text):
    for t, c in lst:
        if t == text:
            return c
    return ''



def sheet_clear_from(sh, first_row=2):
    """Remove every value below the header (keeps cell styles, formulas untouched)."""
    for rn, row in sh.rows.items():
        if rn < first_row:
            continue
        for col, cx in list(row[1].items()):
            if '<f>' in cx or '<f ' in cx:
                continue
            if '<v' in cx or '<is>' in cx:
                m = re.search(r'\bs="(\d+)"', cx)
                row[1][col] = '<c r="%s%d"%s/>' % (col, rn, ' s="%s"' % m.group(1) if m else '')
                row[2] = True


Sheet.clear_from = sheet_clear_from


# ================================================================ constants
DEFAULT_UOM_RULES = [
    ('KGM', 'NET WEIGHT', 'x', 1), ('TNE', 'NET WEIGHT', '/', 1000),
    ('NMB', 'QTY', 'x', 1), ('TEN', 'QTY', '/', 10), ('TPR', 'QTY', '/', 10),
    ('CEN', 'QTY', '/', 100), ('MIL', 'QTY', '/', 1000),
]
# invoice UOMs that are a plain piece count (QTY based HS UOM rules only use these)
COUNT_UOMS = {'EA', 'EACH', 'PC', 'PCS', 'PCE', 'NMB', 'UN', 'UNIT', 'UNITS', 'NOS', 'NO', 'ST', 'SET'}
KG_UOMS = {'KG', 'KGS', 'KGM'}

DEFAULT_DESC_TEMPLATE = '{ITEM} , {DESC} , QTY : {QTY} {UOM} , NW : {NW} KGM , GW : {GW} KGM'
DESC_VARS = ('ITEM', 'DESC', 'QTY', 'UOM', 'NW', 'GW', 'HS', 'COO', 'BATCH', 'INVOICE')
REMARK_VARS = ('INV', 'HAWB', 'DELIVERY', 'JOB', 'PO')

# Permit type -> Declaration type code (master prompt)
PERMIT_SCHEME = collections.OrderedDict([('IG', 'GST'), ('ME', 'APS'), ('IT', 'SFZ'), ('IM', 'SFZ'),
                                         ('OD', 'DRT'), ('OT', 'DRT')])
INSURANCE_PERMITS = ('ME', 'IM', 'IT', 'IG')
TERM_MAP = {'DAP': 'CIF', 'FCA': 'FOB'}

COUNTRY_ALIASES = {
    'SOUTH KOREA': 'KR', 'KOREA': 'KR', 'REPUBLIC OF KOREA': 'KR', 'KOREA, REPUBLIC OF': 'KR',
    'UNITED KINGDOM': 'GB', 'UK': 'GB', 'GREAT BRITAIN': 'GB', 'ENGLAND': 'GB',
    'USA': 'US', 'UNITED STATES': 'US', 'UNITED STATES OF AMERICA': 'US', 'U.S.A.': 'US',
    'NETHERLANDS': 'NL', 'THE NETHERLANDS': 'NL', 'HOLLAND': 'NL', 'VIETNAM': 'VN', 'VIET NAM': 'VN',
    'TAIWAN': 'TW', 'HONG KONG': 'HK', 'CHINA': 'CN', "PEOPLE'S REPUBLIC OF CHINA": 'CN',
    'RUSSIA': 'RU', 'CZECH REPUBLIC': 'CZ', 'UAE': 'AE', 'UNITED ARAB EMIRATES': 'AE',
    'IRELAND': 'IE', 'SINGAPORE': 'SG', 'AUSTRALIA': 'AU', 'NEW ZEALAND': 'NZ', 'JAPAN': 'JP',
    'MALAYSIA': 'MY', 'THAILAND': 'TH', 'INDONESIA': 'ID', 'PHILIPPINES': 'PH', 'INDIA': 'IN',
    'GERMANY': 'DE', 'FRANCE': 'FR', 'ITALY': 'IT', 'SPAIN': 'ES', 'SWITZERLAND': 'CH', 'BELGIUM': 'BE',
    'CANADA': 'CA', 'BRAZIL': 'BR', 'MEXICO': 'MX', 'SWEDEN': 'SE', 'DENMARK': 'DK', 'AUSTRIA': 'AT',
    'PUERTO RICO': 'PR', 'ISRAEL': 'IL', 'TURKEY': 'TR', 'SAUDI ARABIA': 'SA', 'SOUTH AFRICA': 'ZA',
}

UOM_RULES = collections.OrderedDict()


def set_uom_rules(rules=None):
    global UOM_RULES
    out = collections.OrderedDict()
    for r in (rules if rules else DEFAULT_UOM_RULES):
        try:
            uom, basis, calc, factor = r
            uom = str(uom).strip().upper()
            basis = 'NET WEIGHT' if str(basis).strip().upper().startswith('N') else 'QTY'
            calc = '/' if str(calc).strip() in ('/', 'divide', 'DIVIDE') else 'x'
            factor = float(str(factor).replace(',', '').strip())
            if uom and factor > 0:
                out[uom] = (basis, calc, factor)
        except (ValueError, TypeError):
            continue
    UOM_RULES = out
    return out


def rule_text(uom):
    r = UOM_RULES.get(uom)
    if not r:
        return '(no rule)'
    basis, calc, f = r
    b = 'net weight' if basis == 'NET WEIGHT' else 'qty'
    return b if f == 1 else '%s %s %g' % (b, '*' if calc == 'x' else '/', f)


set_uom_rules()


# ================================================================ small helpers
NUM = r'\d[\d,]*(?:\.\d+)?'


def fnum(s):
    try:
        return float(str(s).replace(',', '').strip())
    except (TypeError, ValueError):
        return None


def fmt_plain(v):
    """72000.0 -> '72000', 7567.172 -> '7567.172' (no thousands separator)."""
    if v is None:
        return ''
    s = ('%.6f' % v).rstrip('0').rstrip('.')
    return s


def fmt_thou(v):
    """2088.0 -> '2,088', 10105.602 -> '10,105.602'."""
    if v is None:
        return ''
    s = '{:,.3f}'.format(round(v, 3)).rstrip('0').rstrip('.')
    return s


def clean(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def squash(s):
    return re.sub(r'[^A-Z0-9]', '', str(s or '').upper())


def pdf_text(path):
    """Text of all pages, layout preserved (pypdf layout mode)."""
    from pypdf import PdfReader
    r = PdfReader(path)
    out = []
    for p in r.pages:
        try:
            t = p.extract_text(extraction_mode='layout')
        except Exception:                       # some PDFs (e.g. B/L drafts) break layout mode
            try:
                t = p.extract_text()
            except Exception:
                t = ''
        out.append(t or '')
    return '\n'.join(out)


# ================================================================ country lookup
class Countries:
    def __init__(self, pairs=None):
        self.code2name = {}
        self.name2code = {}
        for code, name in (pairs or []):
            code, name = str(code).strip().upper(), str(name).strip().upper()
            if len(code) == 2:
                self.code2name[code] = name
                self.name2code[name] = code
        for n, c in COUNTRY_ALIASES.items():
            self.name2code.setdefault(n, c)
            self.code2name.setdefault(c, n)
        self.names = sorted(self.name2code, key=len, reverse=True)

    def code(self, text):
        t = clean(text).upper().strip(' .,:')
        if not t:
            return ''
        if len(t) == 2 and t in self.code2name:
            return t
        if t in self.name2code:
            return self.name2code[t]
        for n in self.names:
            if re.search(r'(?<![A-Z])%s(?![A-Z])' % re.escape(n), t):
                return self.name2code[n]
        return ''

    def find_in(self, text, exclude=('SG',)):
        """Country code of the last country name found in text (excluding codes in exclude)."""
        best = None
        U = text.upper()
        for n in self.names:
            for m in re.finditer(r'(?<![A-Z])%s(?![A-Z])' % re.escape(n), U):
                c = self.name2code[n]
                if c in exclude:
                    continue
                if best is None or m.start() > best[0]:
                    best = (m.start(), c)
        return best[1] if best else ''


def block_after_label(lines, label_re, n=8, stop_re=r'Delivery\s*#|Incoterms|Invoice N|Order No|Due Date'):
    """Text of the column that starts at a label (e.g. 'Ship to:') for the next n lines."""
    for i, ln in enumerate(lines):
        m = re.search(label_re, ln, re.I)
        if not m:
            continue
        col = max(0, m.start() - 2)
        parts = [ln[m.end():]]
        for ln2 in lines[i + 1:i + 1 + n]:
            if re.search(stop_re, ln2):
                break
            parts.append(ln2[col:] if len(ln2) > col else '')
        return '\n'.join(parts)
    return ''


COMPANY_RE = re.compile(r'(PTE\.?\s*LTD|GMBH|LIMITED|\bLTD\b|\bINC\b|B\.V\.?|\bBRANCH\b|CORPORATION|\bLLC\b|\bAG\b|S\.A\.)', re.I)
UEN_RE = re.compile(r'Co(?:mpany|\.)?\s*Reg(?:istration|\.)?\s*(?:Number|No\.?)?\s*[:.]?\s*([0-9]{8,9}[A-Z]|[STR]\d{2}[A-Z]{2}\d{4}[A-Z])', re.I)


def find_exporter(lines, text):
    """Seller name (first company line at the top of the document) + UEN (Company Registration No)."""
    name = ''
    for ln in lines[:20]:
        part = re.split(r'\s{3,}', ln.strip())[0].strip()
        if part and COMPANY_RE.search(part) and not re.match(r'(Bill|Ship|Sold)', part, re.I):
            name = clean(part)
            break
    m = UEN_RE.search(text)
    return name, (m.group(1).upper() if m else '')


def block_lines(lines, label_re, n=7):
    """Lines of the column under a label (each cut at the next column gap)."""
    blk = block_after_label(lines, label_re, n)
    out = []
    for ln in blk.split('\n'):
        part = re.split(r'\s{3,}', ln.strip())
        part = [p for p in part if p and not re.fullmatch(r'I[CS]\d{3,}', p)]
        if part:
            out.append(clean(part[0]))
    return out


def consignee_from(lines_, countries):
    """(name, address1, address2) from ship-to block lines."""
    if not lines_:
        return '', '', ''
    name = lines_[0]
    rest = [l for l in lines_[1:] if not countries.code(l) or len(l) > 25]
    if lines_ and countries.find_in(lines_[-1], exclude=()) and len(lines_[-1]) <= 25 and lines_[-1] in rest:
        rest.remove(lines_[-1])
    a1 = rest[0] if rest else ''
    a2 = ', '.join(rest[1:]) if len(rest) > 1 else ''
    return name, a1, a2


# ================================================================ PDF parsers
def classify(text):
    head = text[:600].upper()
    if re.search(r'BILL\s+OF\s+LADING|B\.\s?L\.\s?No|SEA\s+WAYBILL|AIR\s+WAYBILL', text, re.I) and \
            re.search(r'Port\s+of\s+(loading|discharge)|Airport\s+of\s+Departure', text, re.I):
        return 'BL'
    if 'PACKING LIST' in head or ('HU #' in text and 'Material' in text):
        return 'PL'
    if 'INVOICE' in head:
        return 'INV'
    return ''


def find_hawb(text):
    m = re.search(r'H\.?\s?A\.?\s?W\.?\s?B\.?(?:\s*/\s*HBL)?\s*(?:number|no\.?|#)?\s*[:#]?\s*([A-Z]{2,5}\d{4,12})', text, re.I)
    return m.group(1).upper() if m else ''


def find_coo(text):
    for ln in text.splitlines():
        s = ln.strip()
        m = re.match(r'(?:COUNTRY OF ORIGIN|ORIGIN COUNTRY|MADE IN)\s*[:\-]?\s*([A-Za-z][A-Za-z ,.]+?)\s*$', s, re.I)
        if m:
            return m.group(1)
    return ''


def find_hts(text):
    m = re.search(r'(?:EXPORT\s+HTS|HTS|HS)\s*(?:NO\.?|NUMBER|CODE)\s*[:.]?\s*(\d{4}\.?\d{2}\.?\d{2,4})', text, re.I)
    return h8(m.group(1))[:8] if m else ''


DESC_WRAP_WIDTH = 30     # MSD/SAP invoices cut the description every 30 characters, also inside a word


def _join_desc(a, b, last_seg=None):
    """Join a wrapped invoice description.
    - previous line filled the full 30-character width -> the word was cut: no space
      ('CAPVAXIVE 0.5ML SYR 10X1DOSE A' + 'US' -> '...10X1DOSE AUS')
    - after '/', '(' or '-' -> no space ('TAB (FIS/FHP/' + 'SNG)')
    - otherwise one space."""
    a, b = clean(a), clean(b)
    seg = clean(last_seg if last_seg is not None else a)
    if len(seg) >= DESC_WRAP_WIDTH or a.endswith(('/', '(', '-')) or b.startswith(('/', ')')):
        return a + b
    return clean(a + ' ' + b)


LINE_A = re.compile(r'^\s*(?P<desc>\S.*?)(?:\s{2,}(?P<batch>[A-Z0-9]{4,}))?\s{2,}(?P<item>\d{5,10})\s+'
                    r'(?P<qty>%s)\s*(?P<uom>[A-Z]{1,4})\s+(?P<price>%s)\s+(?P<amt>%s)\s*$' % (NUM, NUM, NUM))
LINE_B = re.compile(r'^\s*(?P<item>\d{5,10})\s{2,}(?P<desc>\S.*?)\s{2,}(?P<batch>[A-Z0-9]{4,})\s+'
                    r'(?P<exp>\d{1,2}[-./][A-Z]{3}[-./]\d{4}|\d{2}[./]\d{2}[./]\d{4})\s+'
                    r'(?P<qty>%s)\s+(?P<uom>[A-Z]{1,4})\s+(?P<price>%s)\s+(?P<amt>%s)\s*$' % (NUM, NUM, NUM))
LINE_C = re.compile(r'^\s*(?P<desc>\S.*?)\s{2,}(?P<item>\d{5,10})\s+(?P<price>%s)\s+(?P<amt>%s)\s*$' % (NUM, NUM))
LINE_C_CONT = re.compile(r'^\s*(?:(?P<desc>\S.*?)\s{2,})?(?P<batch>[A-Z0-9]{6,})\s+(?P<qty>%s)\s*(?P<uom>[A-Z]{1,4})\s*$' % NUM)
LINE_B_CONT = re.compile(r'^\s*(?:(?P<desc>\S.*?)\s{2,})?(?P<batch>[A-Z0-9]{6,})\s+'
                         r'(?P<exp>\d{1,2}[-./][A-Z]{3}[-./]\d{4}|\d{2}[./]\d{2}[./]\d{4})\s+(?P<qty>%s)\s*$' % NUM)


def parse_invoice(text, path='', countries=None):
    countries = countries or Countries()
    lines = text.splitlines()
    inv = dict(kind='INV', file=os.path.basename(path), path=path, invoice_no='', date='', incoterm='',
               currency='', rate='', sap_order='', delivery='', total=None, lines=[], hawb=find_hawb(text),
               coo_text=find_coo(text), hs=find_hts(text), dest='', notes=[])
    # invoice number + date (value under the "Invoice No." label, nearest column)
    for i, ln in enumerate(lines):
        m0 = re.search(r'Invoice\s*(?:No\.?|Number)', ln, re.I)
        if not m0:
            continue
        m = re.search(r'Invoice\s*(?:No\.?|Number)\s*[:#]?\s*(\d{6,12})\b', ln, re.I)
        cand = [(0, m, ln)] if m else []
        if not cand:
            for ln2 in lines[i + 1:i + 4]:
                for t in re.finditer(r'(?<![\d.,/-])(\d{8,12})(?![\d.,/-])', ln2):
                    cand.append((abs(t.start() - m0.start()), t, ln2))
                if cand:
                    break
        if cand:
            _, t, src = min(cand, key=lambda x: x[0])
            inv['invoice_no'] = t.group(1)
            d = re.search(r'(\d{2}[./-]\d{2}[./-]\d{4}|\d{1,2}-[A-Z]{3}-\d{4})', src[t.end():])
            if d:
                inv['date'] = d.group(1)
            break
    m = re.search(r'Incoterms?\s*:?\s*([A-Z]{3})\b', text)
    if m:
        inv['incoterm'] = m.group(1)
    # delivery / SAP order / currency (header row + value row)
    for i, ln in enumerate(lines):
        if re.search(r'\bDelivery\b', ln) and re.search(r'\bCurrency\b', ln):
            for ln2 in lines[i + 1:i + 4]:
                if not ln2.strip():
                    continue
                toks = list(re.finditer(r'(?<![\d.,])\d{8,12}(?![\d.,])', ln2))
                if toks:
                    inv['delivery'] = toks[-1].group(0)
                    if len(toks) > 1:
                        inv['sap_order'] = toks[-2].group(0)
                    c = re.search(r'\b([A-Z]{3})\b', ln2[toks[-1].end():])
                    if c:
                        inv['currency'] = c.group(1)
                    r = re.search(r'Rate of\s*(%s)\s*([A-Z]{3})' % NUM, ln2)
                    if r:
                        inv['rate'] = '%s %s' % (r.group(1), r.group(2))
                    break
            break
    if not inv['delivery']:
        m = re.search(r'Delivery\s*(?:No\.?|#)?\s*[:#]\s*(\d{8,12})', text)
        if m:
            inv['delivery'] = m.group(1)
    if not inv['currency']:
        m = re.search(r'Currency\s*[:]\s*([A-Z]{3})', text)
        if m:
            inv['currency'] = m.group(1)
    # item lines
    hdr = None
    for i, ln in enumerate(lines):
        if re.search(r'Description', ln) and re.search(r'Quantity', ln) and re.search(r'Amount', ln):
            hdr = i
            break
    cur = None
    if hdr is not None:
        layout_b = bool(re.match(r'\s*Item\s*no', lines[hdr], re.I))
        for ln in lines[hdr + 1:]:
            s = ln.strip()
            if re.match(r'_{3,}|Remarks|Beneficiary|SALES VALUE|TOTAL VALUE|Payment Instructions', s):
                break
            if not s:
                continue
            m = (LINE_B if layout_b else LINE_A).match(ln)
            if not m and not layout_b:
                mc = LINE_C.match(ln)          # qty is on the batch lines below
                if mc:
                    d = mc.groupdict()
                    cur = dict(item=d['item'], desc=clean(d['desc']), batches=[], qty=0.0, uom='',
                               price=fnum(d['price']), amount=fnum(d['amt']), hs='', batch_qty=[], split_qty=True)
                    inv['lines'].append(cur)
                    continue
            if m:
                d = m.groupdict()
                cur = dict(item=d['item'], desc=clean(d['desc']), batches=[d.get('batch') or ''],
                           qty=fnum(d['qty']), uom=d['uom'], price=fnum(d['price']), amount=fnum(d['amt']),
                           hs='', batch_qty=[fnum(d['qty'])] if layout_b else [])
                inv['lines'].append(cur)
                continue
            if cur is None:
                continue
            hm = re.match(r'HS\s*(?:Number|No\.?|Code)\s*[:.]?\s*(\d[\d.]{5,12})', s, re.I)
            if hm:
                cur['hs'] = h8(hm.group(1))[:8]
                continue
            if cur.get('split_qty'):
                cm = LINE_C_CONT.match(ln)
                if cm:
                    if cm.group('desc'):
                        cur['desc'] = _join_desc(cur['desc'], cm.group('desc'), cur.get('_seg'))
                        cur['_seg'] = cm.group('desc')
                    cur['batches'].append(cm.group('batch'))
                    cur['batch_qty'].append(fnum(cm.group('qty')))
                    cur['qty'] = round(sum(cur['batch_qty']), 6)
                    cur['uom'] = cur['uom'] or cm.group('uom')
                    continue
            if layout_b:
                cm = LINE_B_CONT.match(ln)
                if cm:
                    if cm.group('desc'):
                        cur['desc'] = clean(cur['desc'] + ' ' + cm.group('desc'))
                    cur['batches'].append(cm.group('batch'))
                    cur['batch_qty'].append(fnum(cm.group('qty')))
                    continue
            if not re.search(r'\d{4,}', s) and len(s) < 40:      # wrapped description text
                cur['desc_wrap'] = cur.get('desc_wrap', '') + ' ' + s
                cur['desc'] = _join_desc(cur['desc'], s, cur.get('_seg'))
                cur['_seg'] = s
        for l in inv['lines']:
            if l.get('split_qty'):
                continue
            if l['batch_qty'] and len(l['batch_qty']) > 1:
                tot = round(sum(q or 0 for q in l['batch_qty']), 6)
                l['line_qty_first'] = l['qty']
                l['qty'] = tot
    m = re.search(r'Total\s+Quantity\s*:\s*(%s)\s*([A-Z]{1,4})' % NUM, text, re.I)
    inv['total_qty'] = (fnum(m.group(1)), m.group(2)) if m else None
    m = re.search(r'SAP\s+order\s+No\.?\s+from\s+Supplying\s+site[^:]*:\s*(\d{6,12})', text, re.I)
    inv['sto'] = m.group(1) if m else ''
    # invoice total
    for ln in lines:
        if re.search(r'TOTAL VALUE|TOTAL AMOUNT|GRAND TOTAL|INVOICE TOTAL', ln, re.I):
            nums = re.findall(NUM, ln)
            if nums:
                inv['total'] = fnum(nums[-1])
    if inv['total'] is None and inv['currency']:
        for ln in lines:
            m = re.match(r'^\s*%s\s+(%s)\s*$' % (inv['currency'], NUM), ln)
            if m:
                inv['total'] = fnum(m.group(1))
    # destination (ship-to country)
    blk = block_after_label(lines, r'Ship[\s-]*to\b:?', 7)
    inv['dest'] = countries.find_in(blk) if blk else ''
    inv['exporter'], inv['exporter_uen'] = find_exporter(lines, text)
    inv['consignee'], inv['consignee_addr1'], inv['consignee_addr2'] = consignee_from(
        block_lines(lines, r'Ship[\s-]*to\b:?', 5), countries)
    return inv


HU_RE = re.compile(r'^\s*(\d{9,})\s+(.*?)\s+(%s)\s+(%s)\s+(KG|KGS|LB|LBS)\b' % (NUM, NUM))
MAT_RE = re.compile(r'Material\s*#?\s*:\s*(\d{4,})\s+(\S.*?)(?:\s{2,}.*)?$')
QTY_RE = re.compile(r'\bQty\s*:\s*(%s)\s+([A-Z]{1,4})\b' % NUM)


def parse_packing_list(text, path='', countries=None):
    countries = countries or Countries()
    lines = text.splitlines()
    pl = dict(kind='PL', file=os.path.basename(path), path=path, delivery='', po='', hus=[], materials=collections.OrderedDict(),
              total_gw=None, total_nw=None, pallets=None, cartons=None, container='', seal='', cont_size='',
              hawb=find_hawb(text), coo_text=find_coo(text), hs=find_hts(text), dest='', notes=[])
    m = re.search(r'Delivery\s*#?\s*:\s*(\d{8,12})', text)
    if m:
        pl['delivery'] = m.group(1)
    m = re.search(r'(?:PO|Order No)\s*#?\s*:\s*(\d{6,12})', text)
    if m:
        pl['po'] = m.group(1)
    hu = None
    for ln in lines:
        m = HU_RE.match(ln)
        if m:
            f = 1.0 if m.group(5).startswith('K') else 0.45359237
            hu = dict(id=m.group(1), type=clean(m.group(2)), gw=fnum(m.group(3)) * f, nw=fnum(m.group(4)) * f,
                      mats=[], qtys=[])
            pl['hus'].append(hu)
            continue
        if hu is None:
            continue
        mm = MAT_RE.search(ln)
        if mm:
            hu['mats'].append((mm.group(1), clean(mm.group(2))))
        for qm in QTY_RE.finditer(ln):
            hu['qtys'].append((fnum(qm.group(1)), qm.group(2)))
    # weights per material
    for hu in pl['hus']:
        mats = hu['mats']
        if not mats:
            pl['notes'].append('HU %s has no material line' % hu['id'])
            continue
        qtys = hu['qtys']
        if len(mats) == 1:
            pairs = [(mats[0], (sum(q for q, _ in qtys) if qtys else 0.0, qtys[0][1] if qtys else ''))]
        else:
            pairs = list(zip(mats, qtys + [(0.0, '')] * (len(mats) - len(qtys))))
            pl['notes'].append('HU %s holds %d materials - its weight is split by quantity' % (hu['id'], len(mats)))
        tq = sum(p[1][0] or 0 for p in pairs) or 0
        for (item, desc), (q, u) in pairs:
            share = (q / tq) if tq else 1.0 / len(pairs)
            d = pl['materials'].setdefault(item, dict(item=item, desc=desc, qty=0.0, uom=u, nw=0.0, gw=0.0, hus=0))
            d['qty'] += q or 0
            d['nw'] += hu['nw'] * share
            d['gw'] += hu['gw'] * share
            d['hus'] += 1
            if u and not d['uom']:
                d['uom'] = u
    m = re.search(r'Gross Weight\s*:\s*(%s)\s*(KG|LB)?' % NUM, text, re.I)
    if m:
        pl['total_gw'] = fnum(m.group(1))
    m = re.search(r'Net Weight\s*:\s*(%s)' % NUM, text, re.I)
    if m:
        pl['total_nw'] = fnum(m.group(1))
    m = re.search(r'(?:Total\s+Count\s+)?Pallets\s*:\s*(\d+)', text, re.I)
    if m:
        pl['pallets'] = int(m.group(1))
    m = re.search(r'Cartons\s*:\s*(\d+)', text, re.I)
    if m:
        pl['cartons'] = int(m.group(1))
    m = re.search(r'Container\s*(?:No\.?)?\s*[:#]\s*([A-Z]{4}\s?\d{6,7})', text, re.I)
    if m:
        pl['container'] = m.group(1).replace(' ', '').upper()
    m = re.search(r'Seal\s*(?:No\.?)?\s*[:#]\s*([A-Z0-9-]{4,})', text, re.I)
    if m:
        pl['seal'] = m.group(1).upper()
    m = re.search(r'\b\d+\s*X\s*(20|40|45)\s*(?:FT|\'|FEET)', text, re.I)
    if m:
        pl['cont_size'] = m.group(1)
    blk = block_after_label(lines, r'Ship\s*to\s*:', 7)
    pl['dest'] = countries.find_in(blk) if blk else ''
    beh = []
    for i, ln in enumerate(lines):
        m = re.search(r'Ship on behalf of\s*:', ln, re.I)
        if m:
            for ln2 in lines[i + 1:i + 9]:
                part = re.split(r'\s{3,}', ln2.strip())[0].strip()
                if not part or re.match(r'(Tel|Fax|VAT|GST|Co\.)', part, re.I):
                    break
                beh.append(part)
            break
    pl['exporter'] = clean(' '.join(beh[:2])) if beh and not COMPANY_RE.search(beh[0]) else (beh[0] if beh else '')
    m = UEN_RE.search(text)
    pl['exporter_uen'] = m.group(1).upper() if m else ''
    pl['consignee'], pl['consignee_addr1'], pl['consignee_addr2'] = consignee_from(
        block_lines(lines, r'Ship\s*to\s*:', 6), countries)
    return pl


# ================================================================ bill of lading / waybill
PORT_CODES = {
    'SINGAPORE': 'SGSIN', 'ROTTERDAM': 'NLRTM', 'AMSTERDAM': 'NLAMS', 'ANTWERP': 'BEANR', 'HAMBURG': 'DEHAM',
    'BREMERHAVEN': 'DEBRV', 'FELIXSTOWE': 'GBFXT', 'SOUTHAMPTON': 'GBSOU', 'LONDON GATEWAY': 'GBLGP',
    'LIVERPOOL': 'GBLIV', 'LE HAVRE': 'FRLEH', 'GENOA': 'ITGOA', 'BARCELONA': 'ESBCN', 'VALENCIA': 'ESVLC',
    'BUSAN': 'KRPUS', 'PUSAN': 'KRPUS', 'INCHEON': 'KRINC', 'SYDNEY': 'AUSYD', 'MELBOURNE': 'AUMEL',
    'BRISBANE': 'AUBNE', 'FREMANTLE': 'AUFRE', 'AUCKLAND': 'NZAKL', 'YOKOHAMA': 'JPYOK', 'TOKYO': 'JPTYO',
    'KOBE': 'JPUKB', 'OSAKA': 'JPOSA', 'SHANGHAI': 'CNSHA', 'SHENZHEN': 'CNSZX', 'HONG KONG': 'HKHKG',
    'KAOHSIUNG': 'TWKHH', 'PORT KLANG': 'MYPKG', 'PENANG': 'MYPEN', 'LAEM CHABANG': 'THLCH', 'BANGKOK': 'THBKK',
    'JAKARTA': 'IDJKT', 'MANILA': 'PHMNL', 'HO CHI MINH': 'VNSGN', 'HAIPHONG': 'VNHPH', 'NHAVA SHEVA': 'INNSA',
    'CHENNAI': 'INMAA', 'JEBEL ALI': 'AEJEA', 'DUBAI': 'AEDXB', 'LOS ANGELES': 'USLAX', 'LONG BEACH': 'USLGB',
    'NEW YORK': 'USNYC', 'CHICAGO': 'USCHI', 'SAVANNAH': 'USSAV', 'MEMPHIS': 'USMEM', 'VANCOUVER': 'CAVAN',
    'SANTOS': 'BRSSZ',
}


def port_code(name, extra=None):
    """UN/LOCODE for a port name from the built-in table (+ extra {name: code}). '' when unknown."""
    n = clean(name).upper()
    if not n:
        return ''
    table = dict(PORT_CODES)
    table.update({k.upper(): v for k, v in (extra or {}).items()})
    if re.fullmatch(r'[A-Z]{5}', n) and n[:2].isalpha():
        return n if n in table.values() else ''
    for k in sorted(table, key=len, reverse=True):
        if re.search(r'(?<![A-Z])%s(?![A-Z])' % re.escape(k), n):
            return table[k]
    return ''


def _bl_words(path):
    """Words of page 1 with coordinates (pdfplumber), watermark letters removed. [] if not available."""
    try:
        import pdfplumber
    except ImportError:
        return []
    try:
        with pdfplumber.open(path) as pdf:
            out = []
            for pno, p in enumerate(pdf.pages[:2]):
                q = p.filter(lambda o: o.get('object_type') != 'char' or (o.get('upright', True) and o.get('size', 0) < 14))
                for w in q.extract_words():
                    w['top'] += pno * 2000
                    out.append(w)
            return out
    except Exception:
        return []


def field_below(words, label, lines=1, max_dy=45):
    """Text under a printed label (same column). label = 'Port of discharge' etc."""
    toks = label.lower().split()
    n = len(toks)
    for i in range(len(words) - n + 1):
        seq = words[i:i + n]
        if [w['text'].rstrip(':') for w in seq] != [t.rstrip(':') for t in label.split()]:
            continue
        if any(abs(w['top'] - seq[0]['top']) > 2 for w in seq):
            continue
        x0, top, bot = seq[0]['x0'] - 3, seq[0]['top'], max(w['bottom'] for w in seq)
        right = [w['x0'] for w in words if abs(w['top'] - top) < 2 and w['x0'] > seq[-1]['x1'] + 15]
        x1 = min(right) - 2 if right else 10000
        cand = [w for w in words if bot - 1 < w['top'] <= bot + max_dy and x0 <= w['x0'] < x1]
        rows = collections.OrderedDict()
        for w in sorted(cand, key=lambda w: (round(w['top']), w['x0'])):
            key = next((k for k in rows if abs(k - w['top']) < 2), w['top'])
            rows.setdefault(key, []).append(w)
        out = [' '.join(w['text'] for w in sorted(r, key=lambda w: w['x0'])) for r in rows.values()]
        return out[:lines] if lines > 1 else (out[0] if out else '')
    return [] if lines > 1 else ''


def parse_bl(text, path='', countries=None):
    countries = countries or Countries()
    words = _bl_words(path) if path else []
    plain = ' '.join(w['text'] for w in words) if words else text
    bl = dict(kind='BL', file=os.path.basename(path), path=path, bl_no='', vessel='', voyage='', pol='', pod='',
              place_receipt='', place_delivery='', forwarder='', shipper='', containers=[], gw=None, packages=None,
              package_unit='', deliveries=[], pos=[], hs='', notes=[])
    get = (lambda lab: field_below(words, lab)) if words else (lambda lab: '')
    bl['bl_no'] = (re.search(r'([A-Z]{3,5}\d{5,12})', get('B.L. No.') or '') or re.search(
        r'B\.?\s?L\.?\s*No\.?.*?\b([A-Z]{4}\d{5,10})\b', text, re.S) or re.search(r'([A-Z]{4}\d{5,10})', bl['file']) or
                   re.search('()', '')).group(1)
    bl['vessel'] = get('Vessel')
    bl['voyage'] = get('Voyage No.')
    bl['pol'] = get('Port of loading')
    bl['pod'] = get('Port of discharge')
    bl['place_receipt'] = get('Place of receipt')
    bl['place_delivery'] = get('Place of delivery')
    for k in ('pol', 'pod', 'place_receipt', 'place_delivery'):
        if re.search(r'Tel|Fax|\d{3,}|whichever|Merchant', bl[k] or '', re.I):
            bl[k] = ''
    fw = field_below(words, 'Forwarding agent', lines=3, max_dy=40) if words else []
    bl['forwarder'] = next((l for l in fw if COMPANY_RE.search(l)), fw[0] if fw else '')
    sh = field_below(words, 'Shipper', lines=2) if words else []
    bl['shipper'] = sh[0] if sh else ''
    if not words:                     # plain-text fallback: "Vessel Voyage No." then values on the next line
        m = re.search(r'Vessel\s+Voyage\s+No\.?\s*\n\s*(.+?)\s+(\S+)\s*\n', text)
        if m:
            bl['vessel'], bl['voyage'] = clean(m.group(1)), m.group(2)
        bl['notes'].append('Read without column positions (pdfplumber missing) - check vessel / ports')
    for m in re.finditer(r'\b([A-Z]{4}\d{7})\s+([A-Z0-9]{5,15})\s+(\d{2}[A-Z]{2})\s+(%s)\s*KG' % NUM, plain):
        c = dict(no=m.group(1), seal=m.group(2), type=m.group(3), weight=fnum(m.group(4)))
        if c['no'] not in [x['no'] for x in bl['containers']]:
            bl['containers'].append(c)
    m = re.search(r'(%s)\s*KG\b' % NUM, plain)
    if m:
        bl['gw'] = fnum(m.group(1))
    m = re.search(r'\b(\d+)\s*(PLT|PALLETS?|PKGS?|PACKAGES|CTNS?|CARTONS)\b', plain, re.I)
    if m:
        bl['packages'], bl['package_unit'] = int(m.group(1)), m.group(2).upper()
    m = re.search(r'DELIVERY\s+ORDER\s+NO\.?\s*:?\s*((?:\d{8,12}[,\s/&]*)+)', plain, re.I)
    if m:
        bl['deliveries'] = re.findall(r'\d{8,12}', m.group(1))
    bl['pos'] = sorted(set(re.findall(r'PO\s*#\s*(\d{6,12})', plain, re.I)))
    m = re.search(r'HS\s*CODE\s*:?\s*(\d{6,10})', plain, re.I)
    bl['hs'] = h8(m.group(1))[:8] if m else ''
    bl['hawb'] = bl['bl_no']
    return bl


# ================================================================ DHL e-mail instructions
LOCATION_NAMES = {'PPZ': 'PASIR PANJANG FTZ', 'ME': 'MAJOR EXPORTER SCHEME', 'KZ': 'KEPPEL FTZ', 'JZ': 'JURONG FTZ', 'TZ': 'TUAS TERMINAL GATEWAY', 'CZ': 'CHANGI FTZ', 'O': 'OTHERS'}
EMAIL_FIELDS = [  # (key, label regex)
    ('carrier', r'CARRIER'), ('booking_ref', r'CARRIER\s+BOOKING\s+REF[^:]*'), ('release', r'PLACE\s+OF\s+RELEASE'),
    ('receipt', r'PLACE\s+OF\s+RECEIPT'), ('permit_type', r'TYPE\s+OF\s+PERMIT'), ('co_type', r'CO\s+TYPE[^:]*'),
    ('permit_remarks', r'PERMIT\s+REMARKS'), ('special', r'SPECIAL\s+INSTRUCTIONS?'), ('exporter', r'EXPORTER'),
    ('uen', r'UEN\s*#?'), ('shipment_type', r'SHIPMENT\s+TYPE'), ('job_no', r'JOB\s*(?:NO|NUMBER|#)\.?'),
    ('staff', r'STAFF\s+EMAIL'),
]


def parse_email(text):
    """Read the forwarder's instruction e-mail ("LABEL: value" lines). Unknown lines are ignored."""
    out = {}
    if not text:
        return out
    for raw in text.splitlines():
        ln = raw.strip()
        if not ln:
            continue
        m = re.match(r'COO\s*:\s*([A-Z]{2})\b(?:.*?HS\s*CODE\s*:\s*([\d.]+))?', ln, re.I)
        if m:
            out['coo'] = m.group(1).upper()
            if m.group(2):
                out['hs'] = m.group(2).replace('.', '')
            continue
        m = re.match(r'UEN\s*(?:NO\.?)?\s*#?\s*:?\s*([0-9]{8,9}[A-Z]|[STR]\d{2}[A-Z]{2}\d{4}[A-Z])\b', ln, re.I)
        if m:
            out['uen'] = m.group(1).upper()
            continue
        m = re.match(r'HS\s*CODE\s*:\s*([\d.]+)', ln, re.I)
        if m:
            out['hs'] = m.group(1).replace('.', '')
            continue
        for key, lab in sorted(EMAIL_FIELDS, key=lambda x: -len(x[1])):
            m = re.match(r'(%s)\s*:\s*(.*)$' % lab, ln, re.I)
            if m:
                if key == 'carrier' and re.match(r'CARRIER\s+BOOKING', ln, re.I):
                    continue
                out.setdefault(key, clean(m.group(2)))
                break
        else:
            if re.fullmatch(r'NON[- ]TRANSPORT COMPANY|TRANSPORT COMPANY', ln, re.I):
                out['company_type'] = ln.upper()
    out['permit_type'] = (out.get('permit_type') or '').upper()[:2]
    return {k: v for k, v in out.items() if v}


def pdf_text_plumber(path):
    """Second reader (pdfplumber, layout kept) - used when pypdf gives unusable text."""
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return '\n'.join((p.extract_text(layout=True) or '') for p in pdf.pages)
    except Exception:
        return ''


def parse_pdf(path, countries=None):
    text = pdf_text(path)
    kind = classify(text)
    if kind in ('', 'PL'):
        alt = None
        if kind == '' or not re.search(r'^\s*\d{9,}\s+\S.*?\d+\.\d+\s+\d+\.\d+\s+KG', text, re.M):
            alt = pdf_text_plumber(path)
        if alt and classify(alt):
            text, kind = alt, classify(alt)
    if kind == 'BL':
        d = parse_bl(text, path, countries)
    elif kind == 'PL':
        d = parse_packing_list(text, path, countries)
    elif kind == 'INV':
        d = parse_invoice(text, path, countries)
    else:
        d = dict(kind='', file=os.path.basename(path), path=path,
                 notes=['Not recognised as invoice, packing list or bill of lading'])
    d['text'] = text
    return d


# ================================================================ masters
MASTER_COLS = ['Item No', 'Description', 'HS Code', 'HS UOM', 'COO', 'CASC Product Code 1', 'Brand', 'Remarks']


def read_email_file(path):
    """Text of an instruction e-mail saved as .txt / .eml (body) / .msg (needs extract_msg)."""
    ext = os.path.splitext(path)[1].lower()
    if ext == '.eml':
        import email
        from email import policy
        with open(path, 'rb') as f:
            msg = email.message_from_binary_file(f, policy=policy.default)
        body = msg.get_body(preferencelist=('plain', 'html'))
        t = body.get_content() if body else ''
        return re.sub(r'<[^>]+>', '\n', t)
    if ext == '.msg':
        try:
            import extract_msg
        except ImportError:
            raise RuntimeError('Reading .msg needs the extract_msg package - copy the e-mail text and paste it instead')
        return extract_msg.Message(path).body or ''
    if ext == '.pdf':
        return pdf_text(path)
    with open(path, encoding='utf-8', errors='replace') as f:
        return f.read()


def load_master(path):
    """Product master: Item No -> {hs, uom, coo, casc, brand, desc}."""
    out = collections.OrderedDict()
    if not path or not os.path.exists(path):
        return out
    rows = read_rows(path)
    hi, cols = find_header(rows, ['Item No'])
    def g(r, name):
        c = cols.get(name.lower())
        if c is None or c >= len(r) or r[c] is None:
            return ''
        return str(r[c]).strip()
    for r in rows[hi + 1:]:
        it = g(r, 'Item No')
        if it.endswith('.0'):
            it = it[:-2]
        if not it:
            continue
        hs = g(r, 'HS Code')
        out[it] = dict(desc=g(r, 'Description'), hs=h8(hs)[:8] if hs else '', uom=g(r, 'HS UOM').upper(),
                       coo=g(r, 'COO').upper(), casc=g(r, 'CASC Product Code 1'), brand=g(r, 'Brand'),
                       remarks=g(r, 'Remarks'))
    return out


def save_master(path, master):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'MASTER'
    ws.append(MASTER_COLS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E78')
    for it, d in master.items():
        ws.append([it, d.get('desc', ''), d.get('hs', ''), d.get('uom', ''), d.get('coo', ''), d.get('casc', ''),
                   d.get('brand', ''), d.get('remarks', '')])
    for col, w in zip('ABCDEFGH', (12, 45, 12, 9, 7, 20, 14, 40)):
        ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=2):
        r[0].number_format = '@'
        r[2].number_format = '@'
    ws.freeze_panes = 'A2'
    wb.save(path)


def load_hs_list(path):
    """UPDATE HS CODE list. Accepts either
       'HS Code' + 'HS UOM'  (code -> uom), or
       'Input HS Code' + 'Matched 8-digit HS Code' (+ optional 'HS UOM')."""
    out = {}
    if not path or not os.path.exists(path):
        return out
    rows = read_rows(path)
    lower = None
    for i, r in enumerate(rows[:40]):
        names = {str(v).strip().lower(): j for j, v in enumerate(r) if v is not None and str(v).strip()}
        if any('hs' in n for n in names):
            lower = (i, names)
            if ('hs code' in names and 'hs uom' in names) or any('matched' in n for n in names):
                break
    if not lower:
        raise RuntimeError('HS list: header row with "HS Code" not found')
    hi, names = lower
    c_in = next((j for n, j in names.items() if 'input' in n and 'hs' in n), names.get('hs code'))
    c_mt = next((j for n, j in names.items() if 'matched' in n), None)
    c_uom = next((j for n, j in names.items() if 'uom' in n or n in ('unit', 'hs unit')), None)
    c_desc = next((j for n, j in names.items() if 'desc' in n), None)
    if c_in is None:
        raise RuntimeError('HS list: no "HS Code" / "Input HS Code" column')
    for r in rows[hi + 1:]:
        v = r[c_in] if c_in < len(r) else None
        if v in (None, ''):
            continue
        k = h8(v)
        mt = h8(r[c_mt]) if c_mt is not None and c_mt < len(r) and r[c_mt] not in (None, '') else k
        uom = str(r[c_uom]).strip().upper() if c_uom is not None and c_uom < len(r) and r[c_uom] else ''
        desc = str(r[c_desc]).strip() if c_desc is not None and c_desc < len(r) and r[c_desc] else ''
        for key in {k, k[:8]}:
            out.setdefault(key, (mt[:8] if len(mt) >= 8 else mt, uom, desc))
    return out


# ================================================================ shipments
def build_shipments(docs):
    """Group parsed documents into shipments.
    Packing lists by Delivery #. Invoices join by Delivery #, else by order number
    (invoice "SAP order No from Supplying site" / SAP Order = packing list Order No).
    Bills of lading join by their DELIVERY ORDER NO. / PO# list and give the HAWB."""
    ships = collections.OrderedDict()
    issues = []

    def new_ship(key, delivery):
        return ships.setdefault(key, dict(key=key, delivery=delivery, invs=[], pls=[], bls=[], hawb='', hawb_src='',
                                          file_name='', dest='', gw_override=None, overrides={}, items=[], issues=[]))

    def name_check(d, num):
        nums = [n for n in re.findall(r'\d{8,12}', d['file'])
                if n not in (d.get('sto'), d.get('sap_order'), d.get('po'), d.get('invoice_no'))]
        if num and nums and num not in nums and len(nums[-1]) == len(num):
            issues.append((d['file'], 'MISMATCH', 'File name says %s but the document Delivery # is %s (matched by document)'
                           % (nums[-1], num)))
    pls = [d for d in docs if d['kind'] == 'PL']
    invs = [d for d in docs if d['kind'] == 'INV']
    bls = [d for d in docs if d['kind'] == 'BL']
    for d in docs:
        if d['kind'] not in ('INV', 'PL', 'BL'):
            issues.append((d['file'], 'ERROR', 'Not an invoice, packing list or bill of lading - skipped'))
    for d in pls:
        if not d.get('delivery'):
            issues.append((d['file'], 'CHECK', 'Delivery number not found in this packing list'))
        new_ship(d.get('delivery') or 'FILE:' + d['file'], d.get('delivery', ''))['pls'].append(d)
        name_check(d, d.get('delivery'))
    for d in invs:
        sh = ships.get(d.get('delivery')) if d.get('delivery') else None
        if sh is None:
            keys = {k for k in (d.get('sto'), d.get('sap_order')) if k}
            sh = next((x for x in ships.values() if any(p.get('po') in keys for p in x['pls']) and not x['invs']), None) \
                or next((x for x in ships.values() if any(p.get('po') in keys for p in x['pls'])), None)
            if sh is not None:
                po = next(p['po'] for p in sh['pls'] if p.get('po') in keys)
                sh['issues'].append(('CHECK', 'Invoice %s Delivery %s differs from packing list Delivery %s - matched by '
                                              'order no %s' % (d['invoice_no'], d.get('delivery'), sh['delivery'], po)))
        if sh is None:
            if not d.get('delivery'):
                issues.append((d['file'], 'CHECK', 'Delivery number not found in this invoice'))
            sh = new_ship(d.get('delivery') or 'FILE:' + d['file'], d.get('delivery', ''))
        else:
            name_check(d, d.get('delivery'))
        sh['invs'].append(d)
    for b in bls:
        hit = [x for x in ships.values() if x['delivery'] in b['deliveries']
               or any(p.get('po') in b['pos'] for p in x['pls'])
               or any(i.get('sto') in b['pos'] or i.get('sap_order') in b['pos'] or i.get('delivery') in b['deliveries']
                      for i in x['invs'])]
        if not hit:
            issues.append((b['file'], 'CHECK', 'Bill of lading %s: none of its deliveries %s / POs %s are in the other '
                                               'documents' % (b['bl_no'], ', '.join(b['deliveries']) or '-',
                                                              ', '.join(b['pos']) or '-')))
            continue
        for x in hit:
            x['bls'].append(b)
        found = {x['delivery'] for x in hit}
        miss = [n for n in b['deliveries'] if n not in found and not any(
            n == i.get('delivery') for x in hit for i in x['invs'])]
        if miss:
            issues.append((b['file'], 'MISSING', 'Bill of lading %s lists delivery %s but no invoice / packing list for it '
                                                 'was added' % (b['bl_no'], ', '.join(miss))))
    for sh in ships.values():
        for d in sh['bls'] + sh['invs'] + sh['pls']:
            if d.get('hawb') and not sh['hawb']:
                sh['hawb'], sh['hawb_src'] = d['hawb'], d['file']
        hawbs = {d['hawb'] for d in sh['bls'] + sh['invs'] + sh['pls'] if d.get('hawb')}
        if len(hawbs) > 1:
            sh['issues'].append(('MISMATCH', 'Different HAWB numbers in the documents: %s' % ', '.join(sorted(hawbs))))
        sh['dest'] = next((d['dest'] for d in sh['invs'] + sh['pls'] if d.get('dest')), '')
        for f in ('exporter', 'exporter_uen'):
            sh[f] = next((d[f] for d in sh['invs'] + sh['pls'] if d.get(f)), '')
        src = next((d for d in sh['pls'] + sh['invs'] if d.get('consignee')), {})
        for f in ('consignee', 'consignee_addr1', 'consignee_addr2'):
            sh[f] = src.get(f, '')
        uens = {d['exporter_uen'] for d in sh['invs'] if d.get('exporter_uen')}
        if len(uens) > 1:
            sh['issues'].append(('MISMATCH', 'Exporter UEN differs between the invoices: %s' % ', '.join(sorted(uens))))
        other = next((d for d in sh['pls'] if d.get('exporter_uen') and d['exporter_uen'] != sh['exporter_uen']), None)
        if sh['invs'] and other:
            sh['issues'].append(('CHECK', 'Packing list "ship on behalf of" is %s (%s) - invoice seller %s (%s) used as '
                                          'Exporter' % (other['exporter'], other['exporter_uen'], sh['exporter'],
                                                        sh['exporter_uen'])))
        if not sh['exporter_uen']:
            sh['issues'].append(('MISSING', 'Exporter UEN not found in the documents - fill it on tab 2'))
        dests = {d['dest'] for d in sh['invs'] + sh['pls'] if d.get('dest')}
        if len(dests) > 1:
            sh['issues'].append(('MISMATCH', 'Ship-to country differs between documents: %s' % ', '.join(sorted(dests))))
    return ships, issues


def ship_label(sh):
    inv = ', '.join(i['invoice_no'] for i in sh['invs'] if i.get('invoice_no'))
    return sh['hawb'] or ('INV ' + inv if inv else 'DEL ' + sh['delivery'])


def resolve(sh, ctx):
    """(Re)build the output items of one shipment.
    ctx: master{}, hs_list{}, countries, default_uom, blank_hs_not_in_list, desc_template, desc_max, brand"""
    master, hsl, C = ctx.get('master', {}), ctx.get('hs_list', {}), ctx.get('countries') or Countries()
    em = ctx.get('email') or {}
    issues = list(sh['issues'])
    if em.get('exporter') or em.get('uen'):
        issues = [i for i in issues if 'used as Exporter' not in i[1] and 'Exporter UEN not found' not in i[1]]
        issues.append(('CHECK', 'Exporter from e-mail: %s (UEN %s)' % (em.get('exporter', '-'), em.get('uen', '-'))))
        if em.get('uen') and sh.get('exporter_uen') and em['uen'] != sh['exporter_uen']:
            issues.append(('MISMATCH', 'Exporter UEN: e-mail %s, invoice %s (%s)' % (em['uen'], sh['exporter_uen'],
                                                                                    sh.get('exporter', ''))))
    pl_mats = collections.OrderedDict()
    pl_gw = pl_nw = 0.0
    pallets = 0
    have_pl_tot = False
    for pl in sh['pls']:
        for k, m in pl['materials'].items():
            d = pl_mats.setdefault(k, dict(m, qty=0.0, nw=0.0, gw=0.0))
            d['qty'] += m['qty']
            d['nw'] += m['nw']
            d['gw'] += m['gw']
        if pl['total_gw'] is not None:
            pl_gw += pl['total_gw']
            have_pl_tot = True
        if pl['total_nw'] is not None:
            pl_nw += pl['total_nw']
        pallets += pl['pallets'] or 0
        for n in pl['notes']:
            issues.append(('CHECK', '%s: %s' % (pl['file'], n)))
        hu_gw, hu_nw = sum(h['gw'] for h in pl['hus']), sum(h['nw'] for h in pl['hus'])
        if pl['total_gw'] is not None and pl['hus'] and abs(hu_gw - pl['total_gw']) > 0.01:
            issues.append(('MISMATCH', '%s: total Gross Weight %s KG but pallets add up to %s KG'
                           % (pl['file'], fmt_thou(pl['total_gw']), fmt_thou(hu_gw))))
        if pl['total_nw'] is not None and pl['hus'] and abs(hu_nw - pl['total_nw']) > 0.01:
            issues.append(('MISMATCH', '%s: total Net Weight %s KG but pallets add up to %s KG'
                           % (pl['file'], fmt_thou(pl['total_nw']), fmt_thou(hu_nw))))
    sh['total_gw'] = sh['gw_override'] if sh.get('gw_override') else (pl_gw if have_pl_tot else None)
    sh['total_nw_doc'] = pl_nw or None
    sh['pallets'] = pallets or None
    sh['container'] = next((p['container'] for p in sh['pls'] if p['container']), '')
    sh['seal'] = next((p['seal'] for p in sh['pls'] if p['seal']), '')
    sh['cont_size'] = next((p['cont_size'] for p in sh['pls'] if p['cont_size']), '')
    conts = collections.OrderedDict()
    for b in sh.get('bls', []):
        for c in b['containers']:
            size = re.match(r'(20|40|45)', c['type'] or '')
            conts[c['no']] = dict(no=c['no'], seal=c['seal'], size=size.group(1) if size else sh['cont_size'],
                                  weight=c['weight'])
    if sh['container'] and sh['container'] not in conts:
        conts[sh['container']] = dict(no=sh['container'], seal=sh['seal'], size=sh['cont_size'], weight=None)
    for c in conts.values():                      # weight in tonnes (TNE) for the Container sheet
        c['tne'] = tne_round(c['weight'] / 1000.0) if c.get('weight') else None   # round figure (5.5 and above -> 6)
    if sh.get('containers_edit') is not None:     # containers entered / changed in the tool win
        conts = collections.OrderedDict((c['no'], dict(c)) for c in sh['containers_edit'] if c.get('no'))
    sh['containers'] = list(conts.values())
    for c in conts.values():
        if not c.get('size'):
            issues.append(('MISSING', 'Container %s: size (20/40/45 FT) not in the documents - Size/Type left blank'
                           % c['no']))
    if not conts and sh['cont_size']:
        issues.append(('MISSING', 'Packing list says %s FT container but no container number is in the documents - '
                                  'Container sheet left empty (add the B/L or type it in the Excel)' % sh['cont_size']))
    if conts and not sh['container']:
        sh['container'] = next(iter(conts))
    b0 = sh['bls'][0] if sh.get('bls') else None
    sh['bl'] = b0
    if b0 and b0.get('hs'):
        for i in sh['invs']:
            for l in i['lines']:
                if l['hs'] and l['hs'] != b0['hs']:
                    issues.append(('MISMATCH', 'Item %s HS %s on invoice, %s on bill of lading' % (l['item'], l['hs'], b0['hs'])))
    sh['currency'] = next((i['currency'] for i in sh['invs'] if i['currency']), '')
    sh['incoterm'] = next((i['incoterm'] for i in sh['invs'] if i['incoterm']), '')
    curs = {i['currency'] for i in sh['invs'] if i['currency']}
    if len(curs) > 1:
        issues.append(('MISMATCH', 'More than one invoice currency: %s' % ', '.join(sorted(curs))))
    if not sh['invs']:
        issues.append(('MISSING', 'No invoice for Delivery %s - item values left blank' % sh['delivery']))
    if not sh['pls']:
        issues.append(('MISSING', 'No packing list for Delivery %s' % sh['delivery']))

    # ---- raw lines
    raw = []
    for inv in sh['invs']:
        doc_coo = C.code(inv['coo_text']) if inv['coo_text'] else ''
        if inv['coo_text'] and not doc_coo:
            issues.append(('CHECK', '%s: Country of Origin "%s" not in country list' % (inv['file'], inv['coo_text'])))
        line_sum = round(sum(l['amount'] or 0 for l in inv['lines']), 2)
        if not inv['lines']:
            issues.append(('ERROR', '%s: no item lines could be read' % inv['file']))
        if inv['total'] is not None and abs(line_sum - inv['total']) > 0.01:
            issues.append(('MISMATCH', '%s: invoice total %s but item lines add up to %s'
                           % (inv['file'], '{:,.2f}'.format(inv['total']), '{:,.2f}'.format(line_sum))))
        for l in inv['lines']:
            if l['price'] and l['qty'] and l['amount'] and abs(l['price'] * l['qty'] - l['amount']) > max(1.0, l['amount'] * 0.001):
                issues.append(('CHECK', '%s item %s: qty %s x unit price %s = %s, invoice shows %s'
                               % (inv['file'], l['item'], fmt_plain(l['qty']), fmt_plain(l['price']),
                                  '{:,.2f}'.format(l['price'] * l['qty']), '{:,.2f}'.format(l['amount']))))
            raw.append(dict(item=l['item'], inv_desc=l['desc'], qty=l['qty'], uom=l['uom'], amount=l['amount'],
                            price=l['price'], invoice=inv['invoice_no'], currency=inv['currency'],
                            doc_hs=l['hs'] or inv['hs'], doc_coo=doc_coo, batches=[b for b in l['batches'] if b],
                            src=inv['file']))
    if not raw:
        for k, m in pl_mats.items():
            raw.append(dict(item=k, inv_desc='', qty=m['qty'], uom=m['uom'], amount=None, price=None, invoice='',
                            currency='', doc_hs='', doc_coo='', batches=[], src='packing list'))
    pl_hs = next((p['hs'] for p in sh['pls'] if p['hs']), '')
    pl_coo_txt = next((p['coo_text'] for p in sh['pls'] if p['coo_text']), '')
    pl_coo = C.code(pl_coo_txt) if pl_coo_txt else ''

    # ---- weights per line (packing list material -> split by qty between lines of same item)
    by_item = collections.defaultdict(list)
    for r in raw:
        by_item[r['item']].append(r)
    for item, rs in by_item.items():
        m = pl_mats.get(item)
        tq = sum(r['qty'] or 0 for r in rs)
        if m:
            for r in rs:
                share = (r['qty'] or 0) / tq if tq else 1.0 / len(rs)
                r['nw'], r['gw'], r['nw_src'] = m['nw'] * share, m['gw'] * share, 'Packing list'
                r['pl_desc'] = m['desc']
            if len(rs) > 1:
                issues.append(('CHECK', 'Item %s is on %d invoice lines - packing list weight split by quantity' % (item, len(rs))))
            u_inv = (rs[0]['uom'] or '').upper()
            u_pl = (m['uom'] or '').upper()
            same_u = u_inv == u_pl or (u_inv in KG_UOMS and u_pl in KG_UOMS)
            if same_u and abs(tq - m['qty']) > 0.0005:
                issues.append(('MISMATCH', 'Item %s quantity: invoice %s %s, packing list %s %s, difference %s'
                               % (item, fmt_plain(tq), u_inv, fmt_plain(m['qty']), u_pl, fmt_plain(round(tq - m['qty'], 6)))))
            elif not same_u and m['uom']:
                issues.append(('CHECK', 'Item %s UOM: invoice %s vs packing list %s - quantity not compared' % (item, u_inv, u_pl)))
            if rs[0]['inv_desc'] and squash(rs[0]['inv_desc']) != squash(m['desc']):
                if squash(rs[0]['inv_desc']).startswith(squash(m['desc'])):     # packing list cut the description
                    for r in rs:
                        r['pl_desc'] = r['inv_desc']
                else:
                    issues.append(('CHECK', 'Item %s description differs: invoice "%s" / packing list "%s" (packing list used)'
                                   % (item, rs[0]['inv_desc'], m['desc'])))
        else:
            for r in rs:
                r['nw'] = r['gw'] = None
                r['nw_src'] = ''
            if sh['pls']:
                issues.append(('MISMATCH', 'Item %s is on the invoice but not on the packing list' % item))
    for k in pl_mats:
        if k not in by_item and raw and raw[0]['src'] != 'packing list':
            issues.append(('MISMATCH', 'Material %s is on the packing list but not on the invoice' % k))
    # 85% fallback for lines without net weight
    miss = [r for r in raw if r['nw'] is None]
    if miss:
        known_nw = sum(r['nw'] for r in raw if r['nw'] is not None)
        known_gw = sum(r['gw'] or 0 for r in raw if r['nw'] is not None)
        gw = sh['total_gw']
        if gw:
            tot_nw = gw * 0.85
            rest_nw = tot_nw - known_nw if known_nw else tot_nw
            rest_gw = gw - known_gw if known_gw else gw
            tv = sum(r['amount'] or 0 for r in miss)
            if tv > 0 and rest_nw > 0:
                for r in miss:
                    sh_ = (r['amount'] or 0) / tv
                    r['nw'], r['gw'], r['nw_src'] = rest_nw * sh_, rest_gw * sh_, '85% of GW, split by value'
                issues.append(('NET WEIGHT', 'Net weight not available for %s - calculated %s x 85%% = %s KG, '
                               'allocated by item value' % (', '.join(r['item'] for r in miss), fmt_thou(gw),
                                                             fmt_thou(round(rest_nw, 3)))))
            else:
                issues.append(('NET WEIGHT', 'Net Weight unavailable and Item Value unavailable for weight allocation (%s)'
                               % ', '.join(r['item'] for r in miss)))
        else:
            issues.append(('NET WEIGHT', 'Net weight and gross weight not available for %s - weight left blank'
                           % ', '.join(r['item'] for r in miss)))

    # ---- per line: master / HS / COO / UOM / HS qty / description
    items = []
    tmpl = ctx.get('desc_template') or DEFAULT_DESC_TEMPLATE
    dmax = int(ctx.get('desc_max') or 512)
    for n, r in enumerate(raw, 1):
        ov = sh['overrides'].get('%s|%s' % (r['invoice'], r['item']), {})
        ms = master.get(r['item'], {})
        it = dict(r)
        it['desc'] = r.get('pl_desc') or r['inv_desc'] or ms.get('desc', '')
        # HS code
        hs, hs_src = '', ''
        if r['doc_hs']:
            hs, hs_src = r['doc_hs'], 'Invoice'
        elif pl_hs:
            hs, hs_src = pl_hs, 'Packing list'
        elif ms.get('hs'):
            hs, hs_src = ms['hs'], 'Product master'
        if ms.get('hs') and hs and hs_src != 'Product master' and ms['hs'] != hs:
            issues.append(('MISMATCH', 'Item %s HS code: document %s, product master %s (document used)' % (r['item'], hs, ms['hs'])))
        if 'hs' in ov:                     # typed HS first: HS UOM, HS qty and CASC follow the new code
            hs, hs_src = (h8(ov['hs'])[:8] if ov['hs'] else ''), 'Edited'
        uom, uom_src = '', ''
        if hs and hsl:
            hit = hsl.get(hs) or hsl.get(hs[:8])
            if hit:
                if hit[0] and hit[0] != hs:
                    issues.append(('CHECK', 'Item %s HS %s mapped to %s by the HS list' % (r['item'], hs, hit[0])))
                    hs = hit[0]
                    hs_src += ' > HS list'
                if hit[1]:
                    uom, uom_src = hit[1], 'HS list'
            else:
                issues.append(('CHECK', 'Item %s HS %s not found in the HS list%s' % (
                    r['item'], hs, ' - left blank' if ctx.get('blank_hs_not_in_list') else '')))
                if ctx.get('blank_hs_not_in_list'):
                    hs, hs_src = '', ''
        ehs = em.get('hs', '')
        if ehs:
            if hs and not hs.startswith(ehs[:6]):
                issues.append(('MISMATCH', 'Item %s HS %s (%s) does not match e-mail HS CODE %s' % (r['item'], hs, hs_src, ehs)))
            elif not hs and len(ehs) >= 8:
                hs, hs_src = ehs[:8], 'E-mail'
            elif not hs:
                issues.append(('MISSING', 'Item %s: e-mail HS CODE %s is not 8 digits - HS code left blank' % (r['item'], ehs)))
        if 'hs' in ov:
            hs, hs_src = h8(ov['hs'])[:8] if ov['hs'] else '', 'Edited'
        if hs and not uom and hsl:          # HS from the e-mail / booking form (or typed): its UOM from the HS list
            hit = hsl.get(hs) or hsl.get(hs[:8])
            if hit and hit[1]:
                uom, uom_src = hit[1], 'HS list'
        if hs and (not hs.isdigit() or len(hs) != 8):
            issues.append(('CHECK', 'Item %s HS code "%s" is not 8 digits' % (r['item'], hs)))
        if not hs:
            issues.append(('MISSING', 'HS Code missing for item %s (%s)' % (r['item'], it['desc'])))
        if not uom and ms.get('uom'):
            uom, uom_src = ms['uom'], 'Product master'
        if not uom and hs and ctx.get('default_uom'):
            uom, uom_src = ctx['default_uom'].upper(), 'Default'
        if 'uom' in ov:
            uom, uom_src = ov['uom'].upper(), 'Edited'
        if uom_src == 'Default':
            issues.append(('CHECK', 'Item %s HS UOM not in HS list / product master - default %s used' % (r['item'], uom)))
        # COO
        coo, coo_src = '', ''
        if r['doc_coo']:
            coo, coo_src = r['doc_coo'], 'Invoice'
        elif pl_coo:
            coo, coo_src = pl_coo, 'Packing list'
        elif ms.get('coo'):
            coo, coo_src = ms['coo'], 'Product master'
        if ms.get('coo') and coo and coo_src != 'Product master' and ms['coo'] != coo:
            issues.append(('MISMATCH', 'Item %s COO: document %s, product master %s (document used)' % (r['item'], coo, ms['coo'])))
        if em.get('coo'):
            if coo and coo != em['coo']:
                issues.append(('MISMATCH', 'Item %s COO %s (%s) but e-mail says COO %s' % (r['item'], coo, coo_src, em['coo'])))
            elif not coo:
                coo, coo_src = em['coo'], 'E-mail'
        if 'coo' in ov:
            coo, coo_src = ov['coo'].upper(), 'Edited'
        if not coo:
            issues.append(('MISSING', 'COO missing for item %s (%s)' % (r['item'], it['desc'])))
        casc = ov.get('casc', ms.get('casc', ''))
        brand = ov.get('brand', ms.get('brand', '') or ctx.get('brand', ''))
        if 'nw' in ov and ov['nw'] not in (None, ''):
            it['nw'], it['nw_src'] = fnum(ov['nw']), 'Edited'
        if 'desc' in ov and ov['desc']:
            it['desc'] = ov['desc']
        # HS qty
        hq, hq_note = None, ''
        rule = UOM_RULES.get(uom)
        if rule:
            basis, calc, f = rule
            if basis == 'NET WEIGHT':
                base = it['nw']
                if base is None:
                    hq_note = 'no net weight'
            else:
                base = r['qty'] if (r['uom'] or '').upper() in COUNT_UOMS else None
                if base is None:
                    hq_note = 'invoice UOM %s is not a piece count' % r['uom']
            if base is not None:
                hq = base * f if calc == 'x' else base / f
        elif uom:
            hq_note = 'no HS Qty rule for %s' % uom
        if hq is None and uom:
            issues.append(('CHECK', 'Item %s HS Qty left blank (%s %s): %s' % (r['item'], uom, rule_text(uom), hq_note)))
        it.update(hs=hs, hs_src=hs_src, hs_uom=uom, uom_src=uom_src, coo=coo, coo_src=coo_src, casc=casc,
                  brand=brand, hs_qty=round(hq, 4) if hq is not None else None, n=n)
        # description
        vals = dict(ITEM=r['item'], DESC=it['desc'], QTY=fmt_plain(r['qty']), UOM=r['uom'] or '',
                    NW=fmt_thou(it['nw']), GW=fmt_thou(it['gw']), HS=hs, COO=coo,
                    BATCH=', '.join(r['batches']), INVOICE=r['invoice'])
        it['description'] = render_desc(tmpl, vals)
        if len(it['description']) > dmax:
            issues.append(('CHECK', 'Item %s description %d characters - cut to %d' % (r['item'], len(it['description']), dmax)))
            it['description'] = it['description'][:dmax].rstrip(' ,')
        items.append(it)
    # totals check
    nw_items = [i['nw'] for i in items if i['nw'] is not None]
    if sh['total_nw_doc'] and nw_items and all(i['nw_src'] == 'Packing list' for i in items):
        if abs(sum(nw_items) - sh['total_nw_doc']) > 0.01:
            issues.append(('MISMATCH', 'Item net weights add up to %s KG, packing list total %s KG'
                           % (fmt_thou(sum(nw_items)), fmt_thou(sh['total_nw_doc']))))
    sh['items'] = items
    sh['issues_all'] = issues
    sh['invoice_total'] = sum(i['total'] if i['total'] is not None else sum(l['amount'] or 0 for l in i['lines'])
                              for i in sh['invs']) if sh['invs'] else None
    if not sh['hawb']:
        issues.append(('MISSING', 'HAWB not found in the documents - fill it on tab 2'))
    return sh


def render_desc(tmpl, vals):
    """Fill the description template. A ' , ' segment whose fields are all empty is dropped."""
    out = []
    for seg in re.split(r'\s,\s', tmpl):
        keys = re.findall(r'\{(\w+)\}', seg)
        if keys and all(not str(vals.get(k.upper(), '')).strip() for k in keys):
            continue
        try:
            out.append(seg.format(**{k: vals.get(k.upper(), '') for k in keys}))
        except (KeyError, IndexError, ValueError):
            out.append(seg)
    return clean(' , '.join(out))


DEFAULT_REMARKS = 'job no : {JOB}\nINV : ( {INV} )'


def render_remarks(tmpl, vals):
    """Fill the Trader's Remarks template. A line whose fields are all empty is left out
    (e.g. no Job No -> no "job no :" line)."""
    if not tmpl:
        return ''
    try:
        out = []
        for line in tmpl.split('\n'):
            keys = re.findall(r'\{(\w+)\}', line)
            if keys and all(not str(vals.get(k, '')).strip() for k in keys):
                continue
            out.append(line.format(**vals))
        return '\n'.join(out).strip()
    except (KeyError, IndexError, ValueError) as e:
        raise RuntimeError("Trader's Remarks template error: %s (use only %s)" % (
            e, ' '.join('{%s}' % v for v in REMARK_VARS)))


# ================================================================ excel writer
def _put(sh, H, header, row, value, style=None, log=None):
    col = H.get(header.lower())
    if not col:
        if log:
            log('CHECK template has no "%s" column - not filled' % header)
        return
    sh.set('%s%d' % (col, row), value, style)


EXPORT_HEADER_CHECK = ['Exporter UEN', 'Exporter', 'Declaration Type', 'Cargo Pack Type', 'Outward Transport',
                       'Departure Date', 'Out Transport Identifier', 'Out Flight Number', 'Outward MAWB',
                       'Outward Carrier Agent Name', 'Discharge Port Code',
                       'Final Destination Country Code', 'Release Location Code', 'Receipt Location Code',
                       'Total Gross Weight', 'Total Outer Pack Qty']


def _date(text, tpl):
    """dd/mm/yyyy (or dd.mm.yyyy / yyyy-mm-dd) -> Excel date serial + date style; else the text as typed."""
    t = str(text).strip()
    for fmt in ('%d/%m/%Y', '%d.%m.%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y'):
        try:
            d = datetime.datetime.strptime(t, fmt).date()
            return (d - datetime.date(1899, 12, 30)).days, tpl.SDATE
        except ValueError:
            continue
    return t, None


def write_excel(path, ships, cfg, lists, log=print):
    """One workbook (one HAWB) from one or more resolved shipments."""
    tpl = Template(cfg['template'])
    changed = {}
    exp = True                      # EXPORT permits only - no inward / import fields are written
    permit = (cfg.get('permit_type') or '').upper()
    items = [it for s in ships for it in s['items']]
    hawb = ships[0]['hawb']
    # ---------------- Header
    shd = tpl.sheet('Header')
    shd.clear_from(2)
    H = tpl.headers(shd)
    hd = dict(cfg.get('header', {}))
    em = cfg.get('email') or {}
    if em.get('permit_type') in PERMIT_SCHEME and not hd.get('declaration_type'):
        permit = em['permit_type']
    if em.get('carrier') and not hd.get('carrier_agent'):
        hd['carrier_agent'] = em['carrier']
    if em.get('booking_ref') and not hd.get('mawb'):
        hd['mawb'] = em['booking_ref']
    locs = dict(LOCATION_NAMES)
    locs.update({k.upper(): v for k, v in (cfg.get('location_names') or {}).items()})
    for side, key in (('receipt', 'receipt'), ('release', 'release')):
        v = em.get(key, '')
        if not v or hd.get(side + '_code'):
            continue
        code = v.strip().upper()
        if code in locs:
            hd[side + '_code'], hd[side + '_name'] = code, hd.get(side + '_name') or locs[code]
        elif re.fullmatch(r'[A-Z]{1,4}', code):
            hd[side + '_code'] = code
        else:                             # an address -> "O" (others) + the address as the name
            hd[side + '_code'], hd[side + '_name'] = cfg.get('others_code', 'O'), v.upper()
    if (em.get('shipment_type') or '').upper().startswith('FCL') and not hd.get('cargo_pack_type'):
        hd['cargo_pack_type'] = 'Containerized'
    decl_text = hd.get('declaration_type', '')
    decl_code = lookup(lists.get('List-DeclarationType', []), decl_text) if decl_text else ''
    if not decl_text and permit in PERMIT_SCHEME:
        decl_code = PERMIT_SCHEME[permit]
        decl_text = next((t for t, c in lists.get('List-DeclarationType', []) if c == decl_code), '')
    has_cont = any(s['container'] or s['cont_size'] for s in ships)
    cargo = hd.get('cargo_pack_type') or ('Containerized' if has_cont and cfg.get('auto_cargo', True) else '')
    tmode = hd.get('transport') or ''
    if not tmode and cfg.get('auto_transport', True) and has_cont:
        tmode = 'Sea'
    put = lambda h, v, st=None: _put(shd, H, h, 2, v, st)
    s0 = ships[0]
    bl = next((s['bl'] for s in ships if s.get('bl')), None)
    if bl:                                  # bill of lading values take priority over tab 3
        hd = dict(hd)
        pc = cfg.get('port_codes') or {}
        for key, val in (('transport_id', bl['vessel']), ('flight', bl['voyage'])):
            if val:
                hd[key] = val
        for side, val in (('discharge', bl['pod']),):     # loading port is never auto-filled (not needed for OD)
            if val:
                hd[side + '_port_name'] = clean(val.split(',')[0]).upper()
                code = port_code(val, pc)
                hd[side + '_port_code'] = code
        if bl['forwarder'] and not hd.get('forwarder'):
            hd['forwarder'] = bl['forwarder']
        if bl['containers'] and not tmode:
            tmode = 'Sea'
    ent = lambda f: s0.get(f) if s0.get(f + '_src') == 'Entered' else ''
    put('Exporter UEN', ent('exporter_uen') or em.get('uen') or s0.get('exporter_uen', ''))
    put('Exporter', ent('exporter') or em.get('exporter') or s0.get('exporter', ''))
    put('Outward MAWB', hd.get('mawb', ''))
    put('Cargo Pack Type ' if 'cargo pack type ' in H else 'Cargo Pack Type', cargo)
    put('Cargo Pack Type Code', lookup(lists.get('List-CargoPacktype', []), cargo))
    put('Declaration Type', decl_text)
    put('Declaration Type Code', decl_code)
    put('Outward Carrier Agent Name', hd.get('carrier_agent', ''))
    put('Outward Carrier Agent Uen', hd.get('carrier_agent_uen', ''))
    put('Freight Forwarder', hd.get('forwarder', ''))
    put('Freight Forwarder Uen', hd.get('forwarder_uen', ''))
    if cfg.get('fill_consignee'):
        put('Consignee Name', s0.get('consignee', ''))
        put('Consignee Address 1', s0.get('consignee_addr1', ''))
        put('Consignee Address 2', s0.get('consignee_addr2', ''))
        put('Consignee Address Country Code', s0.get('dest', ''))
    put('Outward Transport', tmode)
    put('Outward Transport Mode', lookup(lists.get('List-TransportMode', []), tmode))
    if hd.get('date'):
        put('Departure Date', *_date(hd['date'], tpl))
    put('Out Flight Number', hd.get('flight', ''))
    put('Out Transport Identifier', hd.get('transport_id', ''))
    if tmode == 'Sea' and hd.get('vessel_type_list'):
        put('Vessel Type List', hd['vessel_type_list'])
        put('Vessel Type', lookup(lists.get('List-VesselType', []), hd['vessel_type_list']))
    for k, h in (('loading_port_code', 'Loading Port Code'), ('loading_port_name', 'Loading Port Name'),
                 ('discharge_port_code', 'Discharge Port Code'), ('discharge_port_name', 'Discharge Port Name')):
        put(h, hd.get(k, ''))
    put('Final Destination Country Code', s0.get('dest', ''))
    for k, h in (('release_code', 'Release Location Code'), ('release_name', 'Release Location Name'),
                 ('receipt_code', 'Receipt Location Code'), ('receipt_name', 'Receipt Location Name')):
        put(h, hd.get(k, ''))
    gws = [s['total_gw'] for s in ships if s.get('total_gw')]
    if gws:
        if tmode == 'Sea':                # ocean permit: gross weight in TNE (KG / 1000)
            put('Total Gross Weight', round(sum(gws) / 1000.0, 3))
            put('Total Gross Weight Uom', 'TNE')
        else:
            put('Total Gross Weight', round(sum(gws), 3))
            put('Total Gross Weight Uom', 'KGM')
    pal = sum(s.get('pallets') or 0 for s in ships)
    if pal and cfg.get('fill_outer_pack', True):
        put('Total Outer Pack Qty', pal)
        put('Total Outer Pack Uom', hd.get('outer_pack_uom') or 'PAT')
    nco = list(dict.fromkeys((c.get('no') if isinstance(c, dict) else str(c)) for s in ships for c in (s.get('containers') or [])
                            if (c.get('no') if isinstance(c, dict) else c)))
    if tmode == 'Sea' and nco and cfg.get('fill_outer_pack', True):      # sea job with containers -> UNT = number of containers
        put('Total Outer Pack Qty', len(nco))
        put('Total Outer Pack Uom', 'UNT')
    invs = [i['invoice_no'] for s in ships for i in s['invs'] if i.get('invoice_no')]
    rem = render_remarks(cfg.get('remarks_template', DEFAULT_REMARKS), dict(
        INV=', '.join(dict.fromkeys(invs)), HAWB=hawb, DELIVERY=', '.join(s['delivery'] for s in ships),
        JOB=cfg.get('job_no', '') or em.get('job_no', ''), PO=', '.join(dict.fromkeys(i['sap_order'] for s in ships for i in s['invs'] if i['sap_order']))))
    if em.get('permit_remarks'):
        rem = (rem + '\n' + em['permit_remarks']).strip()
    if rem:
        put("Trader's Remarks", rem)
    changed[tpl.sheets['Header']] = shd.tostring()
    filled = {h for h in H if (lambda cx: bool(cx) and ('<v' in cx or '<is>' in cx))(shd.get('%s2' % H[h]))}
    header_blank = [h for h in EXPORT_HEADER_CHECK if h.lower() in H and h.lower() not in filled]
    # ---------------- Container
    if 'Container' in tpl.sheets:
        sc = tpl.sheet('Container')
        sc.clear_from(2)
        if cfg.get('fill_container', True):
            HC = tpl.headers(sc)
            r = 2
            seen = set()
            for s in ships:
                for c in s.get('containers', []):
                    if c['no'] in seen:
                        continue
                    seen.add(c['no'])
                    _put(sc, HC, 'Container No', r, c['no'])
                    st = c.get('size_text') or (('FCL' + c['size']) if c.get('size') else '')
                    if st:
                        _put(sc, HC, 'Size/Type', r, st)
                    cw = (cfg.get('container_weight', 'TNE') or '').upper()
                    w = c.get('tne') if cw == 'TNE' else c.get('weight') if cw == 'KG' else None
                    if w not in (None, ''):
                        _put(sc, HC, 'Weight', r, round(float(w), 3))
                    if c.get('seal'):
                        _put(sc, HC, 'Seal No', r, c['seal'])
                    r += 1
        changed[tpl.sheets['Container']] = sc.tostring()
    # ---------------- Invoice
    si = tpl.sheet('Invoice')
    si.clear_from(2)
    HI = tpl.headers(si)
    _put(si, HI, 'Invoice Number', 2, '-')
    cur = ships[0]['currency']
    total = sum(s['invoice_total'] or 0 for s in ships)
    if cur:
        _put(si, HI, 'Total Invoice Currency', 2, cur)
    if total:
        _put(si, HI, 'Total Invoice Amount', 2, round(total, 2), tpl.S2)
    if cfg.get('fill_term') and cfg.get('term_on_export'):     # MSD = export permit: Term Type not captured
        t = TERM_MAP.get(ships[0]['incoterm'], ships[0]['incoterm'])
        if t in [x for x, _ in lists.get('List-TermType', [])]:
            _put(si, HI, 'Term Type', 2, t)
    changed[tpl.sheets['Invoice']] = si.tostring()
    # ---------------- Item
    st = tpl.sheet('Item')
    st.clear_from(2)
    HT = tpl.headers(st)
    casc_rows = []
    for n, it in enumerate(items, 1):
        r = n + 1
        p = lambda h, v, s_=None: _put(st, HT, h, r, v, s_)
        p('Country of Origin', it['coo'])
        if it['hs']:
            p('HS Code', int(it['hs']) if it['hs'].isdigit() and not it['hs'].startswith('0') else it['hs'],
              None if it['hs'].isdigit() and not it['hs'].startswith('0') else tpl.STXT)
        p('Description', it['description'])
        if it['hs_qty'] is not None:
            p('HS Qty', it['hs_qty'], tpl.S4)
        if it['amount'] is not None:
            p('Total Line Amount', round(it['amount'], 2), tpl.S2)
        if it['brand']:
            p('Brand', it['brand'])
        if hawb:
            p('Out HAWB', hawb)
        p('HS UOM', it['hs_uom'])
        p('Invoice Number', it['invoice'])
        p('Item Currency', it['currency'])
        if it['casc']:
            p('CASC Product Code 1', it['casc'])
            casc_rows.append((n, it['casc']))
    changed[tpl.sheets['Item']] = st.tostring()
    # ---------------- Casc
    if 'Casc' in tpl.sheets:
        sk = tpl.sheet('Casc')
        sk.clear_from(2)
        HK = tpl.headers(sk)
        for i, (n, code) in enumerate(casc_rows, 2):
            _put(sk, HK, 'Item Number', i, n)
            _put(sk, HK, 'Product Code Seq Number', i, 1)
            _put(sk, HK, 'Code 1', i, code)
        changed[tpl.sheets['Casc']] = sk.tostring()
    tpl.save(path, changed)
    return header_blank


def safe_name(s):
    s = re.sub(r'[\\/:*?"<>|]+', '_', str(s)).strip(' ._')
    return s or 'NO_NAME'


def file_name_for(group, cfg):
    s0 = group[0]
    if s0.get('file_name'):
        return safe_name(s0['file_name'])
    inv = '_'.join(i['invoice_no'] for s in group for i in s['invs'] if i.get('invoice_no'))
    vals = dict(HAWB=s0['hawb'], INVOICE=inv, DELIVERY='_'.join(s['delivery'] for s in group), DEST=s0.get('dest', ''))
    pat = cfg.get('file_pattern') or '{HAWB}'
    try:
        name = pat.format(**vals)
    except (KeyError, IndexError, ValueError):
        name = vals['HAWB']
    if not s0['hawb'] and '{HAWB}' in pat:
        name = 'NO_HAWB_INV_' + (inv or vals['DELIVERY'])
    return safe_name(name)


def generate(ships, cfg, log=print):
    """Write one Excel per HAWB + CHECK_REPORT.xlsx + EXCEL_ALL.zip. Returns (files, report, zip)."""
    lists = template_lists(cfg['template'])
    out = cfg['out_dir']
    os.makedirs(out, exist_ok=True)
    groups = collections.OrderedDict()
    for s in ships:
        if not s.get('include', True):
            continue
        key = s['hawb'] or ('DEL:' + s['key'])
        groups.setdefault(key, []).append(s)
    files, used = [], set()
    for key, g in groups.items():
        name = file_name_for(g, cfg)
        base, k = name, 2
        while name.upper() in used:
            name = '%s_%d' % (base, k)
            k += 1
        used.add(name.upper())
        path = os.path.join(out, name + '.xlsx')
        bl = next((x['bl'] for x in g if x.get('bl')), None)
        if bl:
            chk = g[0].setdefault('issues_all', [])
            gw = sum(x.get('total_gw') or 0 for x in g)
            if bl['gw'] and gw and abs(bl['gw'] - gw) > 0.01:
                chk.append(('MISMATCH', 'Gross weight: bill of lading %s KG, packing lists %s KG'
                            % (fmt_thou(bl['gw']), fmt_thou(gw))))
            pal = sum(x.get('pallets') or 0 for x in g)
            if bl['packages'] and pal and bl['packages'] != pal:
                chk.append(('MISMATCH', 'Packages: bill of lading %s %s, packing lists %s pallets'
                            % (bl['packages'], bl['package_unit'], pal)))
            if bl['pod'] and not port_code(bl['pod'], cfg.get('port_codes')):
                chk.append(('MISSING', 'Discharge port "%s" has no code in the port table - add it on tab 4' % bl['pod']))
        blank = write_excel(path, g, cfg, lists, log)
        n_items = sum(len(s['items']) for s in g)
        log('%s.xlsx  -  %d item(s), %s %s' % (name, n_items, g[0]['currency'],
                                              '{:,.2f}'.format(sum(s['invoice_total'] or 0 for s in g))))
        files.append(dict(name=name + '.xlsx', path=path, ships=g, header_blank=blank))
        if blank:
            log('   PLEASE FILL (Header, export): ' + ', '.join(blank))
    rep = os.path.join(out, cfg.get('report_name') or 'CHECK_REPORT.xlsx')
    write_report(rep, files, cfg)
    log('%s  -  check report' % os.path.basename(rep))
    zpath = os.path.join(out, 'EXCEL_ALL.zip')
    with zipfile.ZipFile(zpath, 'w', zipfile.ZIP_DEFLATED) as z:
        for f in files:
            z.write(f['path'], f['name'])
    log('EXCEL_ALL.zip  -  %d Excel file(s)' % len(files))
    return files, rep, zpath


def write_report(path, files, cfg):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    B, F = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E78')
    fills = {'MISMATCH': 'F8CBAD', 'MISSING': 'FFE699', 'ERROR': 'FF9999', 'NET WEIGHT': 'BDD7EE', 'CHECK': 'FFF2CC'}

    def sheet(title, cols, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = B, F
            c.alignment = Alignment(wrap_text=True, vertical='center')
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[num2col(i)].width = w
        ws.freeze_panes = 'A2'
        return ws
    ws = sheet('Issues', ['Excel file', 'HAWB', 'Delivery', 'Type', 'Detail'], [30, 14, 12, 13, 120], True)
    for f in files:
        for s in f['ships']:
            seen = set()
            for lvl, msg in s.get('issues_all', []):
                if (lvl, msg) in seen:
                    continue
                seen.add((lvl, msg))
                ws.append([f['name'], s['hawb'], s['delivery'], lvl, msg])
                ws.cell(ws.max_row, 4).fill = PatternFill('solid', fgColor=fills.get(lvl, 'FFFFFF'))
    for fn, lvl, msg in cfg.get('global_issues', []):
        ws.append([fn, '', '', lvl, msg])
        ws.cell(ws.max_row, 4).fill = PatternFill('solid', fgColor=fills.get(lvl, 'FFFFFF'))
    if ws.max_row == 1:
        ws.append(['', '', '', 'OK', 'No missing data or mismatches found'])
    ws.auto_filter.ref = ws.dimensions
    ws = sheet('Header Check', ['Excel file', 'Export Header field left blank', 'Why'], [30, 34, 80])
    why = {'Exporter UEN': 'Not in the documents - enter on tab 2', 'Departure Date': 'Not in the documents - tab 3',
           'Out Transport Identifier': 'Vessel / flight not in the documents - tab 3',
           'Out Flight Number': 'Voyage / flight not in the documents - tab 3',
           'Outward MAWB': 'MAWB / OBL not in the documents - tab 3',
           'Discharge Port Code': 'Invoice shows only "Dest Port" (no name) - tab 3',
           'Declaration Type': 'Choose Permit type / Declaration Type on tab 3'}
    for f in files:
        for h in f.get('header_blank', []):
            ws.append([f['name'], h, why.get(h, 'Not in the documents - fill on tab 3 if required')])
    if ws.max_row == 1:
        ws.append(['', 'All export header fields filled'])
    ws = sheet('Summary', ['Excel file', 'HAWB', 'Exporter', 'Exporter UEN', 'Invoice(s)', 'Delivery', 'Dest', 'Currency', 'Invoice total',
                           'Item value total', 'Items', 'Gross Wt (KG)', 'Net Wt items (KG)', 'Net Wt PL (KG)',
                           'Pallets', 'Container', 'Seal', 'Incoterm', 'Source files'],
               [28, 13, 40, 13, 24, 12, 6, 8, 16, 16, 6, 13, 14, 13, 8, 13, 13, 9, 70])
    for f in files:
        for s in f['ships']:
            iv = sum(i['amount'] or 0 for i in s['items'])
            nw = sum(i['nw'] or 0 for i in s['items'])
            ws.append([f['name'], s['hawb'], s.get('exporter', ''), s.get('exporter_uen', ''),
                       ', '.join(i['invoice_no'] for i in s['invs']), s['delivery'], s.get('dest', ''),
                       s['currency'], s['invoice_total'], round(iv, 2), len(s['items']), s.get('total_gw'),
                       round(nw, 3), s.get('total_nw_doc'), s.get('pallets'), s['container'], s['seal'], s['incoterm'],
                       ', '.join(d['file'] for d in s['invs'] + s['pls'])])
            for c in (9, 10):
                ws.cell(ws.max_row, c).number_format = '#,##0.00'
            for c in (12, 13, 14):
                ws.cell(ws.max_row, c).number_format = '#,##0.000'
    ws = sheet('Items', ['Excel file', 'Row', 'Item No', 'Invoice No', 'Description (Excel)', 'Inv Qty', 'Inv UOM',
                         'Amount', 'Currency', 'Net Wt (KG)', 'NW source', 'Gross Wt (KG)', 'HS Code', 'HS source',
                         'HS UOM', 'UOM source', 'HS Qty', 'COO', 'COO source', 'CASC Product Code 1', 'Batches'],
               [26, 5, 10, 12, 70, 11, 7, 16, 8, 12, 22, 12, 10, 16, 7, 14, 12, 5, 14, 14, 40])
    Y = PatternFill('solid', fgColor='FFE699')
    for f in files:
        row = 1
        for s in f['ships']:
            for it in s['items']:
                row += 1
                ws.append([f['name'], row, it['item'], it['invoice'], it['description'], it['qty'], it['uom'],
                           it['amount'], it['currency'], round(it['nw'], 3) if it['nw'] is not None else None,
                           it['nw_src'], round(it['gw'], 3) if it['gw'] is not None else None, it['hs'], it['hs_src'],
                           it['hs_uom'], it['uom_src'], it['hs_qty'], it['coo'], it['coo_src'], it['casc'],
                           ', '.join(it['batches'])])
                ws.cell(ws.max_row, 8).number_format = '#,##0.00'
                ws.cell(ws.max_row, 17).number_format = '0.0000'
                for c, bad in ((13, not it['hs']), (18, not it['coo']), (17, it['hs_qty'] is None),
                               (15, it['uom_src'] == 'Default')):
                    if bad:
                        ws.cell(ws.max_row, c).fill = Y
    ws = sheet('UOM Rules', ['HS UOM', 'Based on', 'Calc', 'Factor', 'HS Qty ='], [10, 14, 8, 10, 30])
    for u, (basis, calc, f) in UOM_RULES.items():
        ws.append([u, basis, calc, f, rule_text(u)])
    wb.save(path)


def run_folder(pdfs, cfg, log=print):
    """Command-line helper: parse PDFs, resolve, generate."""
    lists = template_lists(cfg['template'])
    C = Countries(lists.get('List-Country', []))
    docs = [parse_pdf(p, C) for p in pdfs]
    ships, gi = build_shipments(docs)
    ctx = dict(master=load_master(cfg.get('master')), hs_list=load_hs_list(cfg.get('hs_list')), countries=C,
               default_uom=cfg.get('default_uom', 'KGM'), blank_hs_not_in_list=cfg.get('blank_hs_not_in_list'),
               desc_template=cfg.get('desc_template'), desc_max=cfg.get('desc_max', 512), brand=cfg.get('brand', ''),
               email=cfg.get('email') or {})
    for k, v in cfg.get('hawb_overrides', {}).items():
        if k in ships:
            ships[k]['hawb'], ships[k]['hawb_src'] = v, 'Entered'
    for s in ships.values():
        resolve(s, ctx)
    cfg = dict(cfg, global_issues=gi)
    return generate(list(ships.values()), cfg, log), ships, gi
