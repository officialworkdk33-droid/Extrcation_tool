# -*- coding: utf-8 -*-
"""INFINITY PERMIT SERVER - the permit tools (Danfoss / Festo / Agilent / any supplier) as one web application
hosted on YOUR computer: your PC is the server, the database is a file next to it, colleagues open it in a browser.

  PERMIT_SERVER.exe          (or: python server.py)
  -> http://localhost:8080   on this PC
  -> http://<this PC's IP>:8080   from other PCs in the office (same network)

Data (all on this PC, folder PERMIT_SERVER_DATA next to the exe):
  permit.db        SQLite database - users, jobs, files, logs, settings, HS map, audit trail
  jobs/<id>/in     uploaded PDFs        jobs/<id>/out   Excel / XML / CHECK_REPORT / RAW_DATA
  jobs/<id>/state.pkl   everything read + all edits of the job
  masters/         product masters (one per supplier tool)       hs_data.json   UPDATE HS CODE list
"""
import os, sys, io, re, json, time, queue, shutil, socket, sqlite3, zipfile, secrets, datetime, threading, traceback
import logging as _logging
_logging.getLogger('pdfminer').setLevel(_logging.ERROR)   # harmless font warnings (FontBBox ...) not shown
import functools, webbrowser

from flask import Flask, request, jsonify, session, send_file, send_from_directory, abort
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

import toolrun as T
import acra

VERSION = '1.0'
APP_NAME = 'INFINITY PERMIT SERVER'


def app_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


DATA = os.environ.get('PERMIT_SERVER_DATA') or os.path.join(app_dir(), 'PERMIT_SERVER_DATA')
P = T.Paths(DATA)
acra.STATE['cache_path'] = os.path.join(DATA, 'acra_cache.json')
DB = os.path.join(DATA, 'permit.db')
CONF = os.path.join(DATA, 'server_config.json')
WEB = os.path.join(T.base_dir(), 'web')

# ======================================================================== database
SCHEMA = """
CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL, name TEXT, pw TEXT NOT NULL,
    role TEXT NOT NULL DEFAULT 'user', active INTEGER NOT NULL DEFAULT 1, created TEXT, last_login TEXT);
CREATE TABLE IF NOT EXISTS jobs (id INTEGER PRIMARY KEY, tool TEXT NOT NULL, title TEXT, job_no TEXT, status TEXT,
    message TEXT, created_by TEXT, created TEXT, updated TEXT, form TEXT, summary TEXT, hawbs TEXT, invoices TEXT,
    value TEXT, items INTEGER DEFAULT 0, problems INTEGER DEFAULT 0, generated TEXT);
CREATE TABLE IF NOT EXISTS files (id INTEGER PRIMARY KEY, job_id INTEGER NOT NULL, kind TEXT NOT NULL, name TEXT NOT NULL,
    path TEXT NOT NULL, size INTEGER, created TEXT, by_user TEXT);
CREATE TABLE IF NOT EXISTS logs (id INTEGER PRIMARY KEY, job_id INTEGER NOT NULL, ts TEXT, line TEXT);
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS audit (id INTEGER PRIMARY KEY, ts TEXT, username TEXT, action TEXT, job_id INTEGER, detail TEXT);
CREATE INDEX IF NOT EXISTS ix_logs_job ON logs(job_id, id);
CREATE INDEX IF NOT EXISTS ix_files_job ON files(job_id);
CREATE INDEX IF NOT EXISTS ix_jobs_upd ON jobs(updated);
"""
_db_local = threading.local()
_db_write = threading.Lock()


def db():
    c = getattr(_db_local, 'c', None)
    if c is None:
        c = sqlite3.connect(DB, timeout=30, check_same_thread=False)
        c.row_factory = sqlite3.Row
        c.execute('PRAGMA journal_mode=WAL')
        c.execute('PRAGMA foreign_keys=ON')
        _db_local.c = c
    return c


def q(sql, args=(), one=False):
    cur = db().execute(sql, args)
    rows = [dict(r) for r in cur.fetchall()]
    return (rows[0] if rows else None) if one else rows


def x(sql, args=()):
    with _db_write:
        c = db()
        cur = c.execute(sql, args)
        c.commit()
        return cur.lastrowid


def now():
    return datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')


JOB_EXTRA = ('decl_type', 'party', 'mawb', 'port', 'eta', 'direction', 'permit_no', 'int_remarks', 'decl_date')


def init_db():
    with _db_write:
        db().executescript(SCHEMA)
        have = {r[1] for r in db().execute('PRAGMA table_info(jobs)')}
        for c in JOB_EXTRA:                       # columns of the job list (added to older databases too)
            if c not in have:
                db().execute('ALTER TABLE jobs ADD COLUMN %s TEXT' % c)
        db().commit()


def get_settings():
    r = q("SELECT value FROM settings WHERE key='global'", one=True)
    return T.merged_settings(json.loads(r['value']) if r else {})


def put_settings(s):
    x("INSERT INTO settings(key, value) VALUES('global', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
      (json.dumps(s),))


def audit(action, job_id=None, detail=''):
    try:
        x('INSERT INTO audit(ts, username, action, job_id, detail) VALUES(?,?,?,?,?)',
          (now(), session.get('user', 'system') if _in_request() else 'system', action, job_id, str(detail)[:2000]))
    except Exception:
        pass


def _in_request():
    try:
        return bool(request)
    except RuntimeError:
        return False


def job_log(jid, line):
    for ln in str(line).splitlines() or ['']:
        x('INSERT INTO logs(job_id, ts, line) VALUES(?,?,?)', (jid, now(), ln))


def set_job(jid, **kw):
    kw['updated'] = now()
    x('UPDATE jobs SET %s WHERE id=?' % ', '.join('%s=?' % k for k in kw), list(kw.values()) + [jid])


def job_row(jid):
    j = q('SELECT * FROM jobs WHERE id=?', (jid,), one=True)
    if not j:
        abort(404)
    j['form'] = json.loads(j['form'] or '{}')
    j['summary'] = json.loads(j['summary'] or '{}')
    return j


# ======================================================================== job states + worker
_states = {}
_states_lock = threading.Lock()
TASKS = queue.Queue()
RUNNING = {}
PROGRESS = {}


def state(jid, tool=None):
    with _states_lock:
        st = _states.get(jid)
        if st is None:
            tool = tool or job_row(jid)['tool']
            st = T.load_state(os.path.join(P.job(jid), 'state.pkl'), tool)
            _states[jid] = st
        return st


def save(jid, st, j=None):
    T.save_state(os.path.join(P.job(jid), 'state.pkl'), st)
    j = j or job_row(jid)
    E = T.engine(st.tool)
    sm = T.summary(E, st, get_settings())
    try:
        lc = T.list_cols(E, st, json.loads(j['form'] or '{}') if isinstance(j['form'], str) else (j['form'] or {}))
    except Exception:
        lc = {}
    set_job(jid, summary=json.dumps(sm), hawbs=sm['hawbs'], invoices=sm['invoices'], value=sm['value'],
            items=sm['items'], problems=sm['problems'], **lc)
    return sm


def add_file(jid, kind, path, by=''):
    return x('INSERT INTO files(job_id, kind, name, path, size, created, by_user) VALUES(?,?,?,?,?,?,?)',
             (jid, kind, os.path.basename(path), path, os.path.getsize(path), now(), by))


def worker():
    while True:
        jid, kind, user = TASKS.get()
        RUNNING[jid] = kind
        log = lambda m: job_log(jid, m)
        try:
            j = job_row(jid)
            st = state(jid, j['tool'])
            if kind == 'extract' and getattr(st, 'tool', j['tool']) != j['tool']:     # other documents -> other reader
                with _states_lock:
                    st = _states[jid] = T.State(j['tool'])
            S = get_settings()
            with st.lock:
                if kind == 'extract':
                    set_job(jid, status='extracting', message='Reading documents ...')
                    log('%s  Extract started by %s' % (now(), user))
                    ins = q("SELECT path FROM files WHERE job_id=? AND kind='input' ORDER BY id", (jid,))
                    PROGRESS[jid] = dict(stage='read', done=0, total=len(ins), file='')
                    r = T.extract(j['tool'], P, S, j['form'], [f['path'] for f in ins if os.path.exists(f['path'])], st, log,
                                  progress=lambda **k: PROGRESS.__setitem__(jid, k))
                    form = j['form']
                    if r.get('permit_type') and not form.get('permit_type'):
                        form['permit_type'] = r['permit_type']
                        form['permit_src'] = 'documents'          # declaration advisor may correct it (never a typed one)
                    if r.get('job_no') and not form.get('job_no'):
                        form['job_no'] = r['job_no']
                    if r.get('brand_mode') and 'brand_mode' not in form:     # remembered for this supplier
                        form['brand_mode'] = r['brand_mode']
                    set_job(jid, form=json.dumps(form), job_no=form.get('job_no', ''))
                    _auto_uen(jid, st, form, log)
                    sm = save(jid, st)
                    set_job(jid, status='review', message='%d shipment(s), %d item(s), %d problem(s)' % (
                        sm['ships'], sm['items'], sm['problems']))
                elif kind == 'generate':
                    set_job(jid, status='generating', message='Making Excel / XML ...')
                    log('=' * 60)
                    log('%s  Generate started by %s' % (now(), user))
                    out = os.path.join(P.job(jid), 'out')
                    x("DELETE FROM files WHERE job_id=? AND kind='output'", (jid,))
                    PROGRESS[jid] = dict(stage='excel', done=0, total=0, file='')
                    T.generate(j['tool'], P, S, j['form'], st, out, log)
                    for f in sorted(os.listdir(out)):
                        add_file(jid, 'output', os.path.join(out, f), user)
                    save(jid, st)
                    n = len(os.listdir(out))
                    set_job(jid, status='done', message='%d file(s) ready to download' % n, generated=now(),
                            decl_date=datetime.datetime.now().strftime('%Y%m%d'))
        except Exception as e:
            log('ERROR ' + traceback.format_exc())
            set_job(jid, status='error', message=str(e)[:300])
        finally:
            RUNNING.pop(jid, None)
            PROGRESS.pop(jid, None)
            TASKS.task_done()


def _auto_uen(jid, st, form, log):
    """Importer / exporter UEN not in the documents -> ACRA open data (data.gov.sg), remembered in the UEN book."""
    if not get_settings().get('acra_lookup', True):
        return
    todo = [sh for sh in st.ships.values() if sh.get('party') and not sh.get('party_uen')]
    if not todo:
        _uen_names(st, form, log)
        return
    S_ = get_settings()
    book = S_.setdefault('uen_book', {})
    found = 0
    for sh in todo:
        name = sh['party'].upper()
        if book.get(name):
            continue
        uen = acra.lookup(name, timeout=10)
        if uen:
            book[name] = uen
            found += 1
            log('UEN %s for %s from ACRA (data.gov.sg open data) - remembered' % (uen, name))
        else:
            log('UEN of %s not found automatically (ACRA open data)%s - type it on the Party tab' % (
                name, ' - ' + acra.STATE['last_error'] if acra.STATE.get('last_error') else ''))
    if found:
        put_settings(S_)
        T.resolve_all(st.tool, P, get_settings(), form, st)
    _uen_names(st, form, log)


def _uen_names(st, form, log):
    """Registered name (+ address) of every UEN of the job (importer / exporter, forwarder, carrier agent) from the
    ACRA open data - kept in Settings (uen_names) and used for the names on the permit."""
    S_ = get_settings()
    if not S_.get('acra_lookup', True) or not (S_.get('options') or {}).get('uen_names', True):
        return
    names = S_.setdefault('uen_names', {})
    hd = form.get('header') or {}
    want = {}
    for sh in st.ships.values():
        cg = sh.get('cargo') or {}
        for u, hint in ((sh.get('party_uen'), sh.get('party')), (cg.get('forwarder_uen') or hd.get('forwarder_uen'), cg.get('forwarder')),
                        (cg.get('carrier_agent_uen') or hd.get('carrier_agent_uen'), cg.get('carrier_agent'))):
            u = str(u or '').strip().upper()
            if u and u not in names:
                want.setdefault(u, hint or '')
    got = 0
    for u, hint in list(want.items())[:10]:
        hit = acra.by_uen(u, hint, timeout=10)
        if hit.get('name'):
            names[u] = dict(name=hit['name'], address=hit.get('address') or [], status=hit.get('status', ''))
            got += 1
            log('UEN %s = %s (ACRA register)' % (u, hit['name']))
    if got:
        put_settings(S_)
        T.resolve_all(st.tool, P, get_settings(), form, st)


def queue_task(jid, kind):
    set_job(jid, status='queued', message='Waiting (%d job(s) before it)' % TASKS.qsize())
    TASKS.put((jid, kind, session.get('user', '')))
    audit(kind, jid)


# ======================================================================== web app
app = Flask(__name__, static_folder=None)
app.config['MAX_CONTENT_LENGTH'] = 300 * 1024 * 1024


def _secret():
    p = os.path.join(DATA, 'secret.key')
    if not os.path.exists(p):
        with open(p, 'w') as f:
            f.write(secrets.token_hex(32))
    return open(p).read().strip()


app.secret_key = _secret()
from werkzeug.middleware.proxy_fix import ProxyFix
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)


def is_local():
    """Request typed on THIS PC (not through the internet tunnel / another PC)."""
    if request.headers.get('Cf-Connecting-Ip') or request.headers.get('X-Forwarded-For'):
        return False
    return (request.remote_addr or '') in ('127.0.0.1', '::1')


_FAILS = {}


def _throttled(key):
    t, n = _FAILS.get(key, (0, 0))
    return n >= 8 and time.time() - t < 900


def _fail(key):
    t, n = _FAILS.get(key, (0, 0))
    _FAILS[key] = (time.time(), (n if time.time() - t < 900 else 0) + 1)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
                  PERMANENT_SESSION_LIFETIME=datetime.timedelta(hours=12))


def login_required(fn):
    @functools.wraps(fn)
    def w(*a, **k):
        if not session.get('user'):
            return jsonify(error='Please log in'), 401
        u = q('SELECT active FROM users WHERE username=?', (session['user'],), one=True)
        if not u or not u['active']:
            session.clear()
            return jsonify(error='Account disabled'), 401
        return fn(*a, **k)
    return w


def admin_required(fn):
    @functools.wraps(fn)
    @login_required
    def w(*a, **k):
        if session.get('role') != 'admin':
            return jsonify(error='Only an admin can do this'), 403
        return fn(*a, **k)
    return w


@app.errorhandler(ValueError)
def _bad(e):
    return jsonify(error=str(e)), 400


@app.after_request
def _nocache(r):
    if request.path.startswith('/api/'):
        r.headers['Cache-Control'] = 'no-store'
    r.headers['X-Frame-Options'] = 'DENY'
    r.headers['X-Content-Type-Options'] = 'nosniff'
    r.headers['Referrer-Policy'] = 'same-origin'
    if request.is_secure:
        r.headers['Strict-Transport-Security'] = 'max-age=31536000'
    return r


@app.route('/')
def index():
    return send_from_directory(WEB, 'index.html')


@app.route('/web/<path:p>')
def web(p):
    return send_from_directory(WEB, p)


@app.route('/favicon.ico')
def favicon():
    if os.path.exists(os.path.join(WEB, 'favicon.png')):
        return send_from_directory(WEB, 'favicon.png')
    return send_from_directory(T.base_dir(), 'logo.png')


# ------------------------------------------------------------------ auth
@app.get('/api/me')
def me():
    if not q('SELECT id FROM users LIMIT 1', one=True):
        return jsonify(setup=True, local=is_local())
    if not session.get('user'):
        return jsonify(user=None)
    u = q('SELECT username, name, role FROM users WHERE username=? AND active=1', (session['user'],), one=True)
    return jsonify(user=u)


@app.post('/api/setup')
def setup():
    if q('SELECT id FROM users LIMIT 1', one=True):
        return jsonify(error='Already set up'), 400
    if not is_local():
        return jsonify(error='Create the first admin on the server PC itself (http://localhost)'), 403
    d = request.get_json(force=True)
    u, pw = (d.get('username') or '').strip().lower(), d.get('password') or ''
    if not re.fullmatch(r'[a-z0-9._-]{2,30}', u) or len(pw) < 6:
        raise ValueError('User name 2-30 letters / digits, password at least 6 characters')
    x('INSERT INTO users(username, name, pw, role, created) VALUES(?,?,?,?,?)',
      (u, d.get('name') or u, generate_password_hash(pw), 'admin', now()))
    session.update(user=u, role='admin')
    session.permanent = True
    audit('setup admin', detail=u)
    return jsonify(ok=True)


@app.post('/api/login')
def login():
    d = request.get_json(force=True)
    un = (d.get('username') or '').strip().lower()
    key = '%s|%s' % (request.headers.get('Cf-Connecting-Ip') or request.remote_addr, un)
    if _throttled(key):
        return jsonify(error='Too many wrong passwords - try again in 15 minutes'), 429
    u = q('SELECT * FROM users WHERE username=?', (un,), one=True)
    if not u or not u['active'] or not check_password_hash(u['pw'], d.get('password') or ''):
        _fail(key)
        audit('login failed', detail='%s from %s' % (un, request.remote_addr))
        time.sleep(0.8)
        return jsonify(error='Wrong user name or password'), 401
    _FAILS.pop(key, None)
    session.update(user=u['username'], role=u['role'])
    session.permanent = True
    x('UPDATE users SET last_login=? WHERE id=?', (now(), u['id']))
    audit('login')
    return jsonify(ok=True)


@app.post('/api/logout')
def logout():
    session.clear()
    return jsonify(ok=True)


@app.post('/api/password')
@login_required
def password():
    d = request.get_json(force=True)
    u = q('SELECT * FROM users WHERE username=?', (session['user'],), one=True)
    if not check_password_hash(u['pw'], d.get('old') or ''):
        raise ValueError('Current password is wrong')
    if len(d.get('new') or '') < 6:
        raise ValueError('New password: at least 6 characters')
    x('UPDATE users SET pw=? WHERE id=?', (generate_password_hash(d['new']), u['id']))
    audit('change password')
    return jsonify(ok=True)


@app.get('/api/users')
@admin_required
def users():
    return jsonify(q('SELECT id, username, name, role, active, created, last_login FROM users ORDER BY username'))


@app.post('/api/users')
@admin_required
def user_save():
    d = request.get_json(force=True)
    role = 'admin' if d.get('role') == 'admin' else 'user'
    if d.get('id'):
        u = q('SELECT * FROM users WHERE id=?', (d['id'],), one=True)
        if not u:
            abort(404)
        if u['username'] == session['user'] and (role != 'admin' or not d.get('active', True)):
            raise ValueError('You cannot remove your own admin rights')
        x('UPDATE users SET name=?, role=?, active=? WHERE id=?', (d.get('name') or u['name'], role,
                                                                   1 if d.get('active', True) else 0, u['id']))
        if d.get('password'):
            if len(d['password']) < 6:
                raise ValueError('Password: at least 6 characters')
            x('UPDATE users SET pw=? WHERE id=?', (generate_password_hash(d['password']), u['id']))
        audit('user update', detail=u['username'])
    else:
        un = (d.get('username') or '').strip().lower()
        if not re.fullmatch(r'[a-z0-9._-]{2,30}', un) or len(d.get('password') or '') < 6:
            raise ValueError('User name 2-30 letters / digits, password at least 6 characters')
        if q('SELECT id FROM users WHERE username=?', (un,), one=True):
            raise ValueError('User name already exists')
        x('INSERT INTO users(username, name, pw, role, created) VALUES(?,?,?,?,?)',
          (un, d.get('name') or un, generate_password_hash(d['password']), role, now()))
        audit('user add', detail=un)
    return jsonify(ok=True)


@app.get('/api/users/template')
@admin_required
def users_template():
    """Excel to add many users at once: User name | Name | Password | Role (user / admin)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'USERS'
    ws.append(['User name', 'Name', 'Password', 'Role'])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E78')
    ws.append(['john.tan', 'John Tan', 'Start#2026', 'user'])
    for col, w in zip('ABCD', (18, 28, 16, 10)):
        ws.column_dimensions[col].width = w
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name='USERS_TEMPLATE.xlsx')


@app.post('/api/users/import')
@admin_required
def users_import():
    """Add / update many users from an Excel (User name, Name, Password, Role). Existing users: name / role /
    password (when given) updated. Returns counts + rows that could not be used."""
    from openpyxl import load_workbook
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith(('.xlsx', '.xlsm')):
        raise ValueError('Choose the users Excel (.xlsx) - download the template first')
    wb = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(rows[:10]) if any(str(v or '').strip().lower() in ('user name', 'username', 'user') for v in r)), None)
    if hi is None:
        raise ValueError('No "User name" column - use the template')
    names = [str(v or '').strip().lower() for v in rows[hi]]
    col = lambda *ks: next((names.index(k) for k in ks if k in names), None)
    cu, cn, cp, cr = col('user name', 'username', 'user'), col('name', 'full name'), col('password'), col('role')
    added, updated, bad = 0, 0, []
    for n, r in enumerate(rows[hi + 1:], hi + 2):
        g = lambda c: str(r[c]).strip() if c is not None and c < len(r) and r[c] is not None else ''
        un, nm, pw, rl = g(cu).lower(), g(cn), g(cp), ('admin' if g(cr).lower() == 'admin' else 'user')
        if not un and not nm:
            continue
        if not re.fullmatch(r'[a-z0-9._-]{2,30}', un):
            bad.append('row %d: user name "%s" (2-30 letters / digits . _ -)' % (n, un))
            continue
        u = q('SELECT * FROM users WHERE username=?', (un,), one=True)
        if u:
            if u['username'] == session['user'] and rl != 'admin':
                rl = 'admin'
            x('UPDATE users SET name=?, role=?, active=1 WHERE id=?', (nm or u['name'], rl, u['id']))
            if pw:
                if len(pw) < 6:
                    bad.append('row %d: %s password shorter than 6 - not changed' % (n, un))
                else:
                    x('UPDATE users SET pw=? WHERE id=?', (generate_password_hash(pw), u['id']))
            updated += 1
        else:
            if len(pw) < 6:
                bad.append('row %d: %s needs a password of at least 6 characters' % (n, un))
                continue
            x('INSERT INTO users(username, name, pw, role, created) VALUES(?,?,?,?,?)',
              (un, nm or un, generate_password_hash(pw), rl, now()))
            added += 1
    audit('users import', detail='%d added, %d updated, %d not used' % (added, updated, len(bad)))
    return jsonify(added=added, updated=updated, bad=bad[:50])


# ------------------------------------------------------------------ meta / dashboard
@app.get('/api/meta')
@login_required
def meta():
    L = T.lists('danfoss')
    E = T.engine('danfoss')
    S = get_settings()
    hs = {}
    try:
        with open(P.hs, encoding='utf-8') as f:
            d = json.load(f)
        hs = dict(count=d.get('count'), pcodes=d.get('pcode_count'), source=d.get('source', ''),
                  updated=datetime.datetime.fromtimestamp(os.path.getmtime(P.hs)).strftime('%d/%m/%Y %H:%M'))
    except Exception:
        pass
    return jsonify(
        version=VERSION, app=APP_NAME, tools=[dict(key=k, **{a: v[a] for a in ('name', 'color', 'charges', 'hint')})
                                              for k, v in T.TOOLS.items()],
        teams=S.get('teams') or ['DGF OCEAN', 'KWE'],
        msg_types=S.get('msg_types') or ['IPTDEC'], msg_default=dict(IMPORT=(S.get('xml') or {}).get('msg_import') or 'IPTDEC',
                                                                      EXPORT=(S.get('xml') or {}).get('msg_export') or ''),
        carrier_book=S.get('carrier_book') or {},
        header_fields=T.HEADER_FIELDS, dir_only=dict(EXPORT=T.engine('invoice').EXPORT_ONLY_RE, IMPORT=T.engine('invoice').IMPORT_ONLY_RE), permits=T.PERMITS, consolidate=list(T.CONSOLIDATE.items()),
        locations=list((S.get('locations') or T.LOCATIONS).items()),
        lists={k: L.get(k, []) for k in ('List-UOM', 'List-Country', 'List-DeclarationType', 'List-TransportMode',
                                          'List-CargoPacktype', 'List-VesselType', 'List-Currency', 'List-BgIndicator')},
        columns=template_columns(),
        permit_scheme=dict(E.PERMIT_SCHEME), ocr=E.OCR.get('status', ''), hs=hs, addresses=addresses(),
        user=session.get('user'), role=session.get('role'))


_COLS = {}


def template_columns():
    """Every column of the Header / Invoice sheets of the dec-upload template (for "More permit fields")."""
    if not _COLS:
        import openpyxl
        wb = openpyxl.load_workbook(T.template_path(), read_only=True)
        for sn in ('Header', 'Invoice'):
            ws = wb[sn]
            _COLS[sn] = [str(c).strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)) if c]
        wb.close()
    return _COLS


@app.get('/api/stats')
@login_required
def stats():
    by = {r['status']: r['n'] for r in q('SELECT status, COUNT(*) n FROM jobs GROUP BY status')}
    tools = {r['tool']: r['n'] for r in q('SELECT tool, COUNT(*) n FROM jobs GROUP BY tool')}
    today = datetime.date.today()
    days = []
    for i in range(13, -1, -1):
        d = (today - datetime.timedelta(days=i)).isoformat()
        n = q("SELECT COUNT(*) n FROM jobs WHERE substr(created,1,10)=?", (d,), one=True)['n']
        g = q("SELECT COUNT(*) n FROM jobs WHERE substr(generated,1,10)=?", (d,), one=True)['n']
        days.append(dict(day=d, created=n, generated=g))
    mine = q('SELECT COUNT(*) n FROM jobs WHERE created_by=?', (session['user'],), one=True)['n']
    return jsonify(status=by, tools=tools, days=days, total=sum(by.values()), mine=mine,
                   queue=TASKS.qsize(), running=list(RUNNING.items()))


# ------------------------------------------------------------------ jobs
def job_public(j):
    j = dict(j)
    j['running'] = RUNNING.get(j['id'])
    return j


@app.get('/api/jobs')
@login_required
def jobs():
    sql = 'SELECT id, tool, title, job_no, status, message, created_by, created, updated, hawbs, invoices, value, items, ' \
          'problems, generated, %s FROM jobs WHERE 1=1' % ', '.join(JOB_EXTRA)
    a = []
    if request.args.get('tool'):
        sql += ' AND tool=?'
        a.append(request.args['tool'])
    if request.args.get('status'):
        sql += ' AND status=?'
        a.append(request.args['status'])
    if request.args.get('direction'):
        sql += ' AND direction=?'
        a.append(request.args['direction'])
    if request.args.get('mine') == '1':
        sql += ' AND created_by=?'
        a.append(session['user'])
    s = (request.args.get('q') or '').strip()
    if s:
        sql += (' AND (title LIKE ? OR job_no LIKE ? OR hawbs LIKE ? OR invoices LIKE ? OR created_by LIKE ? OR party LIKE ? '
                'OR mawb LIKE ? OR permit_no LIKE ? OR int_remarks LIKE ? OR CAST(id AS TEXT)=?)')
        a += ['%' + s + '%'] * 9 + [s]
    sql += ' ORDER BY updated DESC LIMIT ?'
    a.append(min(int(request.args.get('limit') or 3000), 20000))
    return jsonify([job_public(r) for r in q(sql, a)])


DOC_RE = re.compile(r'\.(pdf|xlsx|xls|xlsm|txt|eml|msg|png|jpg|jpeg|tif|tiff)$', re.I)


def _new_dst(jid, name):
    dst = os.path.join(P.job(jid), 'in', name)
    base, ext = os.path.splitext(dst)
    k = 1
    while os.path.exists(dst):
        dst = '%s_%d%s' % (base, k, ext)
        k += 1
    return dst


def _save_uploads(jid, files):
    """Save the uploaded documents; a .zip (e.g. a Satair shipment: one PDF per house B/L) is unpacked."""
    n = 0
    for f in files:
        if not f or not f.filename:
            continue
        name = secure_filename(f.filename) or ('file%d.pdf' % n)
        if name.lower().endswith('.zip'):
            data = io.BytesIO(f.read())
            try:
                z = zipfile.ZipFile(data)
            except zipfile.BadZipFile:
                raise ValueError('Not a valid zip file: ' + f.filename)
            got = 0
            for zi in z.infolist():
                if zi.is_dir() or '__MACOSX' in zi.filename or os.path.basename(zi.filename).startswith('.'):
                    continue
                inner = secure_filename(os.path.basename(zi.filename))
                if not inner or not DOC_RE.search(inner) or zi.file_size > 200 * 1024 * 1024:
                    continue
                dst = _new_dst(jid, inner)
                with z.open(zi) as src, open(dst, 'wb') as out:
                    shutil.copyfileobj(src, out)
                add_file(jid, 'input', dst, session.get('user', ''))
                n += 1
                got += 1
            if not got:
                raise ValueError('No PDF / Excel / e-mail files in ' + f.filename)
            continue
        if not DOC_RE.search(name):
            raise ValueError('Not a PDF / Excel / e-mail / zip file: ' + f.filename)
        dst = _new_dst(jid, name)
        f.save(dst)
        add_file(jid, 'input', dst, session.get('user', ''))
        n += 1
    return n


@app.post('/api/jobs')
@login_required
def job_new():
    tool = request.form.get('tool')
    auto = tool == 'auto'
    if auto:
        tool = 'invoice'
    if tool not in T.TOOLS:
        raise ValueError('Pick the supplier tool')
    form = json.loads(request.form.get('form') or '{}')
    if form.get('permit_type'):
        form['permit_src'] = 'you'                 # picked when the job was made
    title = (request.form.get('title') or '').strip()
    jid = x('INSERT INTO jobs(tool, title, job_no, status, message, created_by, created, updated, form) VALUES(?,?,?,?,?,?,?,?,?)',
            (tool, title, form.get('job_no', ''), 'new', '', session['user'], now(), now(), json.dumps(form)))
    n = _save_uploads(jid, request.files.getlist('files'))
    if auto and n:
        tool = T.detect_tool([f['path'] for f in q("SELECT path FROM files WHERE job_id=? AND kind='input'", (jid,))])
        set_job(jid, tool=tool)
    if not title:
        set_job(jid, title='%s job %d' % (T.TOOLS[tool]['name'], jid))
    audit('job create', jid, '%s, %d file(s)' % (tool, n))
    if n:
        queue_task(jid, 'extract')
    return jsonify(id=jid)


def _ships_list(st, S):
    E = T.engine(st.tool)
    return [T.ship_json(E, k, s, S, full=False) for k, s in st.ships.items()]


@app.get('/api/jobs/<int:jid>')
@login_required
def job_get(jid):
    j = job_row(jid)
    st = state(jid, j['tool'])
    S = get_settings()
    j['running'] = RUNNING.get(jid)
    j['ships'] = _ships_list(st, S)
    j['files'] = q('SELECT id, kind, name, size, created, by_user FROM files WHERE job_id=? ORDER BY kind, name', (jid,))
    j['undo'] = len(st.undo)
    first = next(iter(st.ships.values()), None)
    j['cargo_auto'] = {k: v for k, v in ((first or {}).get('cargo') or {}).items() if isinstance(v, (str, int, float))}
    j['gi'] = [list(g) for g in st.gi]
    return jsonify(j)


@app.get('/api/jobs/<int:jid>/ship/<path:key>')
@login_required
def ship_get(jid, key):
    j = job_row(jid)
    st = state(jid, j['tool'])
    if key not in st.ships:
        abort(404)
    return jsonify(T.ship_json(T.engine(st.tool), key, st.ships[key], get_settings()))


@app.get('/api/jobs/<int:jid>/ship/<path:key>/item/<int:i>')
@login_required
def item_get(jid, key, i):
    j = job_row(jid)
    st = state(jid, j['tool'])
    sh = st.ships.get(key)
    if not sh or i >= len(sh['items']):
        abort(404)
    E = T.engine(st.tool)
    it = sh['items'][i]
    d = T.item_json(E, sh, i, it)
    d['pl'] = T.pl_rows(sh, it)
    d['notes'] = [[l, m] for l, m in sh.get('issues_all', [])
                  if it['item'] and re.search(r'(?<![\w-])%s(?![\w-])' % re.escape(it['item']), m)][:12]
    d['hs_options'] = T.hs_options(P, d['doc_hs'] or it['hs'])
    return jsonify(d)


@app.get('/api/jobs/<int:jid>/log')
@login_required
def job_logs(jid):
    after = int(request.args.get('after') or 0)
    rows = q('SELECT id, ts, line FROM logs WHERE job_id=? AND id>? ORDER BY id LIMIT 2000', (jid, after))
    j = q('SELECT status, message FROM jobs WHERE id=?', (jid,), one=True)
    return jsonify(lines=rows, status=j['status'], message=j['message'], running=RUNNING.get(jid),
                   progress=PROGRESS.get(jid))


@app.get('/api/jobs/<int:jid>/validate')
@login_required
def job_validate(jid):
    j = job_row(jid)
    st = state(jid, j['tool'])
    v = T.validate(T.engine(st.tool), st, j['form'], get_settings())
    v['ships'] = [dict(key=k, name=s.get('hawb') or k, include=s.get('include', True), direction=s.get('direction', ''),
                       items=len(s['items']), bad=sum(1 for i in s['items'] if not i['hs'] or not i['coo'] or
                                                      i.get('hs_qty') is None or not i['hs_uom']),
                       controlled=sum(1 for i in s['items'] if i.get('controlled')),
                       edited=sum(1 for i in s['items'] if any((s.get('overrides') or {}).get(x, {}).keys() - {'sep'}
                                                               for x in i.get('keys', []))))
                  for k, s in st.ships.items()]
    return jsonify(v)


def _busy(jid):
    j = q('SELECT status FROM jobs WHERE id=?', (jid,), one=True)
    if j and j['status'] in ('queued', 'extracting', 'generating'):
        raise ValueError('The job is busy (%s) - wait a moment' % j['status'])


@app.post('/api/jobs/<int:jid>/files')
@login_required
def job_files_add(jid):
    job_row(jid)
    _busy(jid)
    n = _save_uploads(jid, request.files.getlist('files'))
    audit('add files', jid, n)
    if n:                                          # the reader is always chosen from the document layout
        tool = T.detect_tool([f['path'] for f in q("SELECT path FROM files WHERE job_id=? AND kind='input'", (jid,))])
        if tool and tool != job_row(jid)['tool']:
            set_job(jid, tool=tool)
    if n and request.form.get('extract', '1') == '1':
        queue_task(jid, 'extract')
    return jsonify(added=n)


@app.delete('/api/jobs/<int:jid>/files/<int:fid>')
@login_required
def job_file_del(jid, fid):
    _busy(jid)
    f = q("SELECT * FROM files WHERE id=? AND job_id=? AND kind='input'", (fid, jid), one=True)
    if not f:
        abort(404)
    try:
        os.remove(f['path'])
    except OSError:
        pass
    x('DELETE FROM files WHERE id=?', (fid,))
    audit('remove file', jid, f['name'])
    return jsonify(ok=True)


@app.get('/api/jobs/<int:jid>/download/<int:fid>')
@login_required
def job_download(jid, fid):
    f = q('SELECT * FROM files WHERE id=? AND job_id=?', (fid, jid), one=True)
    if not f or not os.path.exists(f['path']):
        abort(404)
    return send_file(f['path'], as_attachment=request.args.get('view') != '1', download_name=f['name'])


@app.get('/api/jobs/<int:jid>/zip')
@login_required
def job_zip(jid):
    j = job_row(jid)
    kind = request.args.get('kind') or 'output'
    fs = q('SELECT * FROM files WHERE job_id=? AND kind=?', (jid, kind))
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, 'w', zipfile.ZIP_DEFLATED) as z:
        for f in fs:
            if os.path.exists(f['path']) and not f['name'].upper().endswith('_ALL.ZIP'):
                z.write(f['path'], f['name'])
    bio.seek(0)
    nm = re.sub(r'[^\w.-]+', '_', (j['job_no'] or j['hawbs'] or 'JOB_%d' % jid).split(',')[0]) + ('_%s.zip' % kind.upper())
    return send_file(bio, as_attachment=True, download_name=nm, mimetype='application/zip')


@app.post('/api/jobs/<int:jid>/extract')
@login_required
def job_extract(jid):
    job_row(jid)
    _busy(jid)
    if not q("SELECT id FROM files WHERE job_id=? AND kind='input'", (jid,), one=True):
        raise ValueError('Upload the PDFs first')
    queue_task(jid, 'extract')
    return jsonify(ok=True)


@app.post('/api/jobs/<int:jid>/generate')
@login_required
def job_generate(jid):
    job_row(jid)
    _busy(jid)
    if not state(jid).ships:
        raise ValueError('Nothing extracted yet')
    queue_task(jid, 'generate')
    return jsonify(ok=True)


@app.post('/api/jobs/<int:jid>/form')
@login_required
def job_form(jid):
    j = job_row(jid)
    _busy(jid)
    d = request.get_json(force=True)
    form = j['form']
    prev_pt = form.get('permit_type') or ''
    for k in ('permit_type', 'job_no', 'party_name', 'party_uen', 'direction', 'cargo_text', 'consolidate',
              'charges_mode', 'split_excel', 'max_items', 'remarks_template', 'header', 'extra_header', 'extra_invoice', 'rules_off',
              'brand_mode', 'exchange_rate', 'j_start', 'cust_list', 'team', 'consolidate_explicit', 'cif', 'uen_names_on', 'split_hs',
              'msg_type', 'decl_auto'):
        if k in d:
            form[k] = d[k]
    if 'permit_type' in d and ((d['permit_type'] or '') != prev_pt or d.get('permit_src')):
        # picked on the page = yours (the advisor keeps it); blank = the advisor decides
        form['permit_src'] = d.get('permit_src') or ('you' if d['permit_type'] else '')
    title = d.get('title')
    set_job(jid, form=json.dumps(form), job_no=form.get('job_no', ''), **({'title': title} if title else {}))
    st = state(jid, j['tool'])
    if 'brand_mode' in d:                       # brand choice of the Item tab = remembered for the supplier
        S_ = get_settings()
        book = S_.setdefault('brand_book', {})
        for sh in st.ships.values():
            nm = T.supplier_of(sh)
            if nm:
                if (d.get('brand_mode') or '').strip():
                    book[nm] = d['brand_mode'].strip().upper()
                else:
                    book.pop(nm, None)
        put_settings(S_)
    with st.lock:
        if st.ships:
            T.resolve_all(j['tool'], P, get_settings(), form, st)
            save(jid, st)
    audit('job details', jid)
    return jsonify(ok=True)


def _ship_edit_response(jid, st, key):
    sm = save(jid, st)
    E = T.engine(st.tool)
    return jsonify(ship=T.ship_json(E, key, st.ships[key], get_settings()) if key in st.ships else None,
                   ships=_ships_list(st, get_settings()), summary=sm, undo=len(st.undo))


# ---------------------------------------------------------------- sheet tabs (one per sheet of the output Excel)
@app.get('/api/jobs/<int:jid>/analysis')
@login_required
def job_analysis(jid):
    """Document extraction analysis: per document what was read + checks, per shipment the cross-document checks."""
    j = job_row(jid)
    st = state(jid, j['tool'])
    if not st.docs and not st.ships:
        return jsonify(documents=[], shipments=[], summary={})
    return jsonify(T.analysis(st))


@app.get('/api/jobs/<int:jid>/sheets')
@login_required
def sheets_get(jid):
    j = job_row(jid)
    st = state(jid, j['tool'])
    if not st.ships:
        return jsonify(files=[], sheets={}, file='')
    if RUNNING.get(jid):
        raise ValueError('The job is busy - wait a moment')
    pdir = os.path.join(P.job(jid), 'preview')
    d = T.preview(st.tool, P, get_settings(), j['form'], st, pdir, os.path.join(P.job(jid), 'state.pkl'))
    names = [f['name'] for f in d['files']]
    name = request.args.get('file') if request.args.get('file') in names else (names[0] if names else '')
    f = next((x for x in d['files'] if x['name'] == name), None)
    out = dict(files=[dict(name=x['name'], ships=x['ships'], orphan=x.get('orphan', 0), direction=x.get('direction', ''))
                      for x in d['files']], file=name, direction=(f or {}).get('direction', ''))
    if f:
        rm = d['rowmaps'].get(name, {})
        out['sheets'] = T.read_sheets(f['path'])
        out['edited'] = T.edited_cells(st, name, rm)
        out['rowmap'] = rm
        out['item_cols'] = T.ITEM_FIELD_COLS
        _learn_ports(out['sheets'].get('Header'))
    if request.args.get('lists') == '1':
        out['lists'] = {k: v for k, v in T.lists(st.tool).items()}
    return jsonify(out)


PORT_COLS = (('Loading Port Code', 'Loading Port Name'), ('Discharge Port Code', 'Discharge Port Name'),
             ('Next Port', 'Next Port Name'), ('Last Port', 'Last Port Name'))


def _learn_ports(sh):
    """Ports read from the documents / typed by the users (code + name) are added to the port list."""
    try:
        H = [str(h or '').strip() for h in sh['headers']]
        row = sh['rows'][0]
        for c, n in PORT_COLS:
            if c in H and n in H:
                T.port_learn(P, row[H.index(c)], row[H.index(n)])
    except Exception:
        pass


@app.get('/api/ports')
@login_required
def ports_find():
    """Port search (code or name) for the magnifier next to Loading / Discharge / Next / Last Port."""
    q = request.args.get('q') or ''
    if request.args.get('exact'):
        L = T.port_list(P)
        c = q.strip().upper().replace(' ', '')
        if c in L:
            return jsonify(found=[c, L[c]])
        n = q.strip().upper()
        hit = next((k for k, v in L.items() if v == n), '')
        return jsonify(found=[hit, n] if hit else None)
    return jsonify(ports=T.port_search(P, q), total=len(T.port_list(P)))


@app.post('/api/ports')
@login_required
def ports_add():
    d = request.get_json(force=True)
    return jsonify(ok=T.port_learn(P, d.get('code'), d.get('name')))


@app.get('/api/uen')
@login_required
def uen_find():
    """CR UEI search by company name (ACRA open data on data.gov.sg; bizfile itself has a reCAPTCHA)."""
    name = (request.args.get('name') or '').strip()
    if len(name) < 3:
        raise ValueError('Type the company name first')
    try:
        hits = acra.search(name, timeout=12)
    except Exception as e:
        raise ValueError('ACRA open data not reachable (%s) - check the internet on the server PC, or look it up on bizfile.gov.sg' % str(e)[:120])
    return jsonify(hits=hits[:15], exact=acra.lookup(name) if hits else '')


@app.get('/api/uen/name')
@login_required
def uen_name():
    """Registered name + address of a UEN (ACRA open data) - remembered in Settings (uen_names)."""
    uen = (request.args.get('uen') or '').strip().upper()
    S_ = get_settings()
    have = (S_.get('uen_names') or {}).get(uen)
    if have and not request.args.get('fresh'):
        return jsonify(dict(have, uen=uen))
    hit = acra.by_uen(uen, request.args.get('hint') or '', timeout=12)
    if not hit.get('name'):
        raise ValueError('UEN %s not found in the ACRA open data%s' % (uen, (' (%s)' % acra.STATE['last_error'][:100]) if acra.STATE.get('last_error') else ''))
    S_.setdefault('uen_names', {})[uen] = dict(name=hit['name'], address=hit.get('address') or [], status=hit.get('status', ''))
    put_settings(S_)
    return jsonify(dict(hit, uen=uen))


@app.get('/api/custmasters')
@login_required
def cust_masters_info():
    """Customer product lists (item / product ID -> HS + CA/SC code) - Settings > Product masters."""
    _, info = T.cust_masters(P)
    return jsonify(files=[dict(file=f, items=n, sheets=sh, updated=datetime.datetime.fromtimestamp(
        os.path.getmtime(os.path.join(P.custdir, f))).strftime('%d/%m/%Y %H:%M')) for f, n, sh in info])


@app.get('/api/custmasters/file')
@login_required
def cust_master_file():
    name = os.path.basename(request.args.get('name') or '')
    if name not in T.cust_master_files(P):
        abort(404)
    return send_from_directory(P.custdir, name, as_attachment=True)


@app.post('/api/custmasters')
@admin_required
def cust_master_put():
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith(('.xlsx', '.xlsm', '.xls')):
        raise ValueError('Choose the Excel file (.xlsx) of the customer product list')
    name = secure_filename(f.filename) or 'customer_products.xlsx'
    tmp = os.path.join(P.custdir, '_upload' + os.path.splitext(name)[1].lower())
    f.save(tmp)
    try:
        m, sheets = T.load_cust_master(tmp)
    except Exception as e:
        os.remove(tmp)
        raise ValueError(str(e))
    dst = os.path.join(P.custdir, name)
    if os.path.exists(dst):
        shutil.copy(dst, dst + '.bak')
    os.replace(tmp, dst)
    if os.path.exists(dst + '.removed'):
        os.remove(dst + '.removed')
    audit('customer product list uploaded', None, '%s (%d items)' % (name, len(m)))
    return jsonify(items=len(m), sheets=sheets)


@app.delete('/api/custmasters')
@admin_required
def cust_master_del():
    name = os.path.basename(request.args.get('name') or '')
    if name not in T.cust_master_files(P):
        abort(404)
    path = os.path.join(P.custdir, name)
    os.replace(path, path + '.removed')          # kept (and not copied back from the program folder)
    audit('customer product list removed', None, name)
    return jsonify(ok=True)


# ---------------------------------------------------------------- SFA team: database + overrides (masters/sfa)
def _sfa():
    import sfa_rules
    return sfa_rules, sfa_rules.setup_folder(P.masters, os.path.join(T.base_dir(), 'masters'))


@app.get('/api/sfa')
@login_required
def sfa_info():
    """SFA team rules - Settings > Product masters > SFA team."""
    R, d = _sfa()
    return jsonify(R.status(d, P.hs))


@app.get('/api/sfa/file')
@login_required
def sfa_file():
    R, d = _sfa()
    name = os.path.basename(request.args.get('name') or '')
    if name in (R.OVR_NAME, R.RULES_NAME, R.DB_NAME) and os.path.exists(os.path.join(d, name)):
        return send_from_directory(d, name, as_attachment=True)
    src = os.path.join(d, R.SRC_NAME)
    if name and os.path.exists(os.path.join(src, name)):
        return send_from_directory(src, name, as_attachment=True)
    abort(404)


def _sfa_rebuild(R, d, why):
    res = R.build_db(os.path.join(d, R.SRC_NAME), os.path.join(d, R.DB_NAME), P.hs)
    audit('SFA database rebuilt', None, '%s: %s' % (why, json.dumps({k: v for k, v in res.items() if k != 'files'})))
    return res


@app.post('/api/sfa')
@admin_required
def sfa_upload():
    """.csv = SFA overrides (replaces the list); .xlsx / .xls = SFA MASTER DATA or an RPT permit report (database rebuilt)."""
    R, d = _sfa()
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith(('.csv', '.xlsx', '.xls', '.xlsm')):
        raise ValueError('Choose the SFA overrides .csv, the SFA MASTER DATA or an RPT permit report (.xlsx / .xls)')
    name = secure_filename(f.filename) or 'upload'
    ext = os.path.splitext(name)[1].lower()
    tmp = os.path.join(d, '_upload' + ext)
    f.save(tmp)
    if ext == '.csv':
        try:
            n = R.check_overrides_file(tmp)
        except Exception as e:
            os.remove(tmp)
            raise ValueError(str(e))
        dst = os.path.join(d, R.OVR_NAME)
        if os.path.exists(dst):
            shutil.copy(dst, dst + '.bak')
        os.replace(tmp, dst)
        audit('SFA overrides uploaded', None, '%d rows' % n)
        return jsonify(**R.status(d, P.hs), message='%d overrides saved' % n)
    kind = R.file_kind(tmp)
    if not kind:
        os.remove(tmp)
        raise ValueError('Not an SFA MASTER DATA (sheet DATA with Prefecture columns) or RPT "HS CODE & PRODUCT CODE" '
                         'report (Description Goods + HS Code)')
    dst = os.path.join(d, R.SRC_NAME, name)
    if os.path.exists(dst):
        shutil.copy(dst, dst + '.bak')
    os.replace(tmp, dst)
    res = _sfa_rebuild(R, d, '%s uploaded (%s)' % (name, kind))
    return jsonify(**R.status(d, P.hs), message='%s read as %s - database rebuilt: %d master rows, %d past permit rows'
                   % (name, 'SFA MASTER DATA' if kind == 'MASTER' else 'permit report', res['sfa_master'], res['permit_history']))


@app.post('/api/sfa/rebuild')
@admin_required
def sfa_rebuild():
    R, d = _sfa()
    res = _sfa_rebuild(R, d, 'rebuild')
    return jsonify(**R.status(d, P.hs), message='Database rebuilt: %d master rows, %d past permit rows (%s)'
                   % (res['sfa_master'], res['permit_history'], ', '.join(res['not_read']) + ' not read' if res['not_read'] else 'all files read'))


@app.post('/api/sfa/importers')
@admin_required
def sfa_importers():
    """Importers whose GLT invoice (Quantity) column is the net weight - one name per line."""
    R, d = _sfa()
    names = (request.get_json(silent=True) or {}).get('names') or []
    if isinstance(names, str):
        names = names.splitlines()
    names = R.write_importers(os.path.join(d, R.QTYNW_NAME), names)
    audit('SFA importers (quantity = net weight)', None, ', '.join(names))
    return jsonify(**R.status(d, P.hs), message='%d importer(s) saved' % len(names))


@app.delete('/api/sfa/source')
@admin_required
def sfa_source_del():
    R, d = _sfa()
    name = os.path.basename(request.args.get('name') or '')
    path = os.path.join(d, R.SRC_NAME, name)
    if not name or not os.path.exists(path):
        abort(404)
    os.replace(path, os.path.join(d, name + '.removed'))
    _sfa_rebuild(R, d, name + ' removed')
    return jsonify(R.status(d, P.hs))


@app.get('/api/cascmaster')
@login_required
def casc_master_info():
    """RS Components product code master (CASC Product Code 1 by HS) - Settings > Product masters."""
    files = sorted(f for f in os.listdir(P.cascdir) if f.lower().endswith(('.xlsx', '.xls', '.xlsm'))) if os.path.isdir(P.cascdir) else []
    m = T.casc_by_hs(P)
    f = files[0] if files else ''
    return jsonify(file=f, hs=len(m), codes=sum(1 for v in m.values() if v),
                   updated=datetime.datetime.fromtimestamp(os.path.getmtime(os.path.join(P.cascdir, f))).strftime('%d/%m/%Y %H:%M') if f else '')


@app.get('/api/cascmaster/file')
@login_required
def casc_master_file():
    files = sorted(f for f in os.listdir(P.cascdir) if f.lower().endswith(('.xlsx', '.xls', '.xlsm')))
    if not files:
        abort(404)
    return send_from_directory(P.cascdir, files[0], as_attachment=True)


@app.post('/api/cascmaster')
@admin_required
def casc_master_put():
    """Upload a new RS product code master (Excel with HS Code + CASC Product Code 1); the old one is kept as .bak."""
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith(('.xlsx', '.xls', '.xlsm')):
        raise ValueError('Choose the Excel file (.xlsx / .xls) with the columns HS Code and CASC Product Code 1')
    ext = os.path.splitext(f.filename)[1].lower()
    tmp = os.path.join(P.cascdir, '_upload' + ext)
    f.save(tmp)
    try:
        m = T.load_casc_master(tmp)
    except Exception as e:
        os.remove(tmp)
        raise ValueError(str(e))
    for old in os.listdir(P.cascdir):
        if old.lower().endswith(('.xlsx', '.xls', '.xlsm')) and not old.startswith('_upload'):
            shutil.move(os.path.join(P.cascdir, old), os.path.join(P.cascdir, old + '.bak'))
    os.replace(tmp, os.path.join(P.cascdir, os.path.splitext(T.CASC_MASTER)[0] + ext))
    audit('RS product code master uploaded', None, '%s (%d HS)' % (f.filename, len(m)))
    return jsonify(hs=len(m), codes=sum(1 for v in m.values() if v))


@app.post('/api/ports/file')
@admin_required
def ports_file():
    """Upload a full port list (Excel: column A = UN/LOCODE e.g. CNASN, column B = name e.g. ANSHAN)."""
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith('.xlsx'):
        raise ValueError('Choose an Excel (.xlsx) file: column A code, column B name')
    dst = os.path.join(P.masters, 'PORTS.xlsx')
    if os.path.exists(dst):
        shutil.copy(dst, dst + '.bak')
    f.save(dst)
    audit('port list uploaded', None, f.filename)
    return jsonify(total=len(T.port_list(P)))


PER_FILE_COLS = {"Trader's Remarks", 'Total Gross Weight'}


@app.post('/api/jobs/<int:jid>/cells')
@login_required
def cells_edit(jid):
    """Values typed on a sheet tab. COO / HS / HS UOM / CASC of an item go through the item rules (HS qty, CASC
    sheet ...); every other value is written as typed into that Excel (kept when the job is read again)."""
    j = job_row(jid)
    _busy(jid)
    d = request.get_json(force=True)
    st = state(jid, j['tool'])
    fname = d.get('file') or ''
    pdir = os.path.join(P.job(jid), 'preview')
    S = get_settings()
    with st.lock:
        try:
            with open(os.path.join(pdir, 'preview.json'), encoding='utf-8') as f:
                rms = json.load(f)['rowmaps']
            rm = rms.get(fname, {})
            all_files = list(rms) or [fname]
        except Exception:
            rm, all_files = {}, [fname]
        rmi = {int(k): v for k, v in rm.items()}
        n_rule = n_cell = 0
        touched = set()
        done_all = set()
        for e in d.get('edits') or []:
            sheet, row, col = e.get('sheet'), int(e.get('row') or 0), str(e.get('col') or '')
            if sheet not in T.SHEET_TABS or sheet == 'Items to be consolidated' or row < 2 or not col:
                raise ValueError('This cell can not be changed')
            val = e.get('value')
            key = T.cell_key(sheet, row, col, rmi)
            field = T.ITEM_FIELD_COLS.get(col) if sheet == 'Item' else None
            if field and row in rmi and not e.get('as_typed'):
                skey, idx = rmi[row][0], rmi[row][1]
                sh = st.ships.get(skey)
                if sh and idx < len(sh['items']):
                    if skey not in touched:
                        T.push_undo(st, skey)
                        touched.add(skey)
                    if e.get('reset'):
                        ov = sh['overrides']
                        for k in sh['items'][idx]['keys']:
                            ov.get(k, {}).pop(field, None)
                    elif d.get('all_files'):                  # Copy to all: every item of every Excel of the job
                        if (field, str(val)) not in done_all:
                            done_all.add((field, str(val)))
                            for k2, sh2 in st.ships.items():
                                if k2 not in touched:
                                    T.push_undo(st, k2)
                                    touched.add(k2)
                                T.set_item_field(sh2, list(range(len(sh2['items']))), field, str(val if val is not None else ''))
                    else:
                        T.set_item_field(sh, [idx], field, str(val if val is not None else ''))
                        if d.get('save_master') and field in ('hs', 'uom', 'coo', 'casc'):
                            _to_master(st.tool, [sh['items'][idx]], field, str(val or ''))
                    st.cells.get(fname, {}).pop(key, None)
                    n_rule += 1
                    continue
            # Header / Invoice / Container values: the same on EVERY Excel of the job (split permits) - except the values
            # that belong to one file (its remarks with J-n / invoices, its gross weight)
            per_file = sheet == 'Item' or (sheet == 'Header' and col.strip() in PER_FILE_COLS)
            if sheet == 'Item' and d.get('all_files') and not e.get('reset'):   # "Copy to all": every item of every Excel
                for fn in all_files:
                    rmf = {int(k): v for k, v in (rms.get(fn) or {}).items()}
                    for r_ in rmf:
                        st.cells.setdefault(fn, {})[T.cell_key(sheet, r_, col, rmf)] = '' if val is None else str(val)
            for fn in ([fname] if per_file else all_files):
                cells = st.cells.setdefault(fn, {})
                if e.get('reset'):
                    cells.pop(key, None)
                else:
                    cells[key] = '' if val is None else str(val)
            n_cell += 1
        for skey in touched:
            T.resolve_one(st.tool, P, S, j['form'], st, st.ships[skey])
        sm = save(jid, st, j)
        audit('sheet edit', jid, '%s: %s' % (fname, '; '.join('%s r%s %s=%s' % (e.get('sheet'), e.get('row'), e.get('col'),
                                                                          '(rules)' if e.get('reset') else e.get('value'))
                                                             for e in (d.get('edits') or [])[:12])))
    return jsonify(ok=True, rules=n_rule, typed=n_cell, summary=sm, undo=len(st.undo))


@app.delete('/api/brand_book/<path:name>')
@admin_required
def brand_book_del(name):
    S = get_settings()
    (S.get('brand_book') or {}).pop(name.upper(), None)
    put_settings(S)
    audit('brand forget', detail=name)
    return jsonify(ok=True, brand_book=S.get('brand_book') or {})


@app.post('/api/jobs/<int:jid>/info')
@login_required
def job_info(jid):
    """Job information panel: permit no. (after approval) and internal remarks."""
    job_row(jid)
    d = request.get_json(force=True)
    upd = {k: str(d[k]).strip()[:300] for k in ('permit_no', 'int_remarks', 'title', 'job_no') if k in d}
    if upd:
        set_job(jid, **upd)
        audit('job info', jid, ', '.join('%s=%s' % kv for kv in upd.items()))
    return jsonify(ok=True)


@app.post('/api/jobs/<int:jid>/items')
@login_required
def items_edit(jid):
    j = job_row(jid)
    _busy(jid)
    d = request.get_json(force=True)
    st = state(jid, j['tool'])
    with st.lock:
        key = d.get('ship')
        sh = st.ships.get(key)
        if not sh:
            abort(404)
        rows = [int(r) for r in d.get('rows') or [] if 0 <= int(r) < len(sh['items'])]
        if not rows:
            raise ValueError('Select item rows first')
        T.push_undo(st, key)
        act = d.get('action') or 'set'
        S = get_settings()
        if act == 'set':
            if d.get('field') not in T.ITEM_FIELDS:
                raise ValueError('Field cannot be edited')
            T.set_item_field(sh, rows, d['field'], str(d.get('value') if d.get('value') is not None else ''))
            if d.get('save_master') and d['field'] in ('hs', 'uom', 'coo', 'casc'):
                _to_master(st.tool, [sh['items'][r] for r in rows], d['field'], str(d.get('value') or ''))
            if d['field'] == 'hs' and d.get('remember') and d.get('value'):
                for r in rows:
                    dh = sh['items'][r].get('doc_hs8') or sh['items'][r].get('doc_hs')
                    if dh:
                        S.setdefault('hs_map', {})[dh] = re.sub(r'\D', '', d['value'])
                put_settings(S)
        elif act == 'clear':
            T.clear_item_edits(sh, rows)
        elif act in ('sep_on', 'sep_off'):
            T.set_sep(sh, rows, act == 'sep_on')
        else:
            raise ValueError('Unknown action')
        T.resolve_one(st.tool, P, S, j['form'], st, sh)
        audit('items ' + act, jid, '%s rows %s %s=%s' % (key, rows[:20], d.get('field', ''), d.get('value', '')))
        return _ship_edit_response(jid, st, key)


def _to_master(tool, items, field, value):
    E = T.engine(tool)
    path = P.master(tool)
    m = E.load_master(path)
    for it in items:
        for item in it['item'].split(', '):
            d = m.setdefault(item, dict(desc=it.get('desc', ''), hs='', uom='', coo='', casc='', brand='', remarks=''))
            d[field] = (E.h8(value)[:8] if value else '') if field == 'hs' else value.upper()
            if not d.get('desc'):
                d['desc'] = it.get('desc', '')
    E.save_master(path, m)


@app.post('/api/jobs/<int:jid>/ship')
@login_required
def ship_edit(jid):
    j = job_row(jid)
    _busy(jid)
    d = request.get_json(force=True)
    st = state(jid, j['tool'])
    with st.lock:
        key = d.get('ship')
        sh = st.ships.get(key)
        if not sh:
            abort(404)
        S = get_settings()
        E = T.engine(st.tool)
        T.push_undo(st, key)
        if d.get('field') == 'container_tne':
            T.container_tne(sh, int(E.fnum(S['rules'].get('cont_dec')) or 2))
        else:
            T.set_ship_field(E, sh, d.get('field'), d.get('value'), S)
            if d.get('field') == 'party_uen':
                put_settings(S)
        if d.get('field') in ('containers', 'container_tne') and sh.get('hawb'):
            for s in st.ships.values():           # same HBL split in parts: same containers
                if s is not sh and s.get('hawb') == sh['hawb']:
                    if 'containers_edit' in sh:
                        s['containers_edit'] = [dict(r) for r in sh['containers_edit']]
                    else:
                        s.pop('containers_edit', None)
                    T.resolve_one(st.tool, P, S, j['form'], st, s)
        T.resolve_one(st.tool, P, S, j['form'], st, sh)
        audit('ship ' + str(d.get('field')), jid, '%s = %s' % (key, json.dumps(d.get('value'))[:300]))
        return _ship_edit_response(jid, st, key)


@app.post('/api/jobs/<int:jid>/undo')
@login_required
def job_undo(jid):
    j = job_row(jid)
    _busy(jid)
    st = state(jid, j['tool'])
    with st.lock:
        key = T.undo(st)
        if key is None:
            raise ValueError('Nothing to undo')
        T.resolve_one(st.tool, P, get_settings(), j['form'], st, st.ships[key])
        audit('undo', jid, key)
        return _ship_edit_response(jid, st, key)


@app.post('/api/jobs/<int:jid>/status')
@login_required
def job_status(jid):
    j = job_row(jid)
    s = request.get_json(force=True).get('status')
    if s not in ('review', 'done', 'archived'):
        raise ValueError('bad status')
    _busy(jid)
    set_job(jid, status=s, message=('archived' if s == 'archived' else j['message']))
    audit('status ' + s, jid)
    return jsonify(ok=True)


@app.delete('/api/jobs/<int:jid>')
@login_required
def job_delete(jid):
    j = job_row(jid)
    if session.get('role') != 'admin' and j['created_by'] != session['user']:
        return jsonify(error='Only the creator or an admin can delete a job'), 403
    _busy(jid)
    with _states_lock:
        _states.pop(jid, None)
    shutil.rmtree(os.path.join(P.jobs, str(jid)), ignore_errors=True)
    for t in ('files', 'logs'):
        x('DELETE FROM %s WHERE job_id=?' % t, (jid,))
    x('DELETE FROM jobs WHERE id=?', (jid,))
    audit('job delete', jid, j['title'])
    return jsonify(ok=True)


@app.post('/api/jobs/<int:jid>/copy')
@login_required
def job_copy(jid):
    """New job with the same PDFs + details (e.g. to try other settings)."""
    j = job_row(jid)
    nid = x('INSERT INTO jobs(tool, title, job_no, status, message, created_by, created, updated, form) VALUES(?,?,?,?,?,?,?,?,?)',
            (j['tool'], (j['title'] or '') + ' (copy)', j['job_no'], 'new', '', session['user'], now(), now(),
             json.dumps(j['form'])))
    for f in q("SELECT * FROM files WHERE job_id=? AND kind='input'", (jid,)):
        if os.path.exists(f['path']):
            dst = os.path.join(P.job(nid), 'in', f['name'])
            shutil.copy(f['path'], dst)
            add_file(nid, 'input', dst, session['user'])
    audit('job copy', nid, 'from %d' % jid)
    queue_task(nid, 'extract')
    return jsonify(id=nid)


# ------------------------------------------------------------------ HS codes
_HS = {'mtime': None, 'data': None}


def hs_data():
    m = os.path.getmtime(P.hs)
    if _HS['mtime'] != m:
        with open(P.hs, encoding='utf-8') as f:
            _HS['data'] = json.load(f)
        _HS['mtime'] = m
    return _HS['data']


@app.get('/api/hs/search')
@login_required
def hs_search():
    s = (request.args.get('q') or '').strip().upper()
    d = hs_data()
    out = []
    if s:
        digits = re.sub(r'\D', '', s)
        words = [w for w in re.split(r'\s+', s) if w]
        for k, r in d['hs'].items():
            if (digits and len(digits) >= 2 and k.startswith(digits)) or \
                    (not digits and all(w in (r[2] or '').upper() for w in words)):
                pc = d.get('pcodes', {}).get(k) or []
                out.append(dict(hs=k, uom=r[1], desc=r[2], ctl_imp=r[3], ctl_exp=r[4], ca_imp=r[5], ca_exp=r[6], duty=r[7],
                                idamisc=any(c[0] == 'IDAMISC' for c in pc), pcodes=[c[0] for c in pc][:30], npc=len(pc)))
                if len(out) >= 300:
                    break
    return jsonify(rows=out, count=d.get('count'))


@app.get('/api/hs/candidates')
@login_required
def hs_cands():
    return jsonify(T.hs_options(P, request.args.get('hs') or ''))


@app.get('/api/hsmap')
@login_required
def hsmap_get():
    return jsonify(sorted((get_settings().get('hs_map') or {}).items()))


@app.post('/api/hsmap')
@login_required
def hsmap_set():
    d = request.get_json(force=True)
    S = get_settings()
    a, b = re.sub(r'\D', '', d.get('from') or ''), re.sub(r'\D', '', d.get('to') or '')
    if not a:
        raise ValueError('Supplier HS code missing')
    if b:
        if len(b) != 8:
            raise ValueError('Singapore HS code must be 8 digits')
        S.setdefault('hs_map', {})[a] = b
    else:
        S.setdefault('hs_map', {}).pop(a, None)
    put_settings(S)
    audit('hs map', detail='%s -> %s' % (a, b or '(removed)'))
    return jsonify(sorted(S['hs_map'].items()))


@app.post('/api/hs/upload')
@admin_required
def hs_upload():
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith(('.xlsx', '.xlsm', '.json')):
        raise ValueError('Choose the UPDATE_HS_CODE.xlsx file')
    tmp = os.path.join(DATA, 'upload_' + secure_filename(f.filename))
    f.save(tmp)
    try:
        if tmp.lower().endswith('.json'):
            json.load(open(tmp, encoding='utf-8'))['hs']
            shutil.copy(tmp, P.hs + '.new')
        else:
            T.engine('danfoss').save_hs_data(tmp, P.hs + '.new')
        if os.path.exists(P.hs):
            shutil.copy(P.hs, P.hs + '.bak')
        os.replace(P.hs + '.new', P.hs)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    d = hs_data()
    audit('hs list update', detail='%s codes' % d.get('count'))
    return jsonify(count=d.get('count'), pcodes=d.get('pcode_count'))


# ------------------------------------------------------------------ settings / masters / backup / audit
@app.get('/api/settings')
@login_required
def settings_get():
    S = get_settings()
    S.pop('uen_book', None) if session.get('role') != 'admin' else None
    return jsonify(S)


@app.post('/api/settings')
@admin_required
def settings_set():
    d = request.get_json(force=True)
    S = get_settings()
    for k in ('header', 'options', 'xml', 'rules', 'control'):
        if isinstance(d.get(k), dict):
            S[k].update(d[k])
    for k in ('insurance_pct', 'insurance_cur', 'remarks_template'):
        if k in d:
            S[k] = d[k]
    if isinstance(d.get('teams'), list):
        tl = [str(t).strip().upper() for t in d['teams'] if str(t).strip()]
        S['teams'] = list(dict.fromkeys(tl)) or ['DGF OCEAN']
    if isinstance(d.get('cust_teams'), dict):
        S['cust_teams'] = {str(a): str(b).strip().upper() for a, b in d['cust_teams'].items()}
    if isinstance(d.get('hs_keywords'), list):
        S['hs_keywords'] = [[str(x[0]).strip().upper(), re.sub(r'\D', '', str(x[1]))[:10]] for x in d['hs_keywords']
                            if isinstance(x, (list, tuple)) and len(x) >= 2 and str(x[0]).strip()]
    if isinstance(d.get('forwarder_map'), dict):
        S['forwarder_map'] = {a.strip().upper(): b.strip() for a, b in d['forwarder_map'].items() if a.strip() and b.strip()}
    if isinstance(d.get('locations'), dict):
        S['locations'] = {a.strip().upper(): b.strip().upper() for a, b in d['locations'].items() if a.strip() and b.strip()}
    if isinstance(d.get('uen_book'), dict):
        S['uen_book'] = d['uen_book']
    if isinstance(d.get('carrier_book'), dict):          # carrier agent UEN -> name
        S['carrier_book'] = {re.sub(r'\s', '', str(a)).upper(): str(b).strip().upper()
                             for a, b in d['carrier_book'].items() if str(a).strip() and str(b).strip()}
    if isinstance(d.get('msg_types'), list):             # XML message types offered on the job
        S['msg_types'] = list(dict.fromkeys(str(x).strip().upper() for x in d['msg_types'] if str(x).strip())) or ['IPTDEC']
    if isinstance(d.get('msg_rules'), list):             # [team / importer / exporter name contains, message type]
        S['msg_rules'] = [[str(x[0]).strip().upper(), str(x[1]).strip().upper()] for x in d['msg_rules']
                          if isinstance(x, (list, tuple)) and len(x) >= 2 and str(x[0]).strip() and str(x[1]).strip()]
    for k in ('extra_header', 'extra_invoice', 'declarant'):
        if isinstance(d.get(k), dict):
            S[k] = {a: b for a, b in d[k].items() if str(b).strip()}
    put_settings(S)
    audit('settings')
    return jsonify(S)


@app.get('/api/masters/<tool>')
@login_required
def master_get(tool):
    if tool not in T.TOOLS or not os.path.exists(P.master(tool)):
        abort(404)
    return send_file(P.master(tool), as_attachment=True, download_name=T.TOOLS[tool]['master'])


@app.post('/api/masters/<tool>')
@admin_required
def master_put(tool):
    if tool not in T.TOOLS:
        abort(404)
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith('.xlsx'):
        raise ValueError('Choose the product master .xlsx')
    tmp = P.master(tool) + '.upload'
    f.save(tmp)
    try:
        n = len(T.engine(tool).load_master(tmp))
    except Exception as e:
        os.remove(tmp)
        raise ValueError('Not a product master: %s' % e)
    if os.path.exists(P.master(tool)):
        shutil.copy(P.master(tool), P.master(tool) + '.bak')
    os.replace(tmp, P.master(tool))
    audit('master upload', detail='%s %d items' % (tool, n))
    return jsonify(items=n)


# ---------------------------------------------------------------- HS lists by description (per customer)
def _hsd_list(S=None):
    S = S or get_settings()
    out = []
    for l in S.get('hs_desc_lists') or []:
        p = os.path.join(P.hsdesc, l.get('file') or '')
        idx = next((ix for nm, rx, ix in T.hs_desc_indexes(P, {'hs_desc_lists': [l]})), None) if os.path.exists(p) else None
        out.append(dict(name=l.get('name', ''), match=l.get('match', ''), file=l.get('file', ''), exists=os.path.exists(p),
                        descriptions=len(idx) if idx else 0,
                        updated=time.strftime('%Y-%m-%d %H:%M', time.localtime(os.path.getmtime(p))) if os.path.exists(p) else ''))
    return out


@app.get('/api/hsdesc')
@login_required
def hsd_get():
    return jsonify(lists=_hsd_list())


@app.post('/api/hsdesc')
@admin_required
def hsd_put():
    """Add a customer HS list or upload a new version of one (same name = replaced, the old file kept as .bak)."""
    name = (request.form.get('name') or '').strip().upper()
    match = (request.form.get('match') or '').strip() or re.escape(name.split()[0] if name else '')
    if not name:
        raise ValueError('Type the customer name (e.g. PANASONIC FACTORY)')
    re.compile(match)
    S = get_settings()
    lists = S.get('hs_desc_lists') or []
    cur = next((l for l in lists if l.get('name', '').upper() == name), None)
    f = request.files.get('file')
    if f and f.filename:
        if not f.filename.lower().endswith(('.xlsx', '.csv')):
            raise ValueError('Choose the HS list .xlsx (columns: description + HS code)')
        fn = (cur or {}).get('file') or re.sub(r'[^A-Z0-9]+', '_', name) + '_HS_LIST' + os.path.splitext(f.filename)[1].lower()
        dst = os.path.join(P.hsdesc, fn)
        tmp = dst + '.upload' + os.path.splitext(f.filename)[1].lower()
        f.save(tmp)
        try:
            n = len(T.engine('invoice').load_hs_desc_list(tmp))
            if not n:
                raise ValueError('no description + HS code rows found')
        except Exception as e:
            os.remove(tmp)
            raise ValueError('Not an HS list: %s' % e)
        if os.path.exists(dst):
            shutil.copy(dst, dst + '.bak')
        os.replace(tmp, dst)
    elif not cur:
        raise ValueError('Choose the HS list file')
    else:
        fn, n = cur['file'], None
    if cur:
        cur.update(match=match, file=fn)
    else:
        lists.append(dict(name=name, match=match, file=fn))
    S['hs_desc_lists'] = lists
    put_settings(S)
    audit('hs list by description', detail='%s %s' % (name, ('%d rows' % n) if n else 'match changed'))
    return jsonify(lists=_hsd_list(S))


@app.get('/api/hsdesc/<path:name>/file')
@login_required
def hsd_file(name):
    l = next((l for l in get_settings().get('hs_desc_lists') or [] if l.get('name', '').upper() == name.upper()), None)
    p = os.path.join(P.hsdesc, (l or {}).get('file') or '')
    if not l or not os.path.exists(p):
        abort(404)
    return send_file(p, as_attachment=True, download_name=l['file'])


@app.delete('/api/hsdesc/<path:name>')
@admin_required
def hsd_del(name):
    S = get_settings()
    S['hs_desc_lists'] = [l for l in S.get('hs_desc_lists') or [] if l.get('name', '').upper() != name.upper()]
    put_settings(S)
    audit('hs list by description removed', detail=name)
    return jsonify(lists=_hsd_list(S))


@app.get('/api/backup')
@admin_required
def backup():
    bio = io.BytesIO()
    with _db_write:
        bk = os.path.join(DATA, 'backup_tmp.db')
        src = db()
        dst = sqlite3.connect(bk)
        src.backup(dst)
        dst.close()
    with zipfile.ZipFile(bio, 'w', zipfile.ZIP_DEFLATED) as z:
        z.write(bk, 'permit.db')
        z.write(P.hs, 'hs_data.json')
        for f in os.listdir(P.masters):
            if f.lower().endswith('.xlsx'):
                z.write(os.path.join(P.masters, f), 'masters/' + f)
        for f in os.listdir(P.hsdesc):
            if f.lower().endswith(('.xlsx', '.csv')):
                z.write(os.path.join(P.hsdesc, f), 'masters/hs_desc/' + f)
        for root, _d, fs in os.walk(PLUGIN_DIR):        # supplier plugins (every backup)
            for f in fs:
                if f.endswith('.py'):
                    p = os.path.join(root, f)
                    z.write(p, os.path.relpath(p, DATA))
        if request.args.get('jobs') == '1':
            for root, _d, fs in os.walk(P.jobs):
                for f in fs:
                    p = os.path.join(root, f)
                    z.write(p, os.path.relpath(p, DATA))
    os.remove(bk)
    bio.seek(0)
    audit('backup', detail='with jobs' if request.args.get('jobs') == '1' else 'database')
    return send_file(bio, as_attachment=True, mimetype='application/zip',
                     download_name='PERMIT_SERVER_BACKUP_%s.zip' % datetime.datetime.now().strftime('%Y%m%d_%H%M'))


@app.post('/api/admin/stop')
def admin_stop():
    """Stop the server: STOP_SERVER.bat (key file) or an admin on the server PC itself."""
    d = request.get_json(silent=True) or {}
    try:
        key = open(STOP_KEY).read().strip()
    except Exception:
        key = None
    ok_key = key and d.get('key') == key and is_local()
    ok_admin = session.get('role') == 'admin' and is_local()
    if not (ok_key or ok_admin):
        return jsonify(error='Only STOP_SERVER.bat or an admin on the server PC can stop the server'), 403
    audit('server stop')
    _shutdown_soon()
    return jsonify(ok=True)


@app.post('/api/admin/online')
def admin_online():
    """GO_ONLINE.bat when the server already runs: switch the internet link on (key file + this PC only)."""
    d = request.get_json(silent=True) or {}
    try:
        key = open(STOP_KEY).read().strip()
    except Exception:
        key = None
    if not (key and d.get('key') == key and is_local()):
        return jsonify(error='Only GO_ONLINE.bat on the server PC can do this'), 403
    if d.get('start'):
        tunnel_start()
        if not d.get('quiet'):
            audit('online on', detail='GO_ONLINE.bat')
    return jsonify(status=TUNNEL.get('status', 'off'), url=TUNNEL.get('url', ''))


# ---------------------------------------------------------------- suppliers: built-in tools, layouts, plugins
PLUGIN_DIR = T.plugin_dir(DATA)


def _plugins_json():
    return dict(folder=PLUGIN_DIR, local=is_local(),
                tools=[dict(key=k, name=v['name'], hint=v['hint'], plugin=v.get('plugin', ''), color=v['color'])
                       for k, v in T.TOOLS.items()],
                layouts=T.builtin_layouts(), plugins=T.PLUGINS,
                files=sorted(f for f in os.listdir(PLUGIN_DIR) if f.endswith('.py')))


@app.get('/api/plugins')
@login_required
def plugins_get():
    return jsonify(_plugins_json())


@app.post('/api/plugins/reload')
@admin_required
def plugins_reload():
    T.load_plugins(PLUGIN_DIR)
    audit('plugins reload', detail=', '.join('%s %s' % (p['file'], p['status']) for p in T.PLUGINS))
    return jsonify(_plugins_json())


@app.post('/api/plugins/upload')
@admin_required
def plugins_upload():
    """A plugin is program code: only an admin sitting at the server PC can add one (never over the internet)."""
    if not is_local():
        return jsonify(error='Add plugins on the server PC itself (http://localhost) - not over the network'), 403
    f = request.files.get('file')
    name = secure_filename(f.filename if f else '')
    if not name.endswith('.py') or name.startswith('_'):
        raise ValueError('Choose the plugin .py file (name not starting with _)')
    dst = os.path.join(PLUGIN_DIR, name)
    if os.path.exists(dst):
        os.makedirs(os.path.join(PLUGIN_DIR, '_old'), exist_ok=True)
        shutil.copy(dst, os.path.join(PLUGIN_DIR, '_old', '%s_%s' % (datetime.datetime.now().strftime('%Y%m%d_%H%M%S'), name)))
    f.save(dst)
    T.load_plugins(PLUGIN_DIR)
    rec = next((p for p in T.PLUGINS if p['file'] == name), {})
    audit('plugin add', detail='%s %s %s' % (name, rec.get('status'), rec.get('detail', '')))
    return jsonify(dict(_plugins_json(), added=rec))


@app.post('/api/plugins/<name>/off')
@admin_required
def plugins_off(name):
    if not is_local():
        return jsonify(error='Switch plugins off on the server PC itself'), 403
    name = secure_filename(name)
    src = os.path.join(PLUGIN_DIR, name)
    if not os.path.exists(src):
        abort(404)
    os.makedirs(os.path.join(PLUGIN_DIR, '_off'), exist_ok=True)
    shutil.move(src, os.path.join(PLUGIN_DIR, '_off', name))
    T.load_plugins(PLUGIN_DIR)
    audit('plugin off', detail=name)
    return jsonify(_plugins_json())


@app.get('/api/plugins/<name>/download')
@admin_required
def plugins_download(name):
    if name == 'TEMPLATE':
        return send_file(os.path.join(T.base_dir(), 'PLUGIN_TEMPLATE.py'), as_attachment=True, download_name='PLUGIN_TEMPLATE.py')
    if name == 'GUIDE':
        return send_file(os.path.join(T.base_dir(), 'HOW_TO_ADD_A_SUPPLIER.txt'), as_attachment=True,
                         download_name='HOW_TO_ADD_A_SUPPLIER.txt')
    p = os.path.join(PLUGIN_DIR, secure_filename(name))
    if not os.path.exists(p):
        abort(404)
    return send_file(p, as_attachment=True, download_name=os.path.basename(p))


@app.post('/api/plugins/test')
@admin_required
def plugins_test():
    """Read one PDF with the chosen tool (or Auto) and show what came out - nothing is saved."""
    f = request.files.get('file')
    if not f or not f.filename:
        raise ValueError('Choose a PDF')
    tmp = os.path.join(DATA, 'tmp_test')
    shutil.rmtree(tmp, ignore_errors=True)
    os.makedirs(tmp)
    p = os.path.join(tmp, secure_filename(f.filename) or 'test.pdf')
    f.save(p)
    log = []
    tool = request.form.get('tool') or 'auto'
    if tool != 'auto' and tool not in T.TOOLS:
        raise ValueError('unknown tool')
    try:
        res = T.test_file(tool, P, get_settings(), p, log.append)
    except Exception as e:
        res = dict(tool=tool, layout='', docs=[], error='%s: %s' % (type(e).__name__, e), trace=traceback.format_exc()[-1500:])
    res['log'] = log[-60:]
    shutil.rmtree(tmp, ignore_errors=True)
    return jsonify(res)


@app.get('/api/server')
@login_required
def server_info():
    return jsonify(since=RUN_PORT.get('since'), background=RUN_PORT.get('background', False), local=is_local(),
                   port=RUN_PORT.get('port'), data=DATA if session.get('role') == 'admin' else '')


@app.get('/api/online')
@login_required
def online_get():
    c = load_conf()
    return jsonify(status=TUNNEL['status'], url=TUNNEL['url'], log=TUNNEL['log'][-15:] if session.get('role') == 'admin' else [],
                   fixed=bool(c.get('tunnel_token')), auto=bool(c.get('online')), link=github_link(c),
                   github_repo=c.get('github_repo', '') if session.get('role') == 'admin' else '',
                   github_token=bool(c.get('github_token')), github=TUNNEL.get('github', ''),
                   public_url=c.get('public_url', ''), local=is_local())


@app.post('/api/online')
@admin_required
def online_set():
    d = request.get_json(force=True)
    conf = load_conf()
    if 'auto' in d:
        conf['online'] = bool(d['auto'])
    if 'github_repo' in d:
        repo = (d['github_repo'] or '').strip().strip('/')
        repo = re.sub(r'^https?://github\.com/', '', repo)
        if repo and not re.fullmatch(r'[\w.-]+/[\w.-]+', repo):
            raise ValueError('GitHub repository as  owner/name  e.g.  infinity-sg/permit-link')
        conf['github_repo'] = repo
        if d.get('github_token'):
            conf['github_token'] = d['github_token'].strip()
        audit('github link', detail=repo)
    restart = False
    if 'tunnel_token' in d or 'public_url' in d:      # own domain (Cloudflare named tunnel) - a secret: server PC only
        if not is_local():
            raise ValueError('Set the company domain on the server PC itself (http://localhost)')
        if 'tunnel_token' in d:
            raw = (d.get('tunnel_token') or '').strip()
            m = re.search(r'(eyJ[A-Za-z0-9_\-=]{40,})', raw)    # the token, also from a pasted "cloudflared ... install <token>"
            if raw and not m:
                raise ValueError('That is not a Cloudflare tunnel token - copy the long text starting with eyJ...')
            if m or d.get('clear'):
                conf['tunnel_token'] = m.group(1) if m else ''
                restart = True
        if 'public_url' in d:
            u = (d.get('public_url') or '').strip().rstrip('/')
            if u and not u.startswith('http'):
                u = 'https://' + u
            if u and not re.fullmatch(r'https?://[A-Za-z0-9.-]+\.[A-Za-z]{2,}(:\d+)?', u):
                raise ValueError('Address like  permits.yourcompany.com')
            conf['public_url'] = u
        audit('company domain', detail=conf.get('public_url', '') + (' (token set)' if conf.get('tunnel_token') else ' (no token)'))
    if 'auto' in d or 'github_repo' in d or 'tunnel_token' in d or 'public_url' in d:
        with open(CONF, 'w', encoding='utf-8') as f:
            json.dump(conf, f, indent=1)
    if restart and TUNNEL.get('proc') and TUNNEL['proc'].poll() is None:
        tunnel_stop()
        tunnel_start()
    if d.get('publish') and TUNNEL['status'] == 'online':
        github_publish(TUNNEL['url'])
    if d.get('on') is True:
        tunnel_start()
        audit('internet on')
    elif d.get('on') is False:
        if not is_local():
            raise ValueError('Switch the internet address off on the server PC itself - you are using it')
        tunnel_stop()
        audit('internet off')
    return online_get()


@app.get('/api/audit')
@admin_required
def audit_list():
    return jsonify(q('SELECT * FROM audit ORDER BY id DESC LIMIT ?', (int(request.args.get('limit') or 500),)))


# ======================================================================== start
def addresses():
    out = []
    try:
        for info in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET):
            ip = info[4][0]
            if not ip.startswith('127.') and ip not in out:
                out.append(ip)
    except Exception:
        pass
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('10.255.255.255', 1))
        ip = s.getsockname()[0]
        s.close()
        if ip not in out and not ip.startswith('127.'):
            out.insert(0, ip)
    except Exception:
        pass
    port = RUN_PORT.get('port') or load_conf()['port']
    return ['http://%s:%d' % (ip, port) for ip in out]


RUN_PORT = {}
TUNNEL = dict(proc=None, url='', status='off', log=[])
CF_URL = 'https://github.com/cloudflare/cloudflared/releases/latest/download/cloudflared-windows-amd64.exe'


def _cloudflared():
    exe = 'cloudflared.exe' if os.name == 'nt' else 'cloudflared'
    for p in (os.path.join(app_dir(), exe), os.path.join(DATA, exe), os.path.join(T.base_dir(), exe)):
        if os.path.exists(p):
            return p
    w = shutil.which('cloudflared')
    if w:
        return w
    if os.name != 'nt':
        return ''
    dst = os.path.join(DATA, exe)
    TUNNEL['status'] = 'downloading cloudflared (one time, about 60 MB) ...'
    print(' Internet    : downloading cloudflared.exe (one time) ...')
    import urllib.request
    urllib.request.urlretrieve(CF_URL, dst + '.part')
    os.replace(dst + '.part', dst)
    return dst


def tunnel_start():
    """Public https address for this PC through Cloudflare Tunnel - no router / firewall change needed.
    No token: a free random address (https://xxxx.trycloudflare.com, new each start).
    tunnel_token in server_config.json: your own fixed address (set up once on dash.cloudflare.com)."""
    if TUNNEL['proc'] and TUNNEL['proc'].poll() is None:
        return
    conf = load_conf()
    TUNNEL.update(url='', status='starting ...', log=[])

    def run():
        import subprocess
        try:
            exe = _cloudflared()
            if not exe:
                TUNNEL['status'] = 'cloudflared not found - install it (see README)'
                return
            port = RUN_PORT.get('port') or conf['port']
            if conf.get('tunnel_token'):
                args = [exe, 'tunnel', '--no-autoupdate', 'run', '--token', conf['tunnel_token']]
                TUNNEL['url'] = conf.get('public_url', '')
            else:
                args = [exe, 'tunnel', '--no-autoupdate', '--url', 'http://127.0.0.1:%d' % port]
            kw = {}
            if os.name == 'nt':
                kw['creationflags'] = 0x08000000          # no extra window
            p = subprocess.Popen(args, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
                                 encoding='utf-8', errors='replace', **kw)
            TUNNEL['proc'] = p
            for ln in p.stdout:
                TUNNEL['log'] = (TUNNEL['log'] + [ln.rstrip()])[-60:]
                m = re.search(r'https://[a-z0-9-]+\.trycloudflare\.com', ln)
                if m and not TUNNEL['url']:
                    TUNNEL['url'] = m.group(0)
                if (m or 'Registered tunnel connection' in ln) and TUNNEL['status'] != 'online':
                    TUNNEL['status'] = 'online'
                    print('=' * 70)
                    print(' INTERNET    : %s' % (TUNNEL['url'] or '(your Cloudflare address)'))
                    print('=' * 70)
                    if TUNNEL['url']:
                        threading.Thread(target=github_publish, args=(TUNNEL['url'],), daemon=True).start()
            TUNNEL['status'] = 'stopped (exit %s)' % p.wait()
        except Exception as e:
            TUNNEL['status'] = 'error: %s' % e
    threading.Thread(target=run, daemon=True).start()


def tunnel_stop():
    p = TUNNEL.get('proc')
    if p and p.poll() is None:
        p.terminate()
    TUNNEL.update(proc=None, url='', status='off')
    threading.Thread(target=github_publish, args=('',), daemon=True).start()


# ---------------------------------------------------------------- fixed link: GitHub Pages page that forwards
LINK_PAGE = '''<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>Permit Server</title>
<meta name="robots" content="noindex">
<style>body{font:16px/1.5 "Segoe UI",system-ui,sans-serif;background:#EDF0F3;color:#1B2A3A;display:grid;place-items:center;min-height:100vh;margin:0}
div{background:#fff;border:1px solid #D3DAE2;border-radius:10px;padding:28px 32px;max-width:440px}h1{font-size:20px;margin:0 0 8px}
a{color:#0E6B6B;font-weight:600}.m{color:#5E6B7A;font-size:13px}</style></head><body><div>
<h1>Permit Server</h1>%(BODY)s<p class="m">Updated %(TIME)s</p></div>%(SCRIPT)s</body></html>
'''


def github_publish(url):
    """Write index.html to your GitHub Pages repo: forwards to the current internet address (or says offline)."""
    conf = load_conf()
    repo, token = (conf.get('github_repo') or '').strip().strip('/'), (conf.get('github_token') or '').strip()
    if not repo or not token:
        return
    import urllib.request, base64
    api = (conf.get('github_api') or 'https://api.github.com').rstrip('/')
    branch = conf.get('github_branch') or 'main'
    t = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    if url:
        body = '<p>Opening the server ...</p><p><a href="%s">%s</a></p>' % (url, url)
        script = '<script>location.replace(%s + location.hash)</script>' % json.dumps(url + '/')
    else:
        body = '<p>The server is offline right now. Try again later or ask the admin to start it.</p>'
        script = ''
    html = LINK_PAGE % dict(BODY=body, TIME=t, SCRIPT=script)
    hd = {'Authorization': 'Bearer ' + token, 'Accept': 'application/vnd.github+json', 'User-Agent': 'permit-server',
          'X-GitHub-Api-Version': '2022-11-28'}
    u = '%s/repos/%s/contents/index.html' % (api, repo)
    try:
        sha = None
        try:
            with urllib.request.urlopen(urllib.request.Request(u + '?ref=' + branch, headers=hd), timeout=20) as r:
                sha = json.load(r).get('sha')
        except Exception:
            pass
        data = dict(message='Permit Server address %s' % (url or 'offline'), branch=branch,
                    content=base64.b64encode(html.encode('utf-8')).decode())
        if sha:
            data['sha'] = sha
        req = urllib.request.Request(u, data=json.dumps(data).encode(), headers=dict(hd, **{'Content-Type': 'application/json'}),
                                     method='PUT')
        with urllib.request.urlopen(req, timeout=30) as r:
            r.read()
        TUNNEL['github'] = 'updated %s' % t
        print(' Fixed link  : %s  (GitHub page updated)' % github_link(conf))
    except Exception as e:
        TUNNEL['github'] = 'GitHub update failed: %s' % e
        print(' Fixed link  : GitHub update failed - %s' % e)


def github_link(conf=None):
    conf = conf or load_conf()
    repo = (conf.get('github_repo') or '').strip().strip('/')
    if '/' not in repo:
        return ''
    o, n = repo.split('/', 1)
    return 'https://%s.github.io/' % o.lower() if n.lower() == '%s.github.io' % o.lower() else 'https://%s.github.io/%s/' % (o.lower(), n)


def load_conf():
    c = dict(port=8080, host='0.0.0.0', open_browser=True, threads=32, workers=2, online=False, tunnel_token='', public_url='',
             github_repo='', github_token='', github_branch='main')
    try:
        with open(CONF, encoding='utf-8') as f:
            c.update(json.load(f))
    except Exception:
        with open(CONF, 'w', encoding='utf-8') as f:
            json.dump(c, f, indent=1)
    return c


def serve(conf):
    if not os.path.exists(os.path.join(WEB, 'index.html')):
        print('*** The web pages are missing (%s). Unzip the WHOLE PERMIT_SERVER.zip (folder "web" next to the program) '
              'and, for the exe, run BUILD_SERVER_EXE.bat again - otherwise the browser shows "Not Found".' % WEB)
    try:
        from waitress import serve as wserve
        th = max(int(conf.get('threads') or 32), 24)     # many people at once (each browser keeps a few requests open)
        print('Web server: waitress, %d threads' % th)
        wserve(app, host=conf['host'], port=conf['port'], threads=th, connection_limit=1000, max_request_body_size=300 * 1024 * 1024)
    except ImportError:
        from wsgiref.simple_server import make_server, WSGIServer, WSGIRequestHandler
        from socketserver import ThreadingMixIn

        class TS(ThreadingMixIn, WSGIServer):
            daemon_threads = True

        class Quiet(WSGIRequestHandler):
            def log_message(self, *a):
                pass
        print('Web server: built-in (threaded)')
        make_server(conf['host'], conf['port'], app, server_class=TS, handler_class=Quiet).serve_forever()


# ---------------------------------------------------------------- run in the background (no window)
STOP_KEY = os.path.join(DATA, 'stop.key')
PIDF = os.path.join(DATA, 'server.pid')
LOGF = os.path.join(DATA, 'server.log')


def _alive(port):
    import urllib.request
    try:
        with urllib.request.build_opener(urllib.request.ProxyHandler({})).open(
                'http://127.0.0.1:%d/api/me' % port, timeout=2) as r:
            return r.status == 200
    except Exception:
        return False


def _self_cmd():
    if getattr(sys, 'frozen', False):
        return [sys.executable]
    return [sys.executable, os.path.abspath(__file__)]


def online_running(port):
    """The server already runs (e.g. started in the background): switch its internet link on and show the address."""
    import urllib.request

    def call(start):
        key = open(STOP_KEY).read().strip()
        req = urllib.request.Request('http://127.0.0.1:%d/api/admin/online' % port,
                                     data=json.dumps({'key': key, 'start': start, 'quiet': not start}).encode(),
                                     headers={'Content-Type': 'application/json'}, method='POST')
        with urllib.request.build_opener(urllib.request.ProxyHandler({})).open(req, timeout=10) as r:
            return json.loads(r.read().decode())
    try:
        st = call(True)
    except Exception as e:
        print(' *** Could not switch the internet link on (%s).' % e)
        print('     Open http://localhost:%d > Settings > Internet address > "Go online",' % port)
        print('     or run STOP_SERVER.bat and then GO_ONLINE.bat again.')
        return 1
    print(' Switching the internet link on (the first time it downloads cloudflared, about 1 minute) ...')
    for _ in range(120):
        if st.get('status') == 'online' or str(st.get('status', '')).startswith(('error', 'stopped', 'cloudflared')):
            break
        time.sleep(1)
        try:
            st = call(False)
        except Exception:
            pass
    print('=' * 70)
    if st.get('status') == 'online':
        print(' INTERNET    : %s' % (st.get('url') or '(your Cloudflare address)'))
        print(' Send this address to your team. You can close this window now.')
    else:
        print(' Internet link: %s' % (st.get('status') or 'not started'))
        print(' Check Settings > Internet address on http://localhost:%d' % port)
    print('=' * 70)
    return 0 if st.get('status') == 'online' else 1


def start_background(conf, argv):
    """Start a copy of the server with no window, then this window can be closed."""
    import subprocess
    port = conf['port']
    if _alive(port):
        print(' The server is already running in the background: http://localhost:%d' % port)
        if '--online' in argv:
            rc = online_running(port)
            if conf.get('open_browser'):
                webbrowser.open('http://localhost:%d' % port)
            return rc
    else:
        args = _self_cmd() + [a for a in argv if a not in ('--background', '--open')] + ['--no-browser', '--child']
        log = open(LOGF, 'a', encoding='utf-8')
        kw = dict(stdout=log, stderr=log, stdin=subprocess.DEVNULL, cwd=app_dir(), close_fds=True)
        if os.name == 'nt':
            kw['creationflags'] = 0x00000008 | 0x00000200 | 0x08000000   # detached, own group, no window
        else:
            kw['start_new_session'] = True
        subprocess.Popen(args, **kw)
        print(' Starting the server in the background ...')
        for _ in range(60):
            time.sleep(1)
            if _alive(port):
                break
        else:
            print(' *** It did not start - see %s' % LOGF)
            return 1
        print('=' * 70)
        print(' %s is running in the BACKGROUND (no window).' % APP_NAME)
        print(' This PC     : http://localhost:%d' % port)
        RUN_PORT['port'] = port
        for u in addresses():
            print(' Other PCs   : %s' % u)
        print(' Stop it     : STOP_SERVER.bat (or Settings > Stop the server)')
        print(' Log         : %s' % LOGF)
        print(' You can close this window now.')
        print('=' * 70)
        if '--online' in argv:
            online_running(port)
    if conf.get('open_browser'):
        webbrowser.open('http://localhost:%d' % port)
    return 0


def _port_pids(port):
    """Process ids listening on this port (to stop a server whose stop.key is in another folder)."""
    import subprocess
    pids = set()
    try:
        if os.name == 'nt':
            out = subprocess.run(['netstat', '-ano', '-p', 'TCP'], capture_output=True, text=True,
                                 creationflags=0x08000000).stdout
            for ln in out.splitlines():
                p = ln.split()
                if len(p) >= 5 and (p[3].upper() == 'LISTENING' or p[2].endswith((':0', ':*'))) and p[1].rsplit(':', 1)[-1] == str(port):
                    pids.add(int(p[4]))
        else:
            out = subprocess.run(['lsof', '-t', '-iTCP:%d' % port, '-sTCP:LISTEN'], capture_output=True, text=True).stdout
            pids.update(int(v) for v in out.split() if v.isdigit())
    except Exception:
        pass
    pids.discard(os.getpid())
    return pids


def _proc_name(pid):
    import subprocess
    try:
        if os.name == 'nt':
            out = subprocess.run(['tasklist', '/FI', 'PID eq %d' % pid, '/FO', 'CSV', '/NH'], capture_output=True,
                                 text=True, creationflags=0x08000000).stdout
            return out.split(',')[0].strip().strip('"').lower()
        return open('/proc/%d/cmdline' % pid, 'rb').read().replace(b'\0', b' ').decode(errors='replace').lower()
    except Exception:
        return ''


def _kill(pid):
    import subprocess
    try:
        if os.name == 'nt':
            subprocess.run(['taskkill', '/F', '/T', '/PID', str(pid)], capture_output=True, creationflags=0x08000000)
        else:
            os.kill(pid, 15)
    except Exception:
        pass


def _wait_down(port, sec):
    for _ in range(sec):
        if not _alive(port):
            return True
        time.sleep(1)
    return not _alive(port)


def stop_server(conf):
    import urllib.request
    import urllib.error
    port = conf['port']
    if not _alive(port):
        print(' No server is running on port %d.' % port)
        return 0
    why = ''
    try:
        key = open(STOP_KEY).read().strip()
        req = urllib.request.Request('http://127.0.0.1:%d/api/admin/stop' % port, data=json.dumps({'key': key}).encode(),
                                     headers={'Content-Type': 'application/json'}, method='POST')
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))     # never through an office proxy
        opener.open(req, timeout=5).read()
    except FileNotFoundError:
        why = 'no stop.key in %s' % DATA
    except urllib.error.HTTPError as e:
        why = 'the server said %s (stop.key is from another folder?)' % e.code
    except Exception as e:
        why = str(e)
    if _wait_down(port, 10):
        print(' The server is stopped.')
        return 0
    # the polite way did not work (server started from another folder / older version): stop the process on the port
    if why:
        print(' Stop request not accepted: %s' % why)
    done = False
    for pid in _port_pids(port):
        name = _proc_name(pid)
        if 'permit_server' in name or 'python' in name or 'server.py' in name:
            print(' Stopping process %d (%s) ...' % (pid, name or '?'))
            _kill(pid)
            done = True
        else:
            print(' Port %d is used by another program: %s (pid %d) - not stopped.' % (port, name or '?', pid))
    if done and _wait_down(port, 10):
        print(' The server is stopped.')
        return 0
    print(' Could not stop it. Open Task Manager > Details > end PERMIT_SERVER.exe, or run in cmd:')
    print('     taskkill /F /IM PERMIT_SERVER.exe')
    return 1


def _shutdown_soon():
    def bye():
        time.sleep(0.8)
        try:
            tunnel_stop()
        except Exception:
            pass
        os._exit(0)
    threading.Thread(target=bye, daemon=True).start()


def main():
    conf = load_conf()
    for a in sys.argv[1:]:
        if a.startswith('--port='):
            conf['port'] = int(a.split('=', 1)[1])
        if a == '--no-browser':
            conf['open_browser'] = False
        if a == '--online':
            conf['online'] = True
    if os.environ.get('PORT'):                     # cloud hosts (Render, Railway, ...) give the port
        conf['port'] = int(os.environ['PORT'])
    if '--stop' in sys.argv:
        sys.exit(stop_server(conf))
    if '--background' in sys.argv:
        sys.exit(start_background(conf, sys.argv[1:]))
    if _alive(conf['port']):                       # already running (e.g. in the background) -> just open it
        print(' The server is already running: http://localhost:%d' % conf['port'])
        if '--online' in sys.argv:
            online_running(conf['port'])
        if conf.get('open_browser'):
            webbrowser.open('http://localhost:%d' % conf['port'])
        return
    RUN_PORT['port'] = conf['port']
    with open(STOP_KEY, 'w') as f:
        f.write(secrets.token_hex(16))
    with open(PIDF, 'w') as f:
        f.write(str(os.getpid()))
    RUN_PORT['since'] = now()
    RUN_PORT['background'] = '--child' in sys.argv
    init_db()
    try:
        T.load_plugins(PLUGIN_DIR)
    except Exception as e:
        print(' Plugins: %s' % e)
    if os.environ.get('PERMIT_ADMIN_PASSWORD') and not q('SELECT id FROM users LIMIT 1', one=True):
        un = (os.environ.get('PERMIT_ADMIN_USER') or 'admin').lower()     # cloud: first admin from settings
        x('INSERT INTO users(username, name, pw, role, created) VALUES(?,?,?,?,?)',
          (un, un, generate_password_hash(os.environ['PERMIT_ADMIN_PASSWORD']), 'admin', now()))
        print(' Admin user "%s" created from PERMIT_ADMIN_PASSWORD' % un)
    x("UPDATE jobs SET status='review', message='server restarted - run again if needed' "
      "WHERE status IN ('queued','extracting','generating')")
    nw = max(1, min(8, int(conf.get('workers') or 2)))      # documents read / Excel made at the same time
    for _ in range(nw):
        threading.Thread(target=worker, daemon=True).start()
    threading.Thread(target=lambda: (T.engine('danfoss'), T.lists('danfoss')), daemon=True).start()
    print('=' * 70)
    print(' %s %s' % (APP_NAME, VERSION))
    print(' Data folder : %s' % DATA)
    print(' This PC     : http://localhost:%d' % conf['port'])
    for u in addresses():
        print(' Other PCs   : %s' % u)
    print(' Running in the background - stop it with STOP_SERVER.bat' if '--child' in sys.argv else
          ' Keep this window open - closing it stops the server. (No window: START_SERVER_BACKGROUND.bat)')
    print('=' * 70)
    if conf.get('online'):
        print(' Internet    : starting the secure tunnel ...')
        tunnel_start()
    if conf.get('open_browser'):
        threading.Timer(1.5, lambda: webbrowser.open('http://localhost:%d' % conf['port'])).start()
    serve(conf)


if __name__ == '__main__':
    main()
